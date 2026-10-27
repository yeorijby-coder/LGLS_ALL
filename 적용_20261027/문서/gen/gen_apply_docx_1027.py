# -*- coding: utf-8 -*-
u"""2026-10-27 적용 안내서 - ★추가분만★ (Word).

   현장 보고(2026-10-27) : 작업 2161(입고) 화물이 입고 통로 PORT:12 까지 가 크레인에 실렸는데
   상태는 15, 단계표는 첫 줄 진행중. [완료처리] 를 누르면 "PORT:23 에 화물이 감지되지 않아 ..." 로만
   끝나 다음으로 넘길 수 없었다.

   앞 안내서(10-15 꾸러미)는 그대로 두고, 이 건만 담는다.
   실행 : python gen_apply_docx_1027.py
"""
import os, sys

HERE = os.path.dirname(os.path.abspath(__file__))        # 적용_20261027/문서/gen
DOCS = os.path.dirname(HERE)                             # 적용_20261027/문서
ROOT = os.path.dirname(os.path.dirname(DOCS))            # Renewal
sys.path.insert(0, os.path.join(ROOT, 'docs', '산출물_20260927', 'gen'))

from docx.shared import Pt, Cm, RGBColor
import gen_common as G

RED = RGBColor(0xB0, 0x2A, 0x20)
SHOT = os.path.join(DOCS, 'shots_1027')


def code(d, lines, size=9):
    t = d.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    c = t.rows[0].cells[0]
    G.shade(c, 'F4F5F7')
    c.text = ''
    for i, ln in enumerate(lines):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(ln)
        r.font.name = 'D2Coding'
        r.font.size = Pt(size)
    d.add_paragraph()


def note(d, text):
    p = d.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    r = p.add_run('※ ' + text)
    r.font.size = Pt(9.5)
    r.font.color.rgb = RED
    return p


G.DATE = '2026-10-27'
G.VER = '1.0'
G.BASE_HISTORY = []
G.FINAL_NOTE = ('2026-10-15 꾸러미에 이어지는 추가분. 작업정보 판넬 [완료처리] 가 화물이 있는 칸을 스스로 찾아 그 단계로 넘긴다.')

d = G.new_doc('2026-10-27 적용 안내서 (추가분)',
              '[완료처리] 경로 탐색 - 화물이 다른 칸에 가 있어도 넘길 수 있게 (현장 작업 2161)')

d.add_heading('0. 이 문서는 무엇인가', 1)
G.para(d, '2026-10-15 꾸러미(적용_20261015)를 이미 받아 보셨다면, 그 뒤에 더해진 것만 여기에 담았다. '
          '바뀐 것은 Client(EcsMain.exe) 한 가지이고 DB·TASK 는 손대지 않았다. 파일은 적용_20261027 폴더에 있다.')
note(d, '앞 문서들(09-29 ~ 10-15)의 내용은 그대로 유효하다. 이 문서는 그 위에 얹는 것이다.')

d.add_heading('1. 현장에서 생긴 일 (2026-10-27, 작업 2161 입고)', 1)
G.para(d, '화물은 입고 통로 PORT:12 까지 가서 크레인에 실렸는데, 작업 상태는 [15] 그대로이고 단계표는 첫 줄(0001 CONVEYOR:12 PORT:24→PORT:23)이 '
          '진행중이었다. 그 자리에서 [완료처리] 를 누르면 "작업 상태가 이미 이 구간의 완료 코드입니다. [15] 그런데 PORT:23 에 화물이 감지되지 않아 '
          '작업번호를 기록할 수 없습니다. 화물이 그 칸에 오면 다시 누르세요." 로 끝나, 다음으로 넘길 길이 없었다.')
d.add_heading('1.1 왜 그랬나', 2)
G.table(d, ['구분', '내용'], [
    ['[완료처리] 의 기준', '고른 단계표 줄(안 고르면 진행중 줄 = 첫 줄)의 ★도착 칸★ 하나만 본다. 그 칸에 화물이 없으면 "없다" 로 끝낸다.'],
    ['실제 화물 위치', '두 단계 뒤(0003 줄의 도착 칸 PORT:12). 첫 줄 기준 PORT:23 에는 당연히 없다.'],
    ['빠져 있던 것', '"그럼 어느 줄을 골라야 하나" 를 화면이 알려 주지 않았다. 10-15 시험안내서의 네 상황에는 "화물이 이미 뒤 칸에 가 있는" 경우가 없었다.'],
], widths=[3.5, 12.5])
note(d, '크레인이 데이터 없이 먼저 실은 것은 별개의 문제다(4절). 이 수정은 "그 뒤에 담당자가 작업을 이어 갈 수 있게" 하는 것이다.')

d.add_heading('2. 고친 것 - [완료처리] 가 화물이 있는 칸을 찾는다', 1)
G.para(d, '고른 줄의 도착 칸에 화물이 없으면, 그 작업 경로의 ★다른 칸들(단계표의 모든 도착 칸)★ 을 CV_DATA 에서 찾아본다. '
          '찾으면 그 단계로 넘기자고 묻고, [예] 한 번으로 상태 이동과 작업번호 기록까지 끝낸다. 어디에도 없으면 무엇을 해야 하는지 알려 준다.')
G.table(d, ['상황', '뜨는 창', '[예] 를 누르면'], [
    ['다른 칸에 내 작업번호가 붙어 있다, 또는 번호 없는 화물(센서만 ON)이 있다',
     '"대신 PORT:8 (단계 0003 CONVEYOR:4) 에 화물이 감지됩니다(작업번호 없음). 그 단계로 넘기고 PORT:8 에 작업번호 0758 를 기록할까요?"',
     '그 단계의 완료 코드로 JOB_STATUS 이동(예: 15→16) + 그 칸에 작업번호 기록(CV_DATA LUGG_NO_OD, TRACKING_WRITE_YN=Y). ★한 번만 묻는다★ (종전의 "기록할까요?" 를 다시 묻지 않는다).'],
    ['경로의 어느 칸에도 없다',
     '"경로의 어느 칸(PORT:31, PORT:7, PORT:8)에도 화물이 감지되지 않습니다. · 설비(크레인/RGV)에 실린 채라면 : PORT:8 에 내린 뒤 단계 0003 줄을 골라 [완료처리] · 이미 마지막까지(랙/작업대) 옮겨졌다면 : 단계 0004 (S/C 2) 줄을 골라 [완료처리]"',
     '(안내만, 아무것도 바꾸지 않는다)'],
    ['고른 줄의 도착 칸에 있다', '종전과 같다 ("작업번호를 기록할까요?")', '종전과 같다'],
], widths=[4.0, 7.5, 4.5], font=8.5)
G.para(d, '찾는 순서 : 내 작업번호가 붙은 칸 > 센서만 ON 인 빈 번호 칸. 여러 칸이면 ★뒤 단계★ 를 고른다(화물은 앞으로만 가므로). '
          '다른 작업의 번호가 붙은 칸은 세지 않는다. [아니오] 를 누르면 아무것도 바꾸지 않는다.')
G.image(d, os.path.join(SHOT, 'panel_seq_0758.png'), 13.0,
        '시험 : 0758(입고 130→902) 상태 15, 첫 줄 진행중. 화물은 PORT:31 이 아니라 PORT:8 에 있다 (CV_DATA 108 센서 ON, 번호 없음)')
G.image(d, os.path.join(SHOT, 'popup_route_found.png'), 11.0, '[완료처리] → 찾은 칸으로 넘기자고 묻는다')
G.image(d, os.path.join(SHOT, 'popup_route_none.png'), 11.0, '어느 칸에도 없을 때 - 무엇을 할지 알려 준다')
d.add_heading('2.1 확인한 것 (개발 PC, 2026-10-27)', 2)
G.table(d, ['확인', '결과'], [
    ['찾은 칸으로 넘기기', '[예] 한 번에 JOB_MST 0758 : 15 → 16, HS_TRACK_NO=107 / CV_DATA 108 : LUGG_NO_OD=0758, TRACKING_WRITE_YN=Y. 추가 창 없음.'],
    ['어느 칸에도 없음', '안내 창만 뜨고 DB 는 그대로(15).'],
    ['[아니오]', '아무것도 바뀌지 않는다.'],
], widths=[4.0, 12.0])

d.add_heading('3. 현장 2161 은 지금 어떻게 하나', 1)
G.table(d, ['Client', '하는 법'], [
    ['이 꾸러미(10-27)를 깐 뒤', '작업 2161 을 고르고 [완료처리] → "대신 PORT:12 (단계 0003 CONVEYOR:6) ..." 창에서 [예]. 화물이 아직 크레인 포크 위라면 2단 참조.'],
    ['아직 10-15 판이면', '단계표에서 ★0003 CONVEYOR:6 줄을 골라★ [완료처리] → 15→16 → "작업번호 기록할까요?" [예]. (첫 줄을 둔 채 누르면 종전처럼 "PORT:23 에 없다" 로 끝난다)'],
    ['화물이 크레인 포크 위에 있다', '먼저 화물을 PORT:12 에 내린다(크레인 수동). 그 뒤 위 순서. 포크 위에 둔 채로는 어느 판에서도 넘길 수 없다 - 포크 위 화물은 C/V 센서에 잡히지 않는다.'],
    ['이미 랙에 넣었다', '단계표 ★0004 S/C 3 줄을 골라★ [완료처리] → 29. 그 뒤 IO_TASK 가 HOST 보고까지 이어 간다. 랙 위치는 작업정보의 도착지(LOC)와 실제가 같은지 확인.'],
], widths=[4.0, 12.0])

d.add_heading('4. 왜 크레인이 데이터 없이 먼저 실었나 - 조사 부탁', 1)
G.para(d, '정상 순서는 "통로 칸에 작업번호가 붙음(트래킹) → IO_TASK 가 크레인 지시 → 크레인이 실음" 이다. 현장은 번호가 붙기 전에 실렸다. '
          '크레인이 수동으로 움직였는지, PLC 가 스스로 당겼는지, 아니면 번호는 붙었는데 상태만 안 넘어갔는지는 아래 결과를 보아야 안다. '
          '현장 DB(LGLS_IO)에서 다음을 돌려 결과(화면 캡처나 텍스트)를 보내 주시면 분석한다.')
code(d, [
    "-- (1) 2161 작업",
    "SELECT * FROM JOB_MST WHERE lugg_no = '2161'",
    "",
    "-- (2) 경로 칸의 트래킹 상태 (입고대 24/23, 입고 통로 11/12)",
    "SELECT mc_no, sensor0_data_rd, lugg_no_rd, lugg_no_od, tracking_write_yn, write_upd_dt",
    "  FROM CV_DATA WHERE mc_no IN ('124','123','111','112')",
    "",
    "-- (3) 크레인 3호기",
    "SELECT sc_no, ucstatus_rd, auto_mode_rd, job_typ_rd, job_typ_od, lugg_no_fk1_rd, lugg_no_fk1_od,",
    "       pallet_on_vehicle_rd, pallet_id_od, transfer_request_od, transfer_ack_rd, load_complete_rd,",
    "       from_01_od, to_01_od, read_upd_dt, od_upd_dt",
    "  FROM SC_DATA_LGLS WHERE sc_no IN ('3','03')",
    "",
    "-- (4) 2161 이 남긴 프로그램 로그 (시간순)",
    "SELECT ins_dt, pgr_nm, track_from, track_to, job_sta, log_kor",
    "  FROM WCS_LOG_PGR WHERE lugg_no = '2161' ORDER BY ins_dt",
])
G.para(d, 'SSMS 에서 LGLS_IO 를 고르고 붙여 넣어 실행하면 된다. sqlcmd 라면 :  sqlcmd -S <서버> -d LGLS_IO -E -W -s "|" -i <파일>')

d.add_heading('5. 바뀐 파일 / 적용 순서', 1)
G.table(d, ['폴더', '파일', '비고'], [
    ['Client', 'EcsMain.exe, DciLib.dll, EcsLib.dll, XmlLib.dll', '★넷을 같이★ 복사 (빌드 2026-10-27 10:49). 10-15 의 Ecs.exe(런처)·설정·EcsDefine.xml 은 그대로'],
    ['소스_변경분', 'PanelJobDlg.cpp / .h + .diff', 'RouteSearchOffer, DoCompleteStep 분리'],
    ['문서', '이 문서, shots_1027', ''],
], widths=[3.0, 7.0, 6.0])
G.numbered(d, [
    'Client 를 닫는다.',
    'Client\\ 의 네 파일을 깔린 폴더(예 C:\\LGLS\\CLIENT)에 덮어쓴다. ini 는 손대지 않는다 (새 설정 없음).',
    '다시 띄워 제목줄 [Build:2026.10.27 10:49:52] 를 확인한다.',
    '작업 2161 을 3절대로 처리한다.',
])
note(d, 'Upload/Download 를 쓰고 있다면 Upload 로 EcsMain.exe 와 DLL 3개를 올려 각 PC 의 런처가 받게 해도 된다 (2026-10-15_Upload_사용매뉴얼).')

out = os.path.join(DOCS, '2026-10-27_적용안내서_추가분.docx')
d.save(out)
print(out)
