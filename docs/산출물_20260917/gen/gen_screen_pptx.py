# -*- coding: utf-8 -*-
"""ECS 화면설계서 (pptx) - 2026-09-17 판.
   09-09 판(docs/산출물_20260909/gen/gen_screen_pptx.py)과 같은 틀 : 표지 / Document history / Contents / 화면별(캡처+설명).
   바뀐 점 : 메인 화면을 전체화면(1920x1080) 캡처로 한 장 가득, 왼쪽 고정 칸·리본·설비 그림 장 추가, Ecs.ini 설정·변경 이력 장 추가."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
from gen_screens import SCREENS

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOTS = os.path.join(BASE, 'shots')
DATE = '2026-09-17'
NAVY = RGBColor(0x1E, 0x27, 0x61); GRAY = RGBColor(0x59, 0x59, 0x59); LIGHT = RGBColor(0xEE, 0xF1, 0xF7); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]; FONT = '맑은 고딕'

def tb(slide, x, y, w, h, text, size=14, bold=False, color=GRAY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = s.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        if isinstance(ln, tuple): txt, b, sz, col = ln
        else: txt, b, sz, col = ln, bold, size, color
        r = p.add_run(); r.text = txt; r.font.size = Pt(sz); r.font.bold = b; r.font.color.rgb = col; r.font.name = FONT
    return s

page = [0]
def footer(slide):
    page[0] += 1
    tb(slide, 0.5, 7.05, 5, 0.35, 'LG 화학 1동 ECS  |  화면설계서  |  Ver 1.2', 10, False, GRAY)
    tb(slide, 12.2, 7.05, 0.8, 0.35, str(page[0]), 10, False, GRAY, PP_ALIGN.RIGHT)

def title(slide, text):
    tb(slide, 0.5, 0.3, 12.3, 0.7, text, 26, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)

def pic(slide, name, x, y, maxw, maxh, border=True):
    p = os.path.join(SHOTS, name)
    if not os.path.exists(p):
        tb(slide, x, y, maxw, 0.4, '[캡처 없음 : %s]' % name, 11, False, GRAY); return
    w, h = Image.open(p).size
    sc = min(maxw / (w / 96.0), maxh / (h / 96.0)); dw, dh = w / 96.0 * sc, h / 96.0 * sc
    px = x + (maxw - dw) / 2
    s = slide.shapes.add_picture(p, Inches(px), Inches(y), Inches(dw), Inches(dh))
    if border: s.line.color.rgb = RGBColor(0xC8, 0xCE, 0xDA)
    return dh

def add_table(slide, x, y, w, rows, colw, size=11, rowh=0.36):
    n, m = len(rows), len(rows[0])
    shp = slide.shapes.add_table(n, m, Inches(x), Inches(y), Inches(w), Inches(rowh * n)); t = shp.table
    for j, cw in enumerate(colw): t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(rowh)
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = ''; c.margin_top = c.margin_bottom = Inches(0.03)
            p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = str(v); r.font.size = Pt(size); r.font.name = FONT
            c.fill.solid()
            if i == 0: c.fill.fore_color.rgb = NAVY; r.font.color.rgb = WHITE; r.font.bold = True
            else: c.fill.fore_color.rgb = WHITE if i % 2 else LIGHT; r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return t

def bullets(slide, x, y, w, h, head, items, size=11):
    lines = [(head, True, 13, NAVY)] + [('· ' + t, False, size, GRAY) for t in items]
    tb(slide, x, y, w, h, lines)

# ── 표지 ──
s = prs.slides.add_slide(blank); page[0] += 1
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height); bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
tb(s, 1.0, 2.4, 11, 0.8, 'LG 화학 1동 자동창고', 22, False, RGBColor(0xCA, 0xDC, 0xFC))
tb(s, 1.0, 3.2, 11, 1.2, 'ECS 화면설계서', 48, True, WHITE)
tb(s, 1.0, 4.6, 11, 0.6, 'ECS Renewal   ·   Ver 1.2   ·   %s' % DATE, 16, False, RGBColor(0xCA, 0xDC, 0xFC))
tb(s, 1.0, 5.3, 11, 0.6, '기준 : EXE_NEWUI\\WCS_CLIENT\\Ecs.exe (Build 2026.09.17 08:18:34) + Ecs.ini', 12, False, RGBColor(0xCA, 0xDC, 0xFC))

# ── Document history ──
s = prs.slides.add_slide(blank); title(s, 'Document history')
add_table(s, 0.7, 1.4, 11.9, [
    ['Vers.', 'Date', 'Author', 'Approver', 'Notes'],
    ['V1.0', '2026-09-03', 'LGLS ECS Renewal', '', '최초 작성'],
    ['V1.1', '2026-09-09', 'LGLS ECS Renewal', '', '명칭 ECS 확정, 범례 화면 내장, 판넬 보기 신설, 캡처 갱신'],
    ['V1.2', DATE, 'LGLS ECS Renewal', '', '메인 화면 새 배치(MAIN_UI=2, 왼쪽 고정 칸) 전체화면 캡처, 리본 메뉴 정리(INI 열기·반자동 TEST·확대 숨김),\n크레인·RGV 색 범례 정합, 레이아웃 그룹 이름 정정, Ecs.ini 설정 장 추가'],
], [1.2, 1.8, 2.6, 1.8, 4.5], rowh=0.5)
footer(s)

# ── Contents ──
FIXED = [('메인', ['메인 화면 (전체화면)', '메인 화면 영역 구성', '리본 메뉴', '왼쪽 고정 칸 ① 통신 상태 · 범례', '왼쪽 고정 칸 ② 설비 반송 · 작업정보', '설비 그림 · 마우스 동작'])]
groups = []
for g, *_ in SCREENS:
    if g not in groups: groups.append(g)
toc = FIXED + [(g, [x[1] for x in SCREENS if x[0] == g]) for g in groups] + [('부록', ['Ecs.ini 화면 설정', '설계 결정 · 변경 이력'])]
s = prs.slides.add_slide(blank); title(s, 'Contents')
cols = [[], [], []]; cnt = [0, 0, 0]
for g, items in toc:
    ci = cnt.index(min(cnt)); cols[ci].append((g, items)); cnt[ci] += len(items) + 2
for ci, col in enumerate(cols):
    lines = []
    for g, items in col:
        lines.append((g, True, 14, NAVY))
        for it in items: lines.append(('   ' + it, False, 11, GRAY))
        lines.append(('', False, 6, GRAY))
    tb(s, 0.6 + ci * 4.15, 1.2, 4.0, 5.7, lines)
footer(s)

# ── 메인 화면 (전체화면 한 장) ──
s = prs.slides.add_slide(blank)
tb(s, 0.3, 0.12, 12.7, 0.5, '메인 화면 (창고 모니터링) - 전체화면 1920×1080', 20, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
pic(s, 'main_full.png', 0.3, 0.68, 12.73, 6.3)
footer(s)

# ── 메인 화면 영역 구성 ──
s = prs.slides.add_slide(blank); title(s, '메인 화면 영역 구성')
pic(s, 'main_full.png', 0.5, 1.2, 5.6, 3.2)
bullets(s, 0.5, 4.3, 5.6, 2.6, '■ 프로그램 개요', [
    '1동 자동창고 전체 배치(랙 10열, 크레인 5대, RGV 1대, C/V#2~#15, 작업대)와 작업 진행을 실시간 표시한다.',
    '실행 : Ecs.exe. 기동 시 [USER] DEFAULT_ID 로 자동 로그인한다.',
    '설비 데이터는 [ETC] COLLECT_INTERVAL_MS=300 주기로 DB 에서 읽는다.'], 10.5)
add_table(s, 6.4, 1.2, 6.5, [
    ['영역', '내용'],
    ['① 제목줄', 'Ecs V1.0 [Build 시각] [DB:DB명@서버][실행 경로] - 길면 흘러감'],
    ['② 리본', 'ECS / MANUAL / LOG 탭 + 오른쪽 끝 통신(EQUIP·HOST·SCH)'],
    ['③ 통신 상태', '설비 통신 / 상위 통신 / 스케줄러 - 정상·느림·끊김'],
    ['④ 범례', '작업 색상 · C/V 상태 · S/C·RGV 상태 · S/C·RGV 레일 (28항목)'],
    ['⑤ 설비 반송', 'S/C#1~5, RGV 현재 반송 작업 (2초 갱신)'],
    ['⑥ 작업정보', '진행 중 작업 목록, 탭 7개, 자동 갱신 (2초)'],
    ['⑦ 설비 그림', 'C/V·크레인·RGV·랙·작업대·입출고 모드 (EcsLayout1.xml)'],
], [1.5, 5.0], size=10.5, rowh=0.42)
tb(s, 6.4, 4.8, 6.5, 1.8, [('■ 크기 조절', True, 13, NAVY),
    ('· 칸 사이 손잡이(점선)를 끌면 크기가 바뀌고 Ecs.ini MAIN_UI_LEFT_W / COMM_H / LEG_H / VEH_H 에 저장된다.', False, 10.5, GRAY),
    ('· 창 크기를 바꾸면 왼쪽 칸 폭도 비례해 다시 저장된다. 기준값 : 1920×1080 최대화에서 LEFT_W=531.', False, 10.5, GRAY),
    ('· MAIN_UI=2(왼쪽 고정 칸 배치)는 기동 시 한 번 읽는다.', False, 10.5, GRAY)])
footer(s)

# ── 리본 ──
s = prs.slides.add_slide(blank); title(s, '리본 메뉴')
y = 1.05
for nm, cap in (('ribbon_ecs_l.png', 'ECS 탭'), ('ribbon_manual_l.png', 'MANUAL 탭'), ('ribbon_log_l.png', 'LOG 탭')):
    tb(s, 0.5, y + 0.3, 1.3, 0.3, cap, 12, True, NAVY)
    pic(s, nm, 1.8, y, 7.0, 0.9); y += 0.95
tb(s, 9.4, 1.35, 2.0, 0.3, '통신 (모든 탭 오른쪽 끝)', 12, True, NAVY)
pic(s, 'ribbon_comm.png', 9.4, 1.7, 1.2, 0.9)
add_table(s, 0.5, 4.0, 12.4, [
    ['탭', '그룹', '버튼', '비고'],
    ['ECS', '환경설정', '통신 연결 정의 · 로그 삭제 설정 · 범례', 'INI 열기(INI_MENU=0) · 사용자(USER_MENU=0) 숨김'],
    ['ECS', '뷰 / 창고 모니터링', '작업정보 · 찾기 / 작업번호 · 트랙번호 · 제품정보', '판넬 보기 · 처리 그룹 숨김(PANEVIEW_MENU / PROCESS_MENU)'],
    ['MANUAL', '수동조작', '작업 · 크레인 · RTV', '반자동 TEST 그룹 숨김(SEMITEST_MENU=0)'],
    ['LOG', '로그 / 알람', '작업로그 · HOST로그 · 설비에러이력 · 유저사용로그 · ECS프로그램로그 / 알람', ''],
    ['공통', '통신', 'EQUIP · HOST · SCH', '파랑=정상, 빨강=단절. 단절 시 클릭하면 핑·포트 점검 안내 (STATUS_POS=RIBBON, COMM_MODE=2)'],
], [1.2, 2.2, 4.8, 4.2], size=10.5, rowh=0.44)
footer(s)

# ── 왼쪽 고정 칸 ① ──
s = prs.slides.add_slide(blank); title(s, '왼쪽 고정 칸 ① 통신 상태 · 범례')
pic(s, 'left_comm.png', 0.5, 1.2, 5.8, 1.3)
pic(s, 'left_legend.png', 0.5, 2.7, 5.8, 3.6)
add_table(s, 6.7, 1.2, 6.2, [
    ['표시', '조회 대상(5초)', '정상 / 느림'],
    ['설비 통신', 'CV_DATA.READ_UPD_DT', '10초 / 60초'],
    ['상위 통신', 'HOST_IF_LOG.INS_DT', '60초 / 300초'],
    ['스케줄러', 'JOB_MST.UPD_DT', '120초 / 900초'],
], [1.5, 2.8, 1.9], size=10.5)
bullets(s, 6.7, 2.9, 6.2, 4.0, '■ 범례 (28항목, 색 = Ecs.ini [USER] USER_COLOR_*)', [
    '작업 색상 10 : 입고 · 출고 · 이동 · 랙투랙 · 호기간이동 + 반자동 5종',
    'C/V 상태 10 : 입고대 · 출고대 · 입고 HS · 출고 HS · 일시정지 · 에러 · 수동 · 통신두절 · 검색 · 작업번호 있음',
    'S/C·RGV 상태 3 : 작업없음(정상) · 미가동(수동) · 에러 - 크레인·RGV 포크 색이 이 범례와 같다',
    'S/C·RGV 레일 5 : 입고 금지 · 출고 금지 · 입출고 정지 · 레일 에러 · 작업중',
    '색은 ECS > 환경설정 > 범례 에서 바꾸며 저장 즉시 반영된다.'], 10.5)
footer(s)

# ── 왼쪽 고정 칸 ② ──
s = prs.slides.add_slide(blank); title(s, '왼쪽 고정 칸 ② 설비 반송 · 작업정보')
pic(s, 'left_veh.png', 0.5, 1.2, 5.8, 1.9)
pic(s, 'left_job.png', 0.5, 3.3, 5.8, 3.5)
bullets(s, 6.7, 1.2, 6.2, 2.2, '■ 설비 반송 (2초 갱신)', [
    '행 : S/C#1 ~ S/C#5, RGV',
    '열 : 작업번호 · 자재 · 팔렛 · 출발(스테이션 + 위치) · 도착(스테이션 + 위치)',
    '크레인·RGV 가 지금 들고 있는 작업을 한눈에 본다.'], 10.5)
bullets(s, 6.7, 3.4, 6.2, 3.4, '■ 작업정보 (자동 갱신 2초)', [
    '탭 : 전체 · 입고 · 출고 · 피킹출고 · 랙투랙 · 호기간이동 · 이동',
    '열 : 작업번호 · 구분 · 상태 · 출발 · 출발위치 · 도착 · 도착위치 · LOT · 제품 · 우선 · 수정시각',
    '상태 표기 예 : [15] CV 구동중 / [16] 통로CV 구동중 / [25] SC 구동중 / [35] RGV 구동중',
    '조건 조회·수정·삭제는 ECS > 뷰 > 작업정보 창에서 한다.'], 10.5)
footer(s)

# ── 설비 그림 ──
s = prs.slides.add_slide(blank); title(s, '설비 그림 · 마우스 동작')
pic(s, 'layout.png', 0.5, 1.2, 5.6, 5.6)
add_table(s, 6.4, 1.2, 6.5, [
    ['동작', '대상', '결과'],
    ['클릭', 'C/V 트랙', 'CV 상태 창'],
    ['클릭', '크레인 (가로 레일 위 칸)', 'SC 상태 창'],
    ['클릭', 'RGV (세로 레일 위 칸)', 'RTV 상태 창'],
    ['더블클릭', '입고 모드 / 출고 모드 글자', '방향 전환 요청 (확인창)'],
    ['우클릭 · 휠', '-', '사용하지 않음 (추후 정의)'],
], [1.4, 2.8, 2.3], size=10.5)
bullets(s, 6.4, 3.6, 6.5, 3.3, '■ 입출고 모드 전환 (더블클릭)', [
    '대상 : C/V#11 입출고 겸용대(트랙 122) / C/V#2 S/C#1 통로(트랙 103)',
    '확인창에 PLC 방향 워드 주소를 보여 준다 : DIR_ADDR_CV11=310, DIR_ADDR_CV2=301',
    '확인하면 설비 통신이 즉시 방향을 바꾸고(DIRM), 스케줄러는 MANUAL_DIR_HOLD_SEC(기본 120초) 동안 자동 전환을 멈춘다.',
    '크레인·RGV 색 : 에러 → 미가동(범례 수동색) → 작업없음(범례 정상색) → 작업구분 색. 색을 지우는 시점은 VEH_CLEAR_MODE=3.',
    '트랙 글자는 리본 [창고 모니터링] 선택(작업번호/트랙번호/제품정보)을 따른다.'], 10.5)
footer(s)

# ── 화면별 ──
def screen_slide(g, name, sh, path, overview, notes):
    s = prs.slides.add_slide(blank); title(s, '%s - %s' % (g, name))
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.3), Inches(8.4), Inches(5.5)); box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xF7, 0xF8, 0xFA); box.line.color.rgb = RGBColor(0xC8, 0xCE, 0xDA)
    p = os.path.join(SHOTS, sh) if sh else None
    if p and os.path.exists(p):
        w, h = Image.open(p).size; maxw, maxh = 8.2, 5.3
        sc = min(maxw / (w / 96.0), maxh / (h / 96.0), 1.0); dw, dh = w / 96.0 * sc, h / 96.0 * sc
        s.shapes.add_picture(p, Inches(0.6 + (8.4 - dw) / 2), Inches(1.3 + (5.5 - dh) / 2), Inches(dw), Inches(dh))
    else:
        tb(s, 0.8, 3.7, 8.0, 0.6, '(캡처 없음 - 오른쪽 설명 참조)', 14, False, GRAY, PP_ALIGN.CENTER)
    lines = [('■ 메뉴 Path', True, 13, NAVY), (path, False, 11, GRAY), ('', False, 6, GRAY), ('■ 프로그램 개요', True, 13, NAVY), (overview, False, 11, GRAY)]
    if notes:
        lines += [('', False, 6, GRAY), ('■ 특기사항', True, 13, NAVY)] + [('· ' + n, False, 10, GRAY) for n in notes]
    tb(s, 9.3, 1.3, 3.6, 5.5, lines)
    footer(s)

for g, name, sh, path, overview, notes in SCREENS:
    screen_slide(g, name, sh, path, overview, notes)

# ── 부록 : Ecs.ini ──
s = prs.slides.add_slide(blank); title(s, '부록 - Ecs.ini 화면 설정 (현재값)')
add_table(s, 0.5, 1.15, 12.4, [
    ['섹션 · 키', '값', '의미', '반영'],
    ['[MENU] MAIN_UI', '2', '1=종전(범례 큰 표) / 2=왼쪽 고정 칸', '재기동'],
    ['[MENU] MAIN_UI_LEFT_W / COMM_H / LEG_H / VEH_H', '531 / 0 / 299 / 159', '왼쪽 칸 크기(px), 0=기본', '자동 저장'],
    ['[MENU] INI_MENU / USER_MENU / SEMITEST_MENU', '0 / 0 / 0', 'INI 열기 · 사용자 · 반자동 TEST 메뉴', '재기동'],
    ['[MENU] PROCESS_MENU / AUTOTIME_MENU / UIMODE_MENU', '0 / 0 / 1', '처리 그룹 · 시간 기반 자동 처리 · 판넬 보기(MAIN_UI=2 이면 숨김)', '재기동'],
    ['[MENU] PRODINFO_MENU / SCWAIT_VIEW', '1 / 0', '제품정보 버튼 · 크레인별 출고 대기 칸', '재기동'],
    ['[MENU] ZOOM_BTN', '0', '설비 대화상자 확대 버튼', '즉시'],
    ['[MENU] STATUS_POS / COMM_MODE', 'RIBBON / 2', '통신 표시 위치 · 모양', '재기동'],
    ['[MENU] LOADBIT_GATE / VEH_CLEAR_MODE', '1 / 3', '크레인·RGV 색 판정(차상 화물 비트) · 색 지우는 시점', '즉시'],
    ['[MENU] DIR_ADDR_CV11 / DIR_ADDR_CV2', '310 / 301', '방향 전환 확인창의 PLC 주소', '재기동'],
    ['[MENU] UI_TRACE', '1', '클릭 경로 로그(LOG\\ECS_UI_yyyymmdd.log)', '즉시'],
    ['[Title] BuildDate / DbInfo / Path', '1 / 1 / 1', '제목줄 표시 항목', '즉시'],
    ['[RibbonMenu] ToolTip', '1', '리본 툴팁에 리소스 경로 표시', '즉시'],
    ['[USER] JOB_STALL_WARN_SEC', '300', '경고창 체류 기준(초)', '즉시'],
    ['[USER] USER_COLOR_*', '범례 창에서 저장', '범례 · 설비 색 (BGR)', '저장 시'],
    ['[ETC] COLLECT_INTERVAL_MS / ViewRetCnt', '300 / 1', '설비 데이터 조회 주기 · 출고 잔여 수 표시', '재기동 / 즉시'],
], [4.4, 2.0, 4.8, 1.2], size=10, rowh=0.36)
footer(s)

# ── 부록 : 변경 이력 ──
s = prs.slides.add_slide(blank); title(s, '부록 - 설계 결정 · 변경 이력 (V1.2)')
add_table(s, 0.5, 1.2, 12.4, [
    ['#', '항목', '결정 / 조치'],
    ['1', '메인 화면 배치', 'MAIN_UI=2 : 왼쪽에 통신 상태 · 범례 · 설비 반송 · 작업정보를 고정. 종전 레이아웃 안 범례 표와 [판넬 보기] 는 쓰지 않는다'],
    ['2', 'INI 열기 버튼', '표시하지 않음 (INI_MENU=0)'],
    ['3', '반자동 TEST 메뉴 그룹', '표시하지 않음 (SEMITEST_MENU=0)'],
    ['4', '설비 대화상자 확대 버튼', '표시하지 않음 (ZOOM_BTN=0)'],
    ['5', '크레인·RGV 색', '범례 기준으로 정합 : 작업없음 = 범례 정상색(RGB 224,224,224), 미가동 = 범례 수동색'],
    ['6', '레이아웃 그룹 이름', 'EcsLayout1.xml 의 SC / RTV 이름을 실제 설비에 맞게 정정 (동작 영향 없음)'],
    ['7', '메인 화면 우클릭 · 휠', '현재 사용하지 않음. 추후 기능 정의'],
    ['8', '리본 [작업정보]', '별도 작업정보 창을 여는 것이 의도된 동작'],
    ['9', '크레인별 출고 대기 칸', '표시하지 않음 (SCWAIT_VIEW=0)'],
    ['10', '화면 반응 속도', '설비 통신 한 바퀴 3초 이상 → 약 1초 (WCS_TASK_CV 개선, 2026-09-17)'],
], [0.6, 3.0, 8.8], size=10.5, rowh=0.42)
tb(s, 0.5, 6.35, 12.4, 0.5, '캡처 : 메인 · 리본 · CV/SC/RTV 상태 · 범례 = 2026-09-17 전체화면(1920×1080). 그 밖의 대화상자 = 2026-09-09 판 캡처(구성 변경 없음).', 10, False, GRAY)
footer(s)

out = os.path.join(BASE, '06_ECS_화면설계서.pptx'); prs.save(out); print('saved', out, 'slides', len(prs.slides))
