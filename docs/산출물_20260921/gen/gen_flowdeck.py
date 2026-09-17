# -*- coding: utf-8 -*-
"""ECS 구성도 및 흐름도 (pptx) - 2026-09-21 최종판 (Ver 2.0)
시스템 구성도 / 입고·출고 플로우차트 / 플로우 다이어그램 / C/V#2 통로 사용권 / 겸용대 보호 / 알람 흐름 / 예외
근거 : TASK/IO_TASK/CLS/Thread/cThread_SCH.cs, TASK/WCS_TASK_HOST/CCliWork.cs·CSrvWork.cs,
       TASK/WCS_TASK_CV_BIN/2_Thread/EQP_THREAD/CvThread.cs·VehThread.cs (2026-09-17 기준)
시뮬레이터 관련 내용은 넣지 않는다.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

OUT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATE = '2026-09-21'
VER = '2.0'
NAVY  = RGBColor(0x1E, 0x27, 0x61)
BLUE  = RGBColor(0x26, 0x6E, 0xBE)
ICE   = RGBColor(0xCA, 0xDC, 0xFC)
LIGHT = RGBColor(0xEE, 0xF1, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY  = RGBColor(0x59, 0x59, 0x59)
LINE  = RGBColor(0xB6, 0xBF, 0xD0)
GREEN = RGBColor(0x2E, 0x86, 0x4B)
AMBER = RGBColor(0xC8, 0x7A, 0x0E)
RED   = RGBColor(0xB0, 0x2E, 0x2E)
FONT  = '맑은 고딕'

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
page = [0]


def tb(slide, x, y, w, h, text, size=12, bold=False, color=GRAY,
       align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = s.text_frame; tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = 0
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if isinstance(ln, tuple): txt, b, sz, col = ln
        else: txt, b, sz, col = ln, bold, size, color
        r = p.add_run(); r.text = txt
        r.font.size = Pt(sz); r.font.bold = b; r.font.color.rgb = col; r.font.name = FONT
    return s


def box(slide, x, y, w, h, text, fill=WHITE, edge=NAVY, txt=NAVY, size=11,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, bold=True):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = edge; sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(text if isinstance(text, list) else [text]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        if isinstance(ln, tuple): t, sz, b, c = ln
        else: t, sz, b, c = ln, size, bold, txt
        r = p.add_run(); r.text = t
        r.font.size = Pt(sz); r.font.bold = b; r.font.color.rgb = c; r.font.name = FONT
    return sh


def arrow(slide, x1, y1, x2, y2, color=BLUE, width=1.5, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    el = c.line._get_or_add_ln()
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls
    tail = parse_xml('<a:tailEnd %s type="triangle" w="med" len="med"/>' % nsdecls('a'))
    el.append(tail)
    if dash:
        d = parse_xml('<a:prstDash %s val="dash"/>' % nsdecls('a'))
        el.insert(0, d)
    return c


def label(slide, x, y, w, text, size=9, color=GRAY, align=PP_ALIGN.CENTER):
    return tb(slide, x, y, w, 0.24, text, size, False, color, align)


def footer(slide, title):
    page[0] += 1
    tb(slide, 0.5, 7.02, 7, 0.3, 'LG 화학 1동 ECS  |  ' + title, 9, False, GRAY)
    tb(slide, 12.3, 7.02, 0.6, 0.3, str(page[0]), 9, False, GRAY, PP_ALIGN.RIGHT)


def head(slide, text, sub=None):
    tb(slide, 0.5, 0.3, 12.3, 0.5, text, 24, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
    if sub:
        tb(slide, 0.5, 0.82, 12.3, 0.3, sub, 11, False, GRAY)


def grid(slide, x, y, widths, rows, font=9, rowh=0.26):
    """rows[0] = 머리행"""
    shp = slide.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y),
                                 Inches(sum(widths)), Inches(rowh * len(rows)))
    t = shp.table
    for j, w in enumerate(widths):
        t.columns[j].width = Inches(w)
    for i2, row in enumerate(rows):
        for j2, v in enumerate(row):
            c = t.cell(i2, j2); c.text = ''
            c.margin_top = c.margin_bottom = Inches(0.02)
            p2 = c.text_frame.paragraphs[0]; r = p2.add_run(); r.text = v
            r.font.size = Pt(font); r.font.name = FONT
            c.fill.solid()
            if i2 == 0:
                c.fill.fore_color.rgb = NAVY; r.font.color.rgb = WHITE; r.font.bold = True
            else:
                c.fill.fore_color.rgb = WHITE if i2 % 2 else LIGHT
                r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return t


def vchain(slide, x, y, w, items, h=0.50, gap=0.16, size=9.5):
    """세로 흐름 : items = [(text|list, fill, edge, side_note|None)]. side_note 는 오른쪽 보류 표시."""
    for i, (t, f, e, note) in enumerate(items):
        box(slide, x, y, w, h, t if isinstance(t, list) else [t], f, e, NAVY, size)
        if note:
            arrow(slide, x + w, y + h / 2, x + w + 0.35, y + h / 2, RED, 1.25)
            tb(slide, x + w + 0.38, y + h / 2 - 0.13, 1.25, 0.3, note, 8.5, True, RED)
        if i < len(items) - 1:
            arrow(slide, x + w / 2, y + h, x + w / 2, y + h + gap, BLUE, 1.25)
        y += h + gap
    return y


# ══════════ 1. 표지 ══════════
s = prs.slides.add_slide(blank); page[0] += 1
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background(); bg.shadow.inherit = False
tb(s, 1.0, 2.5, 11, 0.6, 'LG 화학 1동 자동창고', 20, False, ICE)
tb(s, 1.0, 3.15, 11, 1.1, 'ECS 시스템 구성도 및 업무 흐름도', 40, True, WHITE)
tb(s, 1.0, 4.5, 11, 0.5, 'ECS Renewal   ·   작성일 %s   ·   Ver %s' % (DATE, VER), 14, False, ICE)
tb(s, 1.0, 5.0, 11, 0.5, '최종판 - 현장 기본 시험 완료 기준', 12, False, ICE)

# ══════════ 2. 문서 이력 ══════════
s = prs.slides.add_slide(blank)
head(s, '문서 이력')
grid(s, 0.5, 1.3, [1.3, 1.8, 2.6, 6.6], [
    ['버전', '일자', '작성', '내용'],
    ['1.0', '2026-09-05', 'LGLS ECS Renewal', '최초 작성 (시스템 구성도, 입고·출고 플로우차트/다이어그램, 피킹, 예외) - 파일명 07_WCS_구성도_및_흐름도'],
    [VER, DATE, 'LGLS ECS Renewal', '최종판 - 명칭 ECS 통일, 작업상태 전이 현행화(11/21/30/31 폐기, 16 추가, 입고 25→09), '
                                    'C/V#2 통로 사용권·겸용대 보호·알람 흐름 추가, 이중입고/공출고 작업상태 처리 미사용 반영']],
     font=10, rowh=0.45)
footer(s, '문서 이력')

# ══════════ 3. 시스템 구성도 ══════════
s = prs.slides.add_slide(blank)
head(s, '시스템 구성도', '모든 프로그램은 DB(LGLS_MCS_IO)를 매개로 구동한다. 화살표는 데이터가 흐르는 방향이다.')

box(s, 1.0, 2.20, 2.4, 0.80, ['상위 시스템', ('WMS / IMS', 10, False, GRAY)], LIGHT, NAVY, NAVY, 13)
box(s, 1.0, 3.60, 2.4, 0.85, ['상위 통신', ('WCS_TASK_HOST', 9, False, GRAY)], WHITE, BLUE, NAVY, 13)
box(s, 5.00, 1.60, 2.4, 0.85, ['스케줄러', ('IO_TASK', 9, False, GRAY)], WHITE, BLUE, NAVY, 13)
box(s, 4.85, 3.35, 2.7, 1.35, ['DB', ('Microsoft SQL Server', 9, False, WHITE), ('LGLS_MCS_IO', 11, True, WHITE)],
    NAVY, NAVY, WHITE, 16, MSO_SHAPE.FLOWCHART_MAGNETIC_DISK)
box(s, 5.00, 5.50, 2.4, 0.85, ['설비 통신', ('WCS_TASK_CV', 9, False, GRAY)], WHITE, BLUE, NAVY, 13)
box(s, 9.20, 3.60, 2.6, 0.85, ['운전 화면', ('Ecs.exe (Client)', 9, False, GRAY)], WHITE, BLUE, NAVY, 13)
box(s, 9.20, 5.50, 2.6, 0.85, ['설비 / PLC', ('XGT · C/V · S/C · RGV', 9, False, GRAY)], LIGHT, NAVY, NAVY, 13)

arrow(s, 2.20, 3.00, 2.20, 3.60)
arrow(s, 2.20, 3.60, 2.20, 3.00)
arrow(s, 3.40, 4.02, 4.85, 4.02)
arrow(s, 6.20, 2.45, 6.20, 3.35)
arrow(s, 6.35, 3.35, 6.35, 2.45)
arrow(s, 6.20, 5.50, 6.20, 4.70)
arrow(s, 6.35, 4.70, 6.35, 5.50)
arrow(s, 9.20, 4.02, 7.55, 4.02)
arrow(s, 7.40, 5.92, 9.20, 5.92, GREEN)
arrow(s, 9.20, 6.12, 7.40, 6.12, GREEN)

label(s, 1.0, 4.50, 2.4, 'O/R/M/D 수신 · S/E/F 송신 · TCP')
label(s, 7.55, 6.15, 1.65, 'XGT 1소켓')

tb(s, 8.30, 1.30, 4.6, 2.2, [
    ('■ 각 프로그램의 역할', True, 12, NAVY),
    ('· 상위 통신 : WMS/IMS 전문을 작업(JOB_MST)으로 바꾸고', False, 10.5, GRAY),
    ('  상태·완료·에러를 되보고한다. 완료 응답을 받으면 삭제.', False, 10.5, GRAY),
    ('· 스케줄러 : 작업 상태를 진행시키고 설비 지시(_OD)와', False, 10.5, GRAY),
    ('  방향 지시(CMD_RQ_ID=DIR)를 쓴다.', False, 10.5, GRAY),
    ('· 설비 통신 : PLC 를 읽어 _RD·알람 코드를 쓰고, 지시를 PLC 에 전송.', False, 10.5, GRAY),
    ('· 운전 화면 : DB 를 보여주고 수동 지시·설정을 DB 에 기록한다.', False, 10.5, GRAY),
    ('  설비와 직접 통신하지 않는다.', False, 10.5, GRAY),
])
tb(s, 0.6, 5.55, 3.9, 1.2, [
    ('■ 한 줄 요약', True, 12, NAVY),
    ('프로그램끼리 직접 부르지 않는다.', False, 10.5, GRAY),
    ('모두 DB 를 통해 주고받으므로,', False, 10.5, GRAY),
    ('하나가 죽어도 지시와 상태는 DB 에 남는다.', False, 10.5, GRAY),
])
footer(s, '시스템 구성도')

# ══════════ 4. 입고 플로우차트 ══════════
s = prs.slides.add_slide(blank)
head(s, '입고 플로우차트', '작업 상태(JOB_STATUS)를 따라가는 판단 흐름. 로직1(124) 기준이며 로직2(130)·로직3(122)은 출발 트랙만 다르다.')

X = 4.9; W = 3.9
ys = [1.25, 1.95, 2.65, 3.35, 4.05, 4.75, 5.45, 6.15]
steps = [
    ('시작 : 입고대에 파렛트 적재', LIGHT, '', MSO_SHAPE.ROUNDED_RECTANGLE),
    ('상위 O 전문 수신 · 작업 생성 (99 → 10)', WHITE, '상위 통신 · 스케줄러', MSO_SHAPE.RECTANGLE),
    ('입고대 CV 이송 지시 (10 → 15)\nRGV 픽업 트랙 도착 확인 (15 유지)', WHITE, '스케줄러', MSO_SHAPE.RECTANGLE),
    ('RGV 반송 (15 → 35 → 39)\n픽업 트랙 → 크레인 통로', WHITE, '스케줄러 · RGV', MSO_SHAPE.RECTANGLE),
    ('도착 트랙 실물 확인 (39 → 15 또는 16)', WHITE, '스케줄러', MSO_SHAPE.RECTANGLE),
    ('크레인 반송 (→ 25)\nS/C측 트랙 → 랙 셀', WHITE, '스케줄러 · 크레인', MSO_SHAPE.RECTANGLE),
    ('크레인 완료 신호 → 입고 완료 (25 → 09)', WHITE, '스케줄러', MSO_SHAPE.RECTANGLE),
    ('상위 F 전문 보고 → 응답 수신 시 삭제', LIGHT, '상위 통신', MSO_SHAPE.ROUNDED_RECTANGLE),
]
for i, (t, f, who, shp) in enumerate(steps):
    h = 0.55 if '\n' not in t else 0.58
    box(s, X, ys[i], W, h, t.split('\n'), f, NAVY, NAVY, 11, shp)
    if who:
        tb(s, X - 2.4, ys[i] + 0.13, 2.25, 0.3, who, 10, True, BLUE, PP_ALIGN.RIGHT)
    if i < len(steps) - 1:
        arrow(s, X + W / 2, ys[i] + h, X + W / 2, ys[i + 1])

tb(s, 9.3, 1.25, 3.6, 0.35, '■ RGV 지시 전 확인 조건', 12, True, NAVY)
for i, t in enumerate(['① 드롭칸 비어 있음 (칸 단위 2-deep)',
                       '② 출발·도착 트랙 H/S 신호 ON',
                       '③ RTV 자동·유휴, 다른 작업 반송 중 아님',
                       '④ 픽업 트랙 화물 = 이 작업의 화물',
                       '⑤ 도착 트랙에 착지 기록 대기(39) 작업 없음',
                       '⑥ 1호기행 : C/V#2 입고 사용권 확보']):
    tb(s, 9.3, 1.65 + i * 0.33, 3.6, 0.3, t, 10, False, GRAY)
box(s, 9.3, 3.7, 3.6, 0.6,
    ['하나라도 아니면 지시하지 않고', '로그에 보류 사유를 남긴다'], LIGHT, AMBER, AMBER, 10)

tb(s, 9.3, 4.5, 3.6, 0.35, '■ 상태 코드', 12, True, NAVY)
for i, t in enumerate(['99 신규   10 CV 구동대기   15 CV 구동중',
                       '16 통로CV 구동중 ([CNF] IN_HS_STATUS=16)',
                       '35 RGV 구동중   39 RGV 구동완료',
                       '25 크레인 구동중   09 완료',
                       '※ 11 · 21 · 30 · 31 은 폐기']):
    tb(s, 9.3, 4.9 + i * 0.32, 3.6, 0.3, t, 10, False, GRAY)
footer(s, '입고 플로우차트')

# ══════════ 5. 입고 플로우 다이어그램 ══════════
LANES = [('WMS / IMS', 1.30), ('상위 통신', 2.15), ('스케줄러', 3.00),
         ('설비 통신', 3.85), ('설비 / PLC', 4.70)]


def swimlanes(slide, note):
    for name, y in LANES:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(y), Inches(12.3), Inches(0.72))
        bar.fill.solid(); bar.fill.fore_color.rgb = LIGHT
        bar.line.color.rgb = LINE; bar.line.width = Pt(0.75); bar.shadow.inherit = False
        bar.text_frame.text = ''
        tb(slide, 0.6, y + 0.22, 1.5, 0.3, name, 11, True, NAVY)
    tb(slide, 0.5, 5.6, 12.3, 1.3, note, 11, False, GRAY)


def step(slide, x, lane, w, text, fill=WHITE, edge=BLUE, size=9.5):
    y = LANES[lane][1] + 0.09
    return box(slide, x, y, w, 0.54, text if isinstance(text, list) else [text],
               fill, edge, NAVY, size, MSO_SHAPE.ROUNDED_RECTANGLE, True)


def hop(slide, x, l1, l2, color=BLUE):
    y1 = LANES[l1][1] + (0.72 if l2 > l1 else 0.09)
    y2 = LANES[l2][1] + (0.09 if l2 > l1 else 0.72)
    arrow(slide, x, y1, x, y2, color)


s = prs.slides.add_slide(blank)
head(s, '입고 플로우 다이어그램', '누가 무엇을 하는지 시간 순서(왼쪽 → 오른쪽)로 본 흐름')
swimlanes(s, ['· 스케줄러는 설비에 직접 말하지 않는다. 지시를 DB(_OD)에 쓰면 설비 통신이 PLC 로 전송한다.',
              '· 설비가 올린 완료 신호는 설비 통신이 관측해 Ack 를 돌려주고 DB(_RD)에 남긴다. 스케줄러는 그 값과 화물 위치를 보고 상태를 올린다.',
              '· 입고는 크레인 완료 즉시 09(완료)가 되고, 상위 응답을 받으면 작업이 삭제된다. 응답이 없으면 RE_REPORT_09_SEC 간격으로 재보고한다.'])
xs = [1.95, 3.55, 5.15, 6.75, 8.35, 9.95, 11.35]
step(s, xs[0], 0, 1.45, ['O 전문', '(작업 지시)'], LIGHT)
hop(s, xs[0] + 0.72, 0, 1)
step(s, xs[0], 1, 1.45, ['작업 생성', 'JOB_MST 99'])
hop(s, xs[0] + 0.72, 1, 2)
step(s, xs[1], 2, 1.45, ['CV 지시', '99 → 10 → 15'])
hop(s, xs[1] + 0.72, 2, 3)
step(s, xs[1], 3, 1.45, ['PLC 쓰기', '트래킹/목적지'])
hop(s, xs[1] + 0.72, 3, 4)
step(s, xs[1], 4, 1.45, ['컨베이어 이송', '→ RGV 픽업'], LIGHT, NAVY)
step(s, xs[2], 2, 1.45, ['RGV 지시', '15 → 35'])
hop(s, xs[2] + 0.72, 2, 3)
step(s, xs[2], 3, 1.45, ['반송지시 전송'])
hop(s, xs[2] + 0.72, 3, 4)
step(s, xs[2], 4, 1.45, ['RGV 반송', '통로로 하역'], LIGHT, NAVY)
step(s, xs[3], 4, 1.45, ['하역 완료 신호'], LIGHT, NAVY)
hop(s, xs[3] + 0.72, 4, 3, GREEN)
step(s, xs[3], 3, 1.45, ['완료 관측 · Ack', 'COMPLETE_RD'], WHITE, GREEN)
hop(s, xs[3] + 0.72, 3, 2, GREEN)
step(s, xs[3], 2, 1.45, ['35 → 39', '→ 15/16'], WHITE, GREEN)
step(s, xs[4], 2, 1.45, ['크레인 지시', '→ 25'])
hop(s, xs[4] + 0.72, 2, 3)
step(s, xs[4], 3, 1.45, ['반송지시 전송'])
hop(s, xs[4] + 0.72, 3, 4)
step(s, xs[4], 4, 1.45, ['크레인 적재', '셀에 저장'], LIGHT, NAVY)
step(s, xs[5], 4, 1.25, ['완료 신호'], LIGHT, NAVY)
hop(s, xs[5] + 0.62, 4, 3, GREEN)
step(s, xs[5], 3, 1.25, ['완료 관측 · Ack'], WHITE, GREEN)
hop(s, xs[5] + 0.62, 3, 2, GREEN)
step(s, xs[5], 2, 1.25, ['25 → 09'], WHITE, GREEN)
hop(s, xs[5] + 0.62, 2, 1, GREEN)
step(s, xs[5], 1, 1.25, ['F 전문 보고'], WHITE, GREEN)
hop(s, xs[5] + 0.62, 1, 0, GREEN)
step(s, xs[6], 0, 1.35, ['응답'], LIGHT)
step(s, xs[6], 1, 1.35, ['작업 삭제'], WHITE, GREEN)
arrow(s, xs[5] + 1.25, LANES[0][1] + 0.36, xs[6], LANES[0][1] + 0.36, GREEN)
hop(s, xs[6] + 0.67, 0, 1, GREEN)
footer(s, '입고 플로우 다이어그램')

# ══════════ 6. 출고 플로우차트 ══════════
s = prs.slides.add_slide(blank)
head(s, '출고 플로우차트', '입고의 역순이다. 크레인이 먼저 움직이고, 출고대 RGV측 트랙에 실물이 기록되면 완료다.')
steps = [
    ('시작 : 상위 O 전문(출발 셀 · 도착 작업대)', LIGHT, '상위 통신', MSO_SHAPE.ROUNDED_RECTANGLE),
    ('작업 생성 (99) → 크레인 구동대기 (20)', WHITE, '스케줄러', MSO_SHAPE.RECTANGLE),
    ('크레인 반송 (20 → 25 → 29)\n랙 셀 → S/C측 트랙 · 상위 1차 완료보고', WHITE, '스케줄러 · 크레인', MSO_SHAPE.RECTANGLE),
    ('H/S 트랙 착지 확인 (29 → 16)', WHITE, '스케줄러', MSO_SHAPE.RECTANGLE),
    ('RGV 반송 (16 → 35 → 39)\n통로 → 출고대 RGV측 트랙', WHITE, '스케줄러 · RGV', MSO_SHAPE.RECTANGLE),
    ('도착 트랙 작업번호 기록 확인 (39 → 09)', WHITE, '스케줄러', MSO_SHAPE.RECTANGLE),
    ('상위 F 전문 보고 → 응답 수신 시 삭제', WHITE, '상위 통신', MSO_SHAPE.RECTANGLE),
    ('출고대에서 지게차 반출 (ECS 책임 밖)', LIGHT, '작업자', MSO_SHAPE.ROUNDED_RECTANGLE),
]
for i, (t, f, who, shp) in enumerate(steps):
    h = 0.55 if '\n' not in t else 0.58
    box(s, X, ys[i], W, h, t.split('\n'), f, NAVY, NAVY, 11, shp)
    if who:
        tb(s, X - 2.4, ys[i] + 0.13, 2.25, 0.3, who, 10, True, BLUE, PP_ALIGN.RIGHT)
    if i < len(steps) - 1:
        arrow(s, X + W / 2, ys[i] + h, X + W / 2, ys[i + 1])

tb(s, 9.3, 1.25, 3.6, 0.35, '■ 출고 작업대', 12, True, NAVY)
for i, t in enumerate(['126 원부자재 불출대 (C/V#13, 로직1)',
                       '129 피킹 작업대 (C/V#14, 로직2)',
                       '122 입출고 겸용대 (C/V#11, 로직3)']):
    tb(s, 9.3, 1.65 + i * 0.32, 3.6, 0.3, t, 10, False, GRAY)
box(s, 9.3, 2.72, 3.6, 1.05,
    ['크레인 출고 지시 전 확인', '크레인 유휴 · 출고 H/S ON · 하차 트랙 빔',
     '같은 크레인 출고가 CV 구간에 2건 미만', '크레인당 사이클 1건 · 1호기는 출고 사용권'],
    LIGHT, AMBER, AMBER, 9.5)
tb(s, 9.3, 3.95, 3.6, 0.35, '■ 상태 코드', 12, True, NAVY)
for i, t in enumerate(['99 신규   20 크레인 대기   25 구동중',
                       '29 크레인 완료(출고 1차 - 랙 셀 해제)',
                       '16 통로CV 구동중   35 · 39 RGV',
                       '09 완료 → 응답 시 삭제',
                       '※ 출고대 신호 경로 : 15 → 19 → 09']):
    tb(s, 9.3, 4.35 + i * 0.32, 3.6, 0.3, t, 10, False, GRAY)
footer(s, '출고 플로우차트')

# ══════════ 7. 출고 플로우 다이어그램 ══════════
s = prs.slides.add_slide(blank)
head(s, '출고 플로우 다이어그램', '누가 무엇을 하는지 시간 순서(왼쪽 → 오른쪽)로 본 흐름')
swimlanes(s, ['· 출고는 크레인이 먼저 움직인다. 크레인이 통로에 내려놓으면 컨베이어와 RGV 가 출고대까지 이어 나른다.',
              '· 크레인 완료(29) 때 상위에 1차 완료보고를 보낸다(랙 셀 해제). 상태는 그대로 두고 작업은 계속 진행한다.',
              '· RTV 완료만으로 끝내지 않는다. 출고대 RGV측 트랙에 작업번호가 실제로 기록된 것을 확인한 뒤 09 로 완료한다.'])
step(s, xs[0], 0, 1.45, ['O 전문', '(출고 지시)'], LIGHT)
hop(s, xs[0] + 0.72, 0, 1)
step(s, xs[0], 1, 1.45, ['작업 생성', 'JOB_MST 99'])
hop(s, xs[0] + 0.72, 1, 2)
step(s, xs[1], 2, 1.45, ['크레인 지시', '99 → 20 → 25'])
hop(s, xs[1] + 0.72, 2, 3)
step(s, xs[1], 3, 1.45, ['반송지시 전송'])
hop(s, xs[1] + 0.72, 3, 4)
step(s, xs[1], 4, 1.45, ['크레인 출고', '셀 → 트랙'], LIGHT, NAVY)
step(s, xs[2], 4, 1.45, ['완료 신호'], LIGHT, NAVY)
hop(s, xs[2] + 0.72, 4, 3, GREEN)
step(s, xs[2], 3, 1.45, ['완료 관측 · Ack'], WHITE, GREEN)
hop(s, xs[2] + 0.72, 3, 2, GREEN)
step(s, xs[2], 2, 1.45, ['25 → 29'], WHITE, GREEN)
hop(s, xs[2] + 0.72, 2, 1, GREEN)
step(s, xs[2], 1, 1.45, ['1차 완료보고 F'], WHITE, GREEN)
step(s, xs[3], 2, 1.45, ['착지 29 → 16', 'RGV 지시 → 35'])
hop(s, xs[3] + 0.72, 2, 3)
step(s, xs[3], 3, 1.45, ['반송지시 전송'])
hop(s, xs[3] + 0.72, 3, 4)
step(s, xs[3], 4, 1.45, ['RGV 반송', '→ 출고대'], LIGHT, NAVY)
step(s, xs[4], 4, 1.45, ['하역 완료 신호'], LIGHT, NAVY)
hop(s, xs[4] + 0.72, 4, 3, GREEN)
step(s, xs[4], 3, 1.45, ['완료 관측 · Ack'], WHITE, GREEN)
hop(s, xs[4] + 0.72, 4 - 1, 2, GREEN)
step(s, xs[4], 2, 1.45, ['35 → 39'], WHITE, GREEN)
step(s, xs[5], 4, 1.25, ['트랙 작업번호', '기록'], LIGHT, NAVY)
hop(s, xs[5] + 0.62, 4, 3, GREEN)
step(s, xs[5], 3, 1.25, ['기록 관측'], WHITE, GREEN)
hop(s, xs[5] + 0.62, 3, 2, GREEN)
step(s, xs[5], 2, 1.25, ['39 → 09'], WHITE, GREEN)
hop(s, xs[5] + 0.62, 2, 1, GREEN)
step(s, xs[5], 1, 1.25, ['F 전문 보고'], WHITE, GREEN)
hop(s, xs[5] + 0.62, 1, 0, GREEN)
step(s, xs[6], 0, 1.35, ['응답'], LIGHT)
step(s, xs[6], 1, 1.35, ['작업 삭제'], WHITE, GREEN)
arrow(s, xs[5] + 1.25, LANES[0][1] + 0.36, xs[6], LANES[0][1] + 0.36, GREEN)
hop(s, xs[6] + 0.67, 0, 1, GREEN)
footer(s, '출고 플로우 다이어그램')

# ══════════ 8. S/C#1 통로(C/V#2) 사용권 ══════════
s = prs.slides.add_slide(blank)
head(s, 'S/C#1 통로(C/V#2) 방향 전환 - 통로 사용권',
     '입고 드롭(103→104)과 출고 반출(104→103)이 한 통로를 쓴다. 현장 설정 [CNF] CV2_DIR_MODE=2 : 지시를 낼 때만 방향을 바꾼다.')
tb(s, 0.5, 1.2, 3.9, 0.3, '출고 - DriveSC (크레인 1호기 출고 후보 20)', 11, True, NAVY)
vchain(s, 0.5, 1.55, 2.6, [
    ('입고 사용권 또는 입고 예약?', WHITE, AMBER, '있음 → 보류'),
    (['CV2_FAIR=1 이고', '입고가 더 오래 기다림?'], WHITE, AMBER, '예 → 보류'),
    ('출고 예약 (20초)', WHITE, BLUE, None),
    (['통로 방향 ≠ 출고?', '→ DIR(1) 기록, 다음 사이클'], WHITE, BLUE, None),
    ('출고 H/S ON → 크레인 지시', WHITE, GREEN, None),
    (['25 = 출고 사용권', '(25 · 29 · 16, 35+통로 위 실물)'], LIGHT, GREEN, None),
], h=0.52, gap=0.18, size=9)
tb(s, 0.5, 6.42, 3.9, 0.5, '풀림 : RTV 가 출고 화물을 싣고 통로를 벗어날 때', 9.5, True, GREEN)

tb(s, 4.55, 1.2, 4.1, 0.3, '입고 - DriveRGV (1호기행 입고 후보 15/16)', 11, True, NAVY)
vchain(s, 4.55, 1.55, 2.6, [
    ('출고 사용권 또는 출고 예약?', WHITE, AMBER, '있음 → 보류'),
    (['FAIR=0 : 대기 출고 + 1호기 정상?', 'FAIR=1 : 출고가 더 오래?'], WHITE, AMBER, '예 → 양보'),
    ('입고 예약 (20초)', WHITE, BLUE, None),
    (['통로 방향 ≠ 입고?', '→ DIR(0) 기록, 다음 사이클'], WHITE, BLUE, None),
    ('출발·도착 H/S ON → RTV 지시', WHITE, GREEN, None),
    (['35 = 입고 사용권', '(35 · 39, 16/15+통로 위 실물)'], LIGHT, GREEN, None),
], h=0.52, gap=0.18, size=9)
tb(s, 4.55, 6.42, 4.1, 0.5, '풀림 : 크레인 1호기에 입고 지시(25)를 줄 때', 9.5, True, GREEN)

tb(s, 8.85, 1.2, 4.1, 0.3, '■ 전환 방식 [CNF] CV2_DIR_MODE', 11, True, NAVY)
grid(s, 8.85, 1.52, [0.55, 3.45], [
    ['값', '동작'],
    ['0', '종전 - 입고 진행 → 출고 진행 → 대기 순 수시 전환'],
    ['1', '교착 때만 (코드 기본값) - 현재 방향 작업이 없으면 전환, 크레인 25 가 아니고 CV2_STALL_SEC(60초) 무진행이면 강제 전환(DIRF)'],
    ['2', '통로 사용권 (현장 설정) - 왼쪽 흐름'],
], font=8.5, rowh=0.3)
tb(s, 8.85, 3.35, 4.1, 0.3, '■ 방향 지시 (CV_DATA.CMD_RQ_ID)', 11, True, NAVY)
grid(s, 8.85, 3.67, [0.65, 3.45], [
    ['ID', '의미 / 설비 통신 처리'],
    ['DIR', '일반 - 현재 방향 화물이 남아 있으면 보류'],
    ['DIRF', '교착 해제 - 화물 보류 없이 즉시'],
    ['DIRW', 'HOST 입고 복귀 - 빈 뒤 DIR 로 승격'],
    ['DIRM', '운전 화면 수동 - 즉시, 자동 지시 120초 억제'],
], font=8.5, rowh=0.28)
box(s, 8.85, 5.5, 4.1, 1.4,
    [('함께 동작하는 보호', 10, True, NAVY),
     ('· 크레인당 사이클 1건 지시 (09-16 교착 방지)', 9, False, GRAY),
     ('· 통로에 입고 화물·RTV 1호기행 입고 지시 → 1호기 출고 보류', 9, False, GRAY),
     ('· 1호기 출고 25/29/16 · 통로 출고 화물 → 1호기행 입고 보류', 9, False, GRAY),
     ('· CV2_FAIR 0 = 출고 우선(현장) / 1 = 번갈아', 9, False, GRAY)],
    LIGHT, NAVY, GRAY, 9, MSO_SHAPE.ROUNDED_RECTANGLE, False)
footer(s, 'C/V#2 통로 사용권')

# ══════════ 9. 겸용대 보호 · 모드 전환 ══════════
s = prs.slides.add_slide(blank)
head(s, '겸용 입출고대(122) 입고 전환 보호 · 모드 전환',
     '출고 화물이 막 도착한 순간 입고로 뒤집히면 벨트가 화물을 121 로 되끌어 갇힌다(09-16 실측). 입고(0) 전환은 모든 경로가 아래 검사를 지난다.')
tb(s, 0.5, 1.25, 5.8, 0.3, 'IsDualStnClearForInbound(122)', 12, True, NAVY)
vchain(s, 0.5, 1.62, 3.9, [
    (['입고 전환 요청', '(자동 복귀 · HOST DIRW 승격 · RequestCvDirection)'], LIGHT, NAVY, None),
    ('121 · 122 화물 센서 0, 트래킹 0 ?', WHITE, AMBER, '아니오 → 보류'),
    ('빈 상태가 3초 이상 연속 관측?', WHITE, AMBER, '아니오 → 대기'),
    ('122 행 출고가 15초 안에 완료(09/19/39)?', WHITE, AMBER, '예 → 보류'),
    ('작업대로 오는 출고 진행 중?', WHITE, AMBER, '예 → 보류'),
    ('DIR(0) 기록 → 설비 통신이 방향 워드 기록', LIGHT, GREEN, None),
], h=0.55, gap=0.2, size=9.5)
tb(s, 6.9, 1.25, 6.0, 0.3, '■ 상위 모드 전환(M 전문)', 12, True, NAVY)
for i, t in enumerate(['· WMS 가 101(122) 또는 110(103)의 모드를 보낸다.',
                       '· 출고(1) : 즉시 CMD_RQ_ID=DIR 로 기록.',
                       '· 입고(0) : DIRW(대기)로 기록 → IO_TASK PromotePendingDirection 이',
                       '  작업대와 짝 트랙이 비고, 122 는 위 검사를 통과하면 DIR 로 승격.']):
    tb(s, 6.9, 1.62 + i * 0.3, 6.0, 0.3, t, 10, False, GRAY)
tb(s, 6.9, 3.0, 6.0, 0.3, '■ 겸용대 CV 지시 보류 (DriveCV)', 12, True, NAVY)
for i, t in enumerate(['· 122 로 오는 출고가 RTV 에 지시돼 있으면 122 발 입고 CV 지시 보류',
                       '· 1호기 출고가 끝나지 않았으면 122 발 1호기행 입고 CV 지시 보류',
                       '· 122 행 출고 CV 지시 전에 방향이 입고면 출고(1)로 먼저 전환']):
    tb(s, 6.9, 3.37 + i * 0.3, 6.0, 0.3, t, 10, False, GRAY)
tb(s, 6.9, 4.45, 6.0, 0.3, '■ 수동 전환 우선', 12, True, NAVY)
for i, t in enumerate(['· 운전 화면에서 방향 칸을 더블클릭하면 DIRM 으로 즉시 기록된다.',
                       '· 그 뒤 [CNF] MANUAL_DIR_HOLD_SEC(기본 120초) 동안',
                       '  자동 방향 지시(RequestCvDirection)를 내지 않는다.']):
    tb(s, 6.9, 4.82 + i * 0.3, 6.0, 0.3, t, 10, False, GRAY)
footer(s, '겸용대 보호 · 모드 전환')

# ══════════ 10. 피킹 작업대 ══════════
s = prs.slides.add_slide(blank)
head(s, '피킹 작업대 흐름 (102)', '129(출고)와 130(재입고)은 물리적으로 한 자리다. IMS 는 둘을 하나의 작업대 102 로 본다.')
box(s, 1.0, 1.5, 2.5, 0.9, ['출고 흐름', '랙 → 129 도착'], WHITE, NAVY, NAVY, 12)
box(s, 4.2, 1.5, 2.5, 0.9, ['작업자 피킹', '필요 수량만 꺼냄'], LIGHT, AMBER, AMBER, 12)
box(s, 7.4, 1.5, 2.5, 0.9, ['PLC 인계', '#30 ON → #29 OFF'], WHITE, GREEN, GREEN, 12)
box(s, 10.6, 1.5, 2.2, 0.9, ['재입고 흐름', '130 → 131 → 랙'], WHITE, NAVY, NAVY, 12)
arrow(s, 3.5, 1.95, 4.2, 1.95); arrow(s, 6.7, 1.95, 7.4, 1.95); arrow(s, 9.9, 1.95, 10.6, 1.95)

tb(s, 1.0, 2.85, 11.8, 0.4, '■ PLC 사양(시나리오 슬라이드 13)에 적힌 인계 순서', 13, True, NAVY)
seq = ['① 화물이 #29 에 도착 → W.O 비트 ON, 트래킹 유지',
       '② PLC Unload Request #2 ON → ECS Ack ON   (작업자 피킹 구간)',
       '③ 반출 완료 → PLC Unload Complete #2 ON → ECS Ack ON → 양쪽 OFF',
       '④ Pallet Exist #30 ON   (아직 #29 도 ON)',
       '⑤ Pallet Exist #29 OFF   ※ 사양 주석 : "출고 영역 #29 bit 만 지움"']
for i, t in enumerate(seq):
    tb(s, 1.2, 3.3 + i * 0.38, 11.5, 0.36, t, 11, False, GRAY)

box(s, 1.0, 5.35, 11.8, 1.35,
    [('운영 기준', 12, True, AMBER),
     ('· 재입고는 상위의 새 O 전문(출발 102 → 내부 130)으로 시작한다. RGV 픽업 트랙은 131 이다(C/V#15 슬롯 순서 131,132,130).', 11, False, GRAY),
     ('· 상위 보고 시 129 · 130 은 모두 102 로 되돌린다.', 11, False, GRAY),
     ('· 피킹대 라인은 3칸이라 입고 드롭 판정이 칸 단위로 3-deep 이 된다.', 11, False, GRAY)],
    LIGHT, AMBER, GRAY, 11, MSO_SHAPE.ROUNDED_RECTANGLE, False)
footer(s, '피킹 작업대 흐름')

# ══════════ 11. 알람 · 에러 흐름 ══════════
s = prs.slides.add_slide(blank)
head(s, '알람 · 에러 흐름 (PLC 알람 리스트 2026-09-17)',
     'PLC 알람을 설비 통신이 코드로 바꿔 DB 에 남기고, 상위 통신이 E 전문으로 보고한다. 스케줄러는 에러 설비에 지시하지 않는다.')
cx = [0.5, 3.05, 5.6, 8.15, 10.7]
titles = [['설비 / PLC', ('알람코드 워드', 9, False, GRAY), ('알람 비트 M55xx·M56xx', 9, False, GRAY)],
          ['설비 통신', ('WCS_TASK_CV', 9, False, GRAY), ('VehThread · CvThread', 9, False, GRAY)],
          ['DB', ('ERR_CODE_RD / ERROR_CODE', 9, False, WHITE), ('EQP_ERR_HIS 이력', 9, False, WHITE)],
          ['상위 통신', ('GetErrorReport', 9, False, GRAY), ('HOST_ERR_SEND_YN=N 대상', 9, False, GRAY)],
          ['WMS / IMS', ('E 전문 수신', 9, False, GRAY), ('ErrorKind · 코드', 9, False, GRAY)]]
for i, x in enumerate(cx):
    fill = NAVY if i == 2 else (LIGHT if i in (0, 4) else WHITE)
    box(s, x, 1.35, 2.15, 1.05, titles[i], fill, NAVY, WHITE if i == 2 else NAVY, 12)
    if i < len(cx) - 1:
        arrow(s, x + 2.15, 1.87, cx[i + 1], 1.87, RED, 1.75)
grid(s, 0.5, 2.7, [1.6, 3.0, 3.4, 4.35], [
    ['설비', 'PLC 신호', 'DB 기록', '비고'],
    ['스태커 크레인', '알람코드 워드 (SC_ERR_CODE_BLOCK=0 상시)', 'SC_DATA_LGLS.ERR_CODE_RD', '코드표 EQP_ECD_MST 구분 SC_LGLS'],
    ['메인 C/V', '과부하 비트 M5501 ~ M550D', "CV_DATA.ERROR_CODE '0119' ~ '0131'", 'CvThread.CvMainAlarmBits - 비트 OFF 면 그 코드일 때만 0'],
    ['RGV', '알람 비트 M5601 ~ M560F', 'RTV_DATA_LGLS.ERR_CODE_RD', 'VehThread.ReadRgvAlarmBitCode - 코드 워드가 0 일 때 사용'],
    ['공통', '-', 'EQP_ERR_HIS', '운전 화면 [설비에러이력]. [CNF] MAIN_ALM_BITS=0 이면 비트 읽기 끔'],
], font=9, rowh=0.3)
tb(s, 0.5, 5.05, 6.0, 0.3, '■ 크레인 에러 ErrorKind (EcsComA.ini [Host])', 12, True, NAVY)
grid(s, 0.5, 5.4, [2.2, 1.3, 2.6], [
    ['코드', 'ErrorKind', '의미'],
    ['73 · 74 (ScDualCodes)', '1', '좌 · 우측 렉 이중입고'],
    ['75 (ScEmptyCodes)', '3', '공출고'],
    ['그 밖', '0', '기계적 에러'],
], font=9.5, rowh=0.3)
box(s, 6.9, 5.05, 5.95, 1.85,
    [('스케줄러의 처리', 11.5, True, NAVY),
     ('· MonitorAlarm : 새 알람 코드를 로그로만 남긴다.', 10, False, GRAY),
     ('· DriveSC / DriveRGV / DriveCV 는 에러 코드가 있는 설비에 지시하지 않는다.', 10, False, GRAY),
     ('· 이중입고·공출고 작업상태 처리(MarkErrorJobStatus)와', 10, False, GRAY),
     ('  재지정 재개(ResumeRedirectedJobs)는 호출하지 않는다.', 10, False, GRAY),
     ('  "이 현장 없음"(사용자 확인 2026-09-16) - 정의만 보존.', 10, False, RED)],
    LIGHT, NAVY, GRAY, 10, MSO_SHAPE.ROUNDED_RECTANGLE, False)
footer(s, '알람 · 에러 흐름')

# ══════════ 12. 예외 처리 ══════════
s = prs.slides.add_slide(blank)
head(s, '예외 처리 흐름', '완료는 설비 신호로만 판정한다. 자동으로 밀어내지 않고, 멈춘 작업은 경고를 남겨 사람이 처리한다.')

tb(s, 0.7, 1.20, 5.9, 0.4, '■ 이중입고(73/74) · 공출고(75)', 13, True, NAVY)
tb(s, 0.9, 1.62, 5.7, 0.32, '→ 이 현장은 작업상태 자동 처리를 쓰지 않는다', 10.5, True, AMBER)
seq1 = ['① 크레인 알람코드 → SC_DATA_LGLS.ERR_CODE_RD, EQP_ERR_HIS',
        '② 상위에 E 전문 보고 (ErrorKind 1 / 3)',
        '③ 에러가 난 크레인에는 새 지시를 내지 않는다',
        '④ 08/07(에러) · 06/05(재지정) 상태 전환은 하지 않는다',
        '⑤ 운전자가 설비와 작업을 확인해 처리한다']
for i2, t in enumerate(seq1):
    tb(s, 0.9, 2.05 + i2 * 0.40, 5.7, 0.38, t, 11, False, GRAY)

tb(s, 6.95, 1.20, 5.9, 0.4, '■ 작업 체류 (설비 무응답)', 13, True, NAVY)
tb(s, 7.15, 1.62, 5.7, 0.32, '→ 체류 복구(시간 기반 자동 완료)는 2026-09-17 폐기', 10.5, True, AMBER)
seq2 = ['① CheckStalledJobs : 1분마다, 300초 이상 멈춘 작업 경고',
        '   ([CNF] JOB_STALL_WARN_SEC, 0 = 감시 끔)',
        '② 상태를 자동으로 바꾸지 않는다 - 실물 위치를 잃지 않기 위함',
        '③ 사람이 화물을 제자리에 옮긴 뒤 [강제완료]',
        '   (S/C 25 → 29, RTV 35 → 39, 설비가 화물을 들고 있으면 거부)']
for i2, t in enumerate(seq2):
    tb(s, 7.15, 2.05 + i2 * 0.40, 5.7, 0.38, t, 11, False, GRAY)

box(s, 0.7, 4.28, 12.15, 0.95,
    [('설비 에러 해제는 운전 화면의 기능이 아니다', 11.5, True, NAVY),
     ('S/C · RGV 에러는 설비(조작반)에서 해제한다. 설비가 알람 해제를 알리면 설비 통신이 Ack 하고 에러 코드를 내린다.', 10.5, False, GRAY),
     ('[환경설정] 의 "시간 기반 자동 처리" 는 기본 사용 안 함이다. 켰을 때만 경과시간으로 완료를 추정한다.', 10.5, False, GRAY)],
    LIGHT, NAVY, GRAY, 10.5, MSO_SHAPE.ROUNDED_RECTANGLE, False)

tb(s, 0.7, 5.32, 12.1, 0.4, '■ 그 밖의 정체와 조치', 13, True, NAVY)
grid(s, 0.7, 5.62, [4.3, 7.8], [
    ['상황', '조치'],
    ['작업 삭제 뒤 RTV/크레인에 지시 흔적이 남음', 'SweepOrphanVehicleData 가 자동 정리 (지시 소비 · 차상 빔 · 유휴 · 작업 없음)'],
    ['상위 완료 응답이 오지 않음', '09 로 유지, EcsComA.ini [Host] RE_REPORT_09_SEC(기본 60초) 간격 재보고'],
    ['상태바 EQUIP / HOST / SCH 빨강', '해당 서버 프로그램 재기동 (지시는 DB 에 남아 이어서 진행된다)'],
    ['C/V#2 통로 방향이 한쪽에 묶임', 'CV2_DIR_MODE 확인. 필요 시 운전 화면에서 방향 칸 더블클릭(수동 전환)'],
], font=9, rowh=0.24)
footer(s, '예외 처리 흐름')

out = os.path.join(OUT, '07_ECS_구성도_및_흐름도.pptx')
prs.save(out)
print('saved', out, 'slides', len(prs.slides))
