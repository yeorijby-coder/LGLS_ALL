# -*- coding: utf-8 -*-
"""ECS 함수단 흐름도 (pptx) - 입고/출고 풀사이클을 실제 함수 호출 단위로 그린다. 2026-09-21 최종판 (Ver 2.0)

근거 소스 (2026-09-17 기준, 함수명은 전부 grep 으로 존재 확인)
  TASK/WCS_TASK_HOST/CSrvWork.cs      상위 수신 : ReadRequest ~ ParseOorR / ParseCancel / ParseModeChange
  TASK/WCS_TASK_HOST/CCliWork.cs      상위 송신 : GetSendData ~ GetJobCompleteReport / ReReportComplete09
  TASK/IO_TASK/CLS/Thread/cThread_SCH.cs         스케줄러 : Thread_Doing 1사이클
  TASK/WCS_TASK_CV_BIN/2_Thread/EQP_THREAD/CvThread.cs   C/V PLC
  TASK/WCS_TASK_CV_BIN/2_Thread/EQP_THREAD/VehThread.cs  S/C · RGV PLC

상태 전이 (cThread_SCH.cs, DB common_code JOB_STATUS 사용 코드)
  입고 : 99 -> 10 -> 15 -> 35 -> 39 -> 15(16) -> 25 -> 09
  출고 : 99 -> 20 -> 25 -> 29 -> 16 -> 35 -> 39 -> 09      (출고대 신호 경로 : 15 -> 19 -> 09)
09-05 판 대비 없어진 함수 : ParseArrived, ParsePallet, GetDirOrder, GetLuggOrder, GetWeightReport, GetEmptyPltRequest
호출이 주석 처리된 함수 : DeleteSemiFinished, MarkErrorJobStatus, ResumeRedirectedJobs (정의만 남음)
시뮬레이터 관련 내용은 넣지 않는다.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

OUT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATE = '2026-09-21'
VER = '2.0'
NAVY = RGBColor(0x1E, 0x27, 0x61)
BLUE = RGBColor(0x1B, 0x5F, 0xB0)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
LIGHT = RGBColor(0xEE, 0xF1, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x59, 0x59, 0x59)
LINE = RGBColor(0xB6, 0xBF, 0xD0)
GREEN = RGBColor(0x23, 0x79, 0x4A)
AMBER = RGBColor(0xB5, 0x72, 0x0A)
RED = RGBColor(0xB0, 0x2E, 0x2E)
TEAL = RGBColor(0x6B, 0x3F, 0xA0)   # [LGLS 2026-09-21] 청록 → 보라 (파랑과 구분)
FONT = '맑은 고딕'

# 프로세스별 색 : 범례와 각 단계 상자가 같은 색을 쓴다
PROC = {
    'HOST': (BLUE,  RGBColor(0xCF, 0xE0, 0xFA), 'WCS_TASK_HOST'),
    'SCH':  (GREEN, RGBColor(0xD4, 0xF0, 0xD8), 'IO_TASK'),
    'CV':   (AMBER, RGBColor(0xFD, 0xE3, 0xBF), 'WCS_TASK_CV / CvThread'),
    'VEH':  (TEAL,  RGBColor(0xE8, 0xDA, 0xF6), 'WCS_TASK_CV / VehThread'),
    'EQP':  (GRAY,  LIGHT,                      '설비 (PLC)'),
}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
page = [0]


def tb(slide, x, y, w, h, text, size=12, bold=False, color=GRAY,
       align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = s.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = 0
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if isinstance(ln, tuple):
            txt, b, sz, col = ln
        else:
            txt, b, sz, col = ln, bold, size, color
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(sz)
        r.font.bold = b
        r.font.color.rgb = col
        r.font.name = FONT
    return s


def box(slide, x, y, w, h, text, fill=WHITE, edge=NAVY, txt=NAVY, size=11,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, bold=True, align=PP_ALIGN.CENTER):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = edge
    sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.07)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(text if isinstance(text, list) else [text]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if isinstance(ln, tuple):
            t, sz, b, c = ln
        else:
            t, sz, b, c = ln, size, bold, txt
        r = p.add_run()
        r.text = t
        r.font.size = Pt(sz)
        r.font.bold = b
        r.font.color.rgb = c
        r.font.name = FONT
    return sh


def arrow(slide, x1, y1, x2, y2, color=BLUE, width=1.5, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    el = c.line._get_or_add_ln()
    el.append(parse_xml('<a:tailEnd %s type="triangle" w="med" len="med"/>' % nsdecls('a')))
    if dash:
        el.insert(0, parse_xml('<a:prstDash %s val="dash"/>' % nsdecls('a')))
    return c


def footer(slide, title):
    page[0] += 1
    tb(slide, 0.45, 7.05, 7, 0.3, 'LG 화학 1동 ECS  |  ' + title, 9, False, GRAY)
    tb(slide, 12.3, 7.05, 0.6, 0.3, str(page[0]), 9, False, GRAY, PP_ALIGN.RIGHT)


def head(slide, text, sub=None):
    tb(slide, 0.45, 0.28, 12.4, 0.5, text, 24, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
    if sub:
        tb(slide, 0.45, 0.80, 12.4, 0.3, sub, 10.5, False, GRAY)


def legend(slide, y=1.05):
    """프로세스 색 범례"""
    x = 0.45
    for key in ('HOST', 'SCH', 'CV', 'VEH'):
        edge, fill, name = PROC[key]
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                    Inches(x), Inches(y), Inches(0.20), Inches(0.16))
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        sh.line.color.rgb = edge
        sh.line.width = Pt(1.0)
        sh.shadow.inherit = False
        tb(slide, x + 0.26, y - 0.03, 2.6, 0.24, name, 9, False, GRAY)
        x += 2.95


def step(slide, x, y, w, h, kind, no, fn, desc, badge=None):
    """함수 단계 상자 : 번호 + 함수명 + 설명 (+ 상태 전이 배지)"""
    edge, fill, _ = PROC[kind]
    lines = [('%s  %s' % (no, fn), 9.5, True, NAVY),
             (desc, 8, False, GRAY)]
    if badge:
        lines.append((badge, 8, True, edge))
    return box(slide, x, y, w, h, lines, fill, edge, NAVY, 9.5)


def chain(slide, steps, cols=3, x0=0.45, y0=1.42, w=3.94, h=0.74, gap=0.80, colgap=4.24):
    """단계 목록을 열 단위로 세로 배치하고 열 안에서 아래로 화살표를 잇는다."""
    rows = (len(steps) + cols - 1) // cols
    for i, st in enumerate(steps):
        c, r = divmod(i, rows)
        x = x0 + c * colgap
        y = y0 + r * gap
        step(slide, x, y, w, h, st[0], st[1], st[2], st[3], st[4] if len(st) > 4 else None)
        if r < rows - 1 and i + 1 < len(steps):
            arrow(slide, x + w / 2, y + h, x + w / 2, y + gap, LINE, 1.25)
    return rows


def grouplist(slide, x0, y, groups, pitch=0.30, gw=1.25, fw=2.35, dw=2.45):
    """묶음 상자 + (함수, 설명) 목록. 다음 y 를 돌려준다."""
    for title, col, items in groups:
        h = 0.14 + pitch * len(items)
        box(slide, x0, y, gw, h, title, LIGHT, col, NAVY, 10)
        yy = y + 0.08
        for fn, desc in items:
            tb(slide, x0 + gw + 0.1, yy, fw, 0.26, fn, 9, True, NAVY)
            tb(slide, x0 + gw + 0.1 + fw + 0.05, yy + 0.01, dw, 0.26, desc, 8.5, False, GRAY)
            yy += pitch
        y += h + 0.08
    return y


# ══════════════════════════════════════════════════════════════════
# 1. 표지
# ══════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
page[0] += 1
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()
bg.shadow.inherit = False
tb(s, 1.0, 2.45, 11, 0.6, 'LG 화학 1동 자동창고', 20, False, ICE)
tb(s, 1.0, 3.10, 11, 1.1, 'ECS 함수단 흐름도', 40, True, WHITE)
tb(s, 1.0, 4.30, 11, 0.5, '입고 · 출고 풀사이클을 실제 함수 호출 단위로 정리', 15, False, ICE)
tb(s, 1.0, 4.95, 11, 0.5, 'ECS Renewal   ·   작성일 %s   ·   Ver %s' % (DATE, VER), 13, False, ICE)
tb(s, 1.0, 5.45, 11, 0.5, '최종판 - 현장 기본 시험 완료 기준 (소스 2026-09-17)', 12, False, ICE)

# ══════════════════════════════════════════════════════════════════
# 2. 문서 이력
# ══════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
head(s, '문서 이력')
rows = [['버전', '일자', '작성', '내용'],
        ['1.0', '2026-09-05', 'LGLS ECS Renewal', '최초 작성 (입고/출고 풀사이클, 스케줄러·C/V 통신·상위 통신 1사이클, 예외 지점)'],
        [VER, DATE, 'LGLS ECS Renewal',
         '최종판 - 명칭 ECS 통일, 상태 전이 현행화(입고 25→09, 출고 29→16→35→39→09), 없어진 함수 정리'
         '(ParseArrived·ParsePallet·GetDirOrder·GetLuggOrder 등), C/V#2 통로 사용권·방향 지시 관문·알람 함수 흐름 추가, '
         '주석 처리 함수(DeleteSemiFinished·MarkErrorJobStatus·ResumeRedirectedJobs) 표기']]
shp = s.shapes.add_table(len(rows), 4, Inches(0.45), Inches(1.3), Inches(12.4), Inches(0.5 * len(rows)))
t = shp.table
for j, w in enumerate([1.2, 1.7, 2.4, 7.1]):
    t.columns[j].width = Inches(w)
for i2, row in enumerate(rows):
    for j2, v in enumerate(row):
        c = t.cell(i2, j2); c.text = ''
        r = c.text_frame.paragraphs[0].add_run(); r.text = v
        r.font.size = Pt(10); r.font.name = FONT
        c.fill.solid()
        if i2 == 0:
            c.fill.fore_color.rgb = NAVY; r.font.color.rgb = WHITE; r.font.bold = True
        else:
            c.fill.fore_color.rgb = WHITE if i2 % 2 else LIGHT
            r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
footer(s, '문서 이력')

# ══════════════════════════════════════════════════════════════════
# 3. 읽는 법 · 프로세스별 진입 함수
# ══════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
head(s, '읽는 법 · 프로세스별 진입 함수',
     '네 개의 처리 루프가 DB(LGLS_MCS_IO)를 매개로 각자 돈다. 서로 직접 호출하는 관계는 없다.')

box(s, 0.45, 1.35, 2.85, 1.30,
    [('WCS_TASK_HOST', 12, True, NAVY), ('상위 통신', 8.5, False, GRAY),
     ('CSrvWork.ReadRequest()', 9, False, BLUE),
     ('CCliWork.GetSendData()', 9, False, BLUE)],
    PROC['HOST'][1], BLUE, NAVY)
box(s, 3.65, 1.35, 2.85, 1.30,
    [('IO_TASK', 12, True, NAVY), ('스케줄러', 8.5, False, GRAY),
     ('cThread_SCH.Thread_Doing()', 9, False, GREEN),
     ('200ms 주기 · 단일 루프', 8.5, False, GRAY)],
    PROC['SCH'][1], GREEN, NAVY)
box(s, 6.85, 1.35, 2.85, 1.30,
    [('WCS_TASK_CV', 12, True, NAVY), ('C/V 통신', 8.5, False, GRAY),
     ('CvThread.Thread_Doing()', 9, False, AMBER),
     ('15설비 슬롯 순회 · 약 16초', 8.5, False, GRAY)],
    PROC['CV'][1], AMBER, NAVY)
box(s, 10.05, 1.35, 2.85, 1.30,
    [('WCS_TASK_CV', 12, True, NAVY), ('S/C · RGV 통신', 8.5, False, GRAY),
     ('VehThread.Thread_Doing()', 9, False, TEAL),
     ('관측 폴링 + 명령 소비', 8.5, False, GRAY)],
    PROC['VEH'][1], TEAL, NAVY)

for x in (3.30, 6.50, 9.70):
    arrow(s, x, 2.00, x + 0.35, 2.00, LINE, 1.25)
    arrow(s, x + 0.35, 2.20, x, 2.20, LINE, 1.25)

box(s, 3.65, 3.05, 6.05, 0.62,
    [('DB : LGLS_MCS_IO', 12, True, WHITE),
     ('JOB_MST · CV_DATA · SC_DATA_LGLS · RTV_DATA_LGLS · EQP_ERR_HIS', 9, False, WHITE)],
    NAVY, NAVY, WHITE, 12, MSO_SHAPE.FLOWCHART_MAGNETIC_DISK)
for x in (1.87, 5.07, 8.27, 11.47):
    arrow(s, x, 2.65, x, 3.05, LINE, 1.25)

tb(s, 0.45, 3.95, 6.05, 2.9, [
    ('■ 표기 규칙', True, 12, NAVY),
    ('', False, 6, GRAY),
    ('· 상자 안 윗줄 = 실제 함수명. 소스 그대로다.', False, 10.5, GRAY),
    ('· 아랫줄 = 그 함수가 하는 일과 건드리는 테이블/신호.', False, 10.5, GRAY),
    ('· 마지막 줄 = JOB_MST.JOB_STATUS 전이. 없는 단계는', False, 10.5, GRAY),
    ('  상태를 바꾸지 않고 설비 신호만 주고받는다.', False, 10.5, GRAY),
    ('· 상자 색 = 그 함수가 사는 프로세스.', False, 10.5, GRAY),
    ('· 번호는 시간 순서다. 열이 넘어가면 다음 열 맨 위로 이어진다.', False, 10.5, GRAY),
], anchor=MSO_ANCHOR.TOP)

tb(s, 6.85, 3.95, 6.05, 3.0, [
    ('■ 작업 상태 전이 (cThread_SCH.cs)', True, 12, NAVY),
    ('', False, 6, GRAY),
    ('입고   99 → 10 → 15 → 35 → 39 → 15(16) → 25 → 09', True, 11, GREEN),
    ('출고   99 → 20 → 25 → 29 → 16 → 35 → 39 → 09', True, 11, BLUE),
    ('', False, 6, GRAY),
    ('· 16 = 통로CV 구동중. 입고는 [CNF] IN_HS_STATUS=16 일 때만.', False, 10, GRAY),
    ('· 11 · 21 · 30 · 31 은 폐기(코드표 사용 N).', False, 10, GRAY),
    ('· 출고 29 는 1차 완료보고만 나가고 작업은 계속 간다.', False, 10, GRAY),
    ('· 09 = 완료. 상위 F 보고 → 응답 수신 시 삭제, 미응답이면 재보고.', False, 10, GRAY),
    ('· 반자동(11/12)은 완료 시점에 DeleteJobNow() 로 즉시 삭제.', False, 10, GRAY),
    ('· 출고대 신호 경로 : 15 → 19(ReportOutStationArrival) → 09.', False, 10, GRAY),
], anchor=MSO_ANCHOR.TOP)
footer(s, '읽는 법')

# ══════════════════════════════════════════════════════════════════
# 4. 입고 풀사이클
# ══════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
head(s, '입고 풀사이클 — 함수 호출 순서',
     '상위 O 전문 수신부터 F 완료보고·삭제까지. JOB_TYP=\'1\', 입고대 124(C/V#12) → S/C 랙 셀 기준.')
legend(s)
chain(s, [
    ('HOST', '01', 'CSrvWork.ReadRequest()',
     'TCP 수신 → CheckHeader() / CheckBody() 로 전문 길이 확인'),
    ('HOST', '02', 'CSrvWork.Parsing() → ParseOorR()',
     'O 전문 고정자리수 파싱 (작업번호·출발/도착 스테이션·셀)'),
    ('HOST', '03', 'IsValidStartStation() / IsValidDestStation()',
     '작업대 번호와 랙 셀 유효성 검증. 실패 시 MakeResponse() 로 반려'),
    ('HOST', '04', 'frmMain.InsertJobMst()',
     'JOB_MST INSERT', 'JOB_STATUS = 99'),
    ('HOST', '05', 'MakeResponse() → SendSock()',
     '상위에 응답 전문(11byte) 회신'),
    ('SCH', '06', 'AcceptNewJob()',
     '입고는 C/V 가 먼저 받는다', '99 → 10'),
    ('SCH', '07', 'DriveCV() → UpdateCvData()',
     'CV_DATA 에 이송 명령 기록 (OD_RQ_YN=\'Y\', DEST_POS_OD)', '10 → 15'),
    ('CV', '08', 'ScanPendingWork() → CvChg_OD_RQ_YN()',
     '쓰기 대기 설비만 골라 PLC 로 이송 명령 WRITE'),
    ('CV', '09', 'CvStatusScenario() → CvTrackingWrite()',
     'PLC 상태 READ → CV_DATA 미러, R 영역 작업번호 트래킹 WRITE'),
    ('SCH', '10', 'CompleteCV()',
     'RGV 픽업 트랙(123)에 이 작업 화물 도착 확인 → 지시값 정리', '15 유지'),
    ('SCH', '11', 'DriveRGV() → Cv2TryTakeForIn()',
     '드롭칸·H/S·RTV 유휴·픽업 화물 소유 확인. 1호기행은 입고 사용권'),
    ('SCH', '12', 'UpdateRtvData() → UpdateJobStatusHs()',
     'RTV 이송 명령 + 도착 트랙을 HS_TRACK_NO 에 기록', '15 → 35'),
    ('VEH', '13', 'VehThread.ConsumeCommands()',
     'RTV 에 출발/도착 트랙 WRITE'),
    ('VEH', '14', 'VehThread.PollObservations()',
     '하역 완료 이벤트 Ack → RTV_DATA_LGLS.COMPLETE_RD 래치'),
    ('SCH', '15', 'CompleteRGVReal()',
     'RtvCompleteFor() 로 완료 신호 소비 → RtvResetComplete()', '35 → 39'),
    ('SCH', '16', 'LandRgvDrop() → RequestArrivalTrackingWrite()',
     '도착 트랙 실물 확인(번호 비었으면 기록 요청). IN_HS_STATUS', '39 → 15 / 16'),
    ('SCH', '17', 'DriveSC() → UpdateScData()',
     '트랙 화물 = 이 작업 · 입고 H/S ON · 크레인당 사이클 1건', '15/16 → 25'),
    ('VEH', '18', 'ConsumeCommands() / PollObservations()',
     'S/C 에 이송 명령 WRITE, 적재 완료 → SC_DATA_LGLS.COMPLETE_RD'),
    ('SCH', '19', 'CompleteSC()',
     '완료 신호 + 유휴 + 8초 경과. 반자동(11)은 DeleteJobNow()', '25 → 09'),
    ('HOST', '20', 'ReReportComplete09() → RequestSrv()',
     '처음 보는 09 는 즉시 F 전송, 이후 RE_REPORT_09_SEC 간격'),
    ('HOST', '21', 'frmMain.DeleteJobMst()',
     '응답 수신 → 작업 삭제(이력 이관). 미응답이면 09 유지', '09 → 삭제'),
])
footer(s, '입고 풀사이클')

# ══════════════════════════════════════════════════════════════════
# 5. 출고 풀사이클
# ══════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
head(s, '출고 풀사이클 — 함수 호출 순서',
     'JOB_TYP=\'2\'. 출고는 S/C 가 먼저 움직인다. 랙 셀 → 출고대 126(C/V#13) 기준.')
legend(s)
chain(s, [
    ('HOST', '01', 'ReadRequest() → Parsing() → ParseOorR()',
     'O 전문 수신·파싱. 출발이 랙 셀이면 출고'),
    ('HOST', '02', 'IsValid*Station() → frmMain.InsertJobMst()',
     '출발 셀과 도착 작업대 검증 후 JOB_MST INSERT', 'JOB_STATUS = 99'),
    ('SCH', '03', 'AcceptNewJob()',
     '출고는 S/C 가 먼저 받는다', '99 → 20'),
    ('SCH', '04', 'DriveSC()',
     '크레인 유휴·출고 H/S·하차 트랙 빔·출고 2-deep·크레인당 1건'),
    ('SCH', '05', 'Cv2TryTakeForOut()  (S/C#1)',
     '입고 사용권·예약이 없을 때만 출고 예약. CV2_FAIR 판정'),
    ('SCH', '06', 'RequestCvDirection(103, "1")',
     'C/V#2 방향이 출고가 아니면 DIR 기록 → 다음 사이클에 지시'),
    ('SCH', '07', 'UpdateScData() → UpdateJobStatusHs()',
     'S/C 이송 명령 + 하차 트랙을 HS_TRACK_NO 에 기록', '20 → 25'),
    ('VEH', '08', 'VehThread.ConsumeCommands()',
     'S/C 에 출발 셀 / 도착 트랙 WRITE'),
    ('VEH', '09', 'VehThread.PollObservations()',
     '하역 완료 → SC_DATA_LGLS.COMPLETE_RD 래치'),
    ('SCH', '10', 'CompleteSC()',
     '출고 1차 완료(랙 셀 해제). 901 은 출고 유지 유예 OUT_HOLD_SEC', '25 → 29'),
    ('HOST', '11', 'GetJobCompleteReport(29)',
     '1차 완료보고 F 전송 → WC_STEP=1 (상태 유지)'),
    ('SCH', '12', 'LandScDrop() → LuggLandedTrack()',
     '하차 트랙/짝 트랙 실물 확인. 없으면 LandedByCargoOnly() + StampHsTrackingForOut()', '29 → 16'),
    ('SCH', '13', 'DriveRGV() → IsRtvBusyWithOwnJob()',
     'RTV 1대 배타, 픽업 화물 소유·H/S 확인'),
    ('SCH', '14', 'UpdateRtvData() → UpdateJobStatusHs()',
     'RTV 이송 명령(통로 → 출고대 RGV측 트랙)', '16 → 35'),
    ('VEH', '15', 'ConsumeCommands() / PollObservations()',
     'RTV 명령 WRITE, 하역 완료 → COMPLETE_RD'),
    ('SCH', '16', 'CompleteRGVReal()',
     '출고도 여기서 끝내지 않는다 - 도착 기록 대기', '35 → 39'),
    ('CV', '17', 'CvStatusScenario() → UpdateCvData()',
     '출고대 RGV측 트랙의 화물·작업번호를 CV_DATA 에 미러'),
    ('SCH', '18', 'LandRgvDrop() → LuggLandedTrack()',
     '도착 트랙에 작업번호 실제 기록 확인. 반자동(12)은 DeleteJobNow()', '39 → 09'),
    ('HOST', '19', 'ReReportComplete09() → RequestSrv()',
     'F 전송 → 응답 시 frmMain.DeleteJobMst()', '09 → 삭제'),
    ('SCH', '20', '(보조) ReportOutStationArrival()',
     '15 인 출고가 출고대에 도착 + RET_READY_RD=1', '15 → 19'),
    ('HOST', '21', '(보조) GetJobCompleteReport(19)',
     'UpdateJobStatusTo(\'09\') → F 전송 → 삭제. 실패 시 19 로 롤백', '19 → 09'),
])
footer(s, '출고 풀사이클')

# ══════════════════════════════════════════════════════════════════
# 6. IO_TASK 스케줄러 1사이클
# ══════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
head(s, 'IO_TASK 스케줄러 — Thread_Doing() 1사이클',
     '200ms 마다 아래 순서를 그대로 한 번 돈다 (왼쪽 열 → 오른쪽 열). Drive* 는 대기 상태 + 유휴 설비에만 작용하는 멱등 함수다.')
yl = grouplist(s, 0.45, 1.25, [
    ('접수', GREEN, [
        ('AcceptNewJob()', '99 → 10(입고·이동) / 20(그 외)'),
    ]),
    ('① 구동 완료', BLUE, [
        ('CompleteCV()', '출고 15 → 19 / 입고는 15 유지'),
        ('CompleteSC()', '입고 25 → 09 / 출고 25 → 29'),
        ('CompleteRGVReal()', '35 → 39 (COMPLETE_RD 소비)'),
        ('CompleteRGVManual()', 'RTV 수동지시(9998) 정리'),
        ('CompleteSCManual()', 'SC 수동지시(9999) 정리'),
    ]),
    ('② 착지 처리', TEAL, [
        ('LandRgvDrop()', '입고 39 → 15/16, 출고 39 → 09'),
        ('LandScDrop()', '출고 29 → 16'),
    ]),
])
box(s, 0.45, yl + 0.1, 6.1, 2.1,
    [('호출하지 않는 함수 (주석 처리, 정의만 보존)', 10.5, True, RED),
     ('DeleteSemiFinished() - 반자동 완료 일괄 삭제. 09-16 비활성', 9, False, GRAY),
     ('   (반자동은 완료 시점에 DeleteJobNow() 로 한 건씩 삭제)', 9, False, GRAY),
     ('MarkErrorJobStatus() - 이중입고/공출고 → 08/07', 9, False, GRAY),
     ('ResumeRedirectedJobs() - 재지정 06/05 → 새 셀 재개', 9, False, GRAY),
     ('   "이 현장은 이중입고/공출고가 없다" (사용자 확인 09-16)', 9, False, GRAY),
     ('※ 2026-09-17 미사용 함수 13개 삭제, 체류 복구(시간 기반 자동 완료) 폐기', 9, False, GRAY)],
    WHITE, RED, GRAY, 9, MSO_SHAPE.ROUNDED_RECTANGLE, False, PP_ALIGN.LEFT)
grouplist(s, 6.85, 1.25, [
    ('③ 구동 지시', AMBER, [
        ('SyncDualCvDirection()', '122 출고 정합·입고 복귀, 103 은 모드 0/1'),
        ('PromotePendingDirection()', 'HOST 입고 복귀 DIRW → DIR 승격'),
        ('DriveCV()', '10 → 15'),
        ('DriveSC()', '입고 15/16 · 출고 20 → 25'),
        ('DriveRGV()', '15/16 → 35'),
    ]),
    ('④ 마무리 · 감시', RED, [
        ('CheckStalledJobs()', '체류 경고만 (1분 주기, 300초)'),
        ('ReportOutStationArrival()', '출고 15 + 출고대 신호 → 19'),
        ('MonitorAlarm() → CheckAlarm()', 'CV/SC/RGV 새 알람 코드 로그'),
        ('ConsumeForceComplete()', '[강제완료] FCMP → 25→29 / 35→39'),
        ('SweepOrphanVehicleData()', '작업 없는 RTV/SC 지시 흔적 정리'),
        ('SyncDisplayTyp()', '표시용 작업구분 보강 (제어 미사용)'),
        ('Heartbeat(true)', '사이클 완주 → 상태표시줄 SCH 녹색'),
    ]),
])
tb(s, 6.85, 5.55, 6.05, 1.3,
   ['※ 사이클이 예외로 깨지면 Heartbeat(false) 후 DB 재연결을 시도한다.',
    '※ Drive* 를 완료·착지 뒤에 한 번만 부르는 것은 방금 인계된 작업도 같은 호출이 잡기 때문이다.',
    '※ ENV_IOSCH.INI [CNF] 값(CV2_DIR_MODE · CV2_FAIR · IN_HS_STATUS · OUT_HOLD_SEC 등)은 호출마다 읽어 재기동 없이 반영된다.'],
   9, False, GRAY)
footer(s, '스케줄러 1사이클')

# ══════════════════════════════════════════════════════════════════
# 7. C/V#2 통로 사용권 함수 흐름
# ══════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
head(s, 'S/C#1 통로(C/V#2) 사용권 — 함수 흐름 ([CNF] CV2_DIR_MODE=2)',
     '왼쪽 = 1호기 출고(DriveSC), 오른쪽 = 1호기행 입고(DriveRGV). 사용권은 작업 상태(DB)로 판정하고, 방향은 지시를 낼 때만 바꾼다.')
legend(s)
chain(s, [
    ('SCH', 'O1', 'DriveSC()  - 1호기 출고 후보(20)',
     '크레인 유휴·하차 트랙 빔 등 일반 게이트 통과'),
    ('SCH', 'O2', 'Cv2TryTakeForOut() → GetSc1AisleToken()',
     '입고 사용권(IN_TOK) 또는 입고 예약(20초)이면 보류'),
    ('SCH', 'O3', 'CV2_FAIR=1 : GetJobInsDt() 비교',
     '사용권이 비고 입고가 더 오래 기다렸으면 보류'),
    ('SCH', 'O4', 'Cv2Reserve("1")',
     '출고 예약 20초 (CV2_RESV_SEC) - 방향 반영 대기 중 선점 방지'),
    ('SCH', 'O5', 'GetCvStockMode(103) ≠ 1 → RequestCvDirection()',
     'Cv2TokenAllowsDir() 통과 시 CMD_RQ_ID=DIR 기록, 이번 사이클 종료'),
    ('CV', 'O6', 'CvChg_CMD_RQ_YN()',
     'IsDualCvDirChangeHeld() 통과 → 방향 D 워드 WRITE → STOCK_MODE 갱신'),
    ('SCH', 'O7', 'IsHsOn(RETHS) → UpdateScData()',
     '다음 사이클에 크레인 지시 = 출고 사용권(25·29·16)', '20 → 25'),
    ('SCH', 'O8', '출고 사용권 해제',
     'RTV 가 출고 화물을 싣고 103/104 를 벗어남(35 + 통로 실물 없음)'),
    ('SCH', 'I1', 'DriveRGV()  - 1호기행 입고 후보(15/16)',
     'CanDropInbound() · IsRtvBusyWithOwnJob() · 픽업 화물 소유 확인'),
    ('SCH', 'I2', 'Cv2TryTakeForIn() → GetSc1AisleToken()',
     '출고 사용권(OUT_TOK) 또는 출고 예약이면 보류'),
    ('SCH', 'I3', 'CV2_FAIR=0 : 대기 출고 + Sc1Ready',
     '1호기 자동·정상이고 대기 출고가 있으면 입고 양보 (FAIR=1 은 대기 시각 비교)'),
    ('SCH', 'I4', 'Cv2Reserve("0")',
     '입고 예약 20초'),
    ('SCH', 'I5', 'GetCvStockMode(103) ≠ 0 → RequestCvDirection()',
     'CMD_RQ_ID=DIR(0) 기록, 이번 사이클 종료'),
    ('CV', 'I6', 'CvChg_CMD_RQ_YN()',
     '통로 출고 화물이 남아 있으면 보류, 없으면 방향 WRITE'),
    ('SCH', 'I7', 'IsHsOn(DEPART/ARRIVE) → UpdateRtvData()',
     'RTV 지시 = 입고 사용권(35·39, 15/16+통로 실물)', '15/16 → 35'),
    ('SCH', 'I8', 'DriveSC() 입고 지시',
     '크레인 1호기에 입고 지시(25)를 주면 입고 사용권 해제', '→ 25'),
], cols=2, w=5.95, h=0.62, gap=0.68, colgap=6.45)
tb(s, 0.45, 6.88, 12.4, 0.2,
   '모드 1(교착 때만) : RequestCvDirection → Cv2SwitchAllowed() - 현재 방향 점유 0 이면 전환, 크레인 25 가 아니고 CV2_STALL_SEC 무진행이면 CMD_RQ_ID=DIRF(즉시).',
   8.5, False, GRAY)
footer(s, 'C/V#2 통로 사용권')

# ══════════════════════════════════════════════════════════════════
# 8. 방향 지시 공통 관문
# ══════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
head(s, '방향 지시 공통 관문 — RequestCvDirection()',
     '겸용대(C/V#2 103/104, C/V#11 121/122)의 자동 방향 지시는 모두 이 함수를 지난다. HOST 입고 복귀(DIRW)만 별도 경로.')
tb(s, 0.45, 1.2, 3.0, 0.3, '호출하는 곳', 12, True, NAVY)
callers = [('SyncDualCvDirection()', '122 출고 정합 / 입고 복귀'),
           ('DriveCV()', '122 행 출고 CV 지시 전 → 출고(1)'),
           ('DriveSC()', '1호기 지시 전 → 103 을 작업 방향으로'),
           ('DriveRGV()', '1호기행 입고 드롭 전 → 103 입고(0)')]
y = 1.55
for fn, desc in callers:
    box(s, 0.45, y, 3.0, 0.62, [(fn, 10, True, NAVY), (desc, 8, False, GRAY)], PROC['SCH'][1], GREEN, NAVY, 10)
    arrow(s, 3.45, y + 0.31, 3.95, 2.95, LINE, 1.0)
    y += 0.78

gates = [
    ('① 중복 억제', '같은 방향 재지시는 DIR_REQ_HOLD_MS(15초) 동안 막는다. 반대 방향은 즉시'),
    ('② IsManualDirHold()', '운전 화면 수동 전환(MANUAL_DIR) 뒤 MANUAL_DIR_HOLD_SEC(120초) 동안 자동 지시 보류'),
    ('③ 103 모드 2 : Cv2TokenAllowsDir()', '상대 사용권·예약이 있으면 보류'),
    ('③ 103 모드 1 : Cv2SwitchAllowed()', '점유 0 이면 허용, 교착(CV2_STALL_SEC)이면 강제(bForceDir)'),
    ('④ 그 외 : IsDualCvBusyWithJob()', '현재 방향의 작업 화물이 겸용대에 있으면 보류'),
    ('⑤ 입고(0) : IsDualStnClearForInbound()', '122 : 121/122 빈 상태 3초 연속 + RecentOutboundArrivedAt() 15초 유예'),
    ('⑥ UPDATE CV_DATA', "CMD_RQ_ID = 'DIR' 또는 'DIRF', CMD_RQ_PARM = 방향, CMD_RQ_YN = 'Y'"),
]
y = 1.25
for i, (fn, desc) in enumerate(gates):
    col = RED if i < 5 else GREEN
    box(s, 3.95, y, 3.55, 0.50, fn, WHITE, col, NAVY, 9.5)
    tb(s, 7.62, y + 0.06, 5.3, 0.45, desc, 9, False, GRAY)
    if i < len(gates) - 1:
        arrow(s, 5.72, y + 0.50, 5.72, y + 0.62, LINE, 1.25)
    y += 0.62
tb(s, 3.95, y + 0.02, 9.0, 0.3, '① ~ ⑤ 중 하나라도 걸리면 false 를 돌려주고 호출부는 다음 사이클에 다시 시도한다.', 9, False, RED)

box(s, 0.45, 6.05, 6.1, 0.85,
    [('PromotePendingDirection()  (HOST M 전문 입고 복귀)', 10, True, NAVY),
     ("CMD_RQ_ID='DIRW' → 작업대·짝 트랙 빔 + (122) IsDualStnClearForInbound() → 'DIR'", 8.5, False, GRAY)],
    PROC['SCH'][1], GREEN, NAVY, 10)
box(s, 6.85, 6.05, 6.05, 0.85,
    [('CvThread.CvChg_CMD_RQ_YN()  (설비 통신)', 10, True, NAVY),
     ('DIR : IsDualCvDirChangeHeld() 보류 가능 / DIRF · DIRM : 즉시 → D 워드 → STOCK_MODE', 8.5, False, GRAY)],
    PROC['CV'][1], AMBER, NAVY, 10)
footer(s, '방향 지시 관문')

# ══════════════════════════════════════════════════════════════════
# 9. WCS_TASK_CV 1사이클
# ══════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
head(s, 'WCS_TASK_CV — CvThread.Thread_Doing() 1사이클',
     '마스터 PLC 1소켓으로 15설비 슬롯을 순회한다. 한 바퀴가 약 16초라, 3초짜리 펄스는 DB 에서 놓칠 수 있다.')

box(s, 0.45, 1.30, 2.55, 0.62, ['ScanPendingWork()', ('쓰기 대기 1회 스캔', 8.5, False, GRAY)],
    PROC['CV'][1], AMBER, NAVY, 10.5)
arrow(s, 3.00, 1.61, 3.45, 1.61, LINE)
box(s, 3.45, 1.30, 2.55, 0.62, ['PreloadStatusBlocks()', ('상태 구간 일괄 조회 (왕복 45→3회)', 8.5, False, GRAY)],
    PROC['CV'][1], AMBER, NAVY, 10.5)
arrow(s, 6.00, 1.61, 6.45, 1.61, LINE)
box(s, 6.45, 1.30, 2.75, 0.62, ['지시 걸린 설비 우선 정렬', ('m_setPendCmd / Od / Trk', 8.5, False, GRAY)],
    LIGHT, GRAY, NAVY, 10.5)
arrow(s, 9.20, 1.61, 9.65, 1.61, LINE)
box(s, 9.65, 1.30, 3.25, 0.62, ['슬롯 순회 후 CvMainAlarmBits()', ('M5501~M550D 과부하 → ERROR_CODE (사이클당 1회)', 8.5, False, GRAY)],
    PROC['CV'][1], AMBER, NAVY, 10.5)

tb(s, 0.45, 2.12, 12.4, 0.3, '슬롯(설비) 하나마다 아래 6단계를 순서대로 수행한다. 실패가 소켓 끊김이면 전체 재접속, 데이터 문제면 그 설비만 건너뛴다.',
   10, False, GRAY)

slot_steps = [
    ('GetFirstAddress(Idx)', '그 설비의 CV_DATA 로 읽을 첫 주소를 구한다. 실패하면 이 설비만 skip'),
    ('CvStatusScenario(Idx)', 'PLC 상태 READ → UpdateCvData() 로 CV_DATA 미러 (센서·작업번호·H/S·방향)'),
    ('CvChg_CMD_RQ_YN(Idx)', 'CMD 지시 WRITE - 방향 전환 DIR / DIRF / DIRM (방향 D 워드, 부호는 [PLC] DIR_CODE)'),
    ('CvChg_OD_RQ_YN(Idx)', 'OD 지시 WRITE (이송 명령)'),
    ('CvEventCheck(Idx)', 'Load / Unload Complete 이벤트 ACK'),
    ('CvTrackingWrite(Idx)', 'R 영역 작업번호 트래킹 WRITE (EncodeJobNoR / GetRTrackingAddr)'),
]
y = 2.50
for i, (fn, desc) in enumerate(slot_steps, 1):
    box(s, 0.45, y, 0.42, 0.56, str(i), PROC['CV'][1], AMBER, NAVY, 11)
    box(s, 0.98, y, 4.35, 0.56, fn, WHITE, AMBER, NAVY, 10.5)
    tb(s, 5.50, y + 0.14, 7.4, 0.3, desc, 9.5, False, GRAY)
    if i < len(slot_steps):
        arrow(s, 0.66, y + 0.56, 0.66, y + 0.68, LINE, 1.25)
    y += 0.68

tb(s, 0.45, 6.55, 12.4, 0.45,
   ['※ 공통 알람 비트 처리 CvAlarmCheck(1) 은 기본 건너뛴다(WCS_DB.INI [CNF] CV_GLOBAL_ALARM=1 일 때만) - S/C#1 알람은 VehThread 가 호기별로 처리한다.',
    '※ 연속 통신 실패가 임계를 넘으면 소켓을 닫고 스레드를 종료한다 — SYS_MAIN.Thread_Tick 이 재생성해 접속 루프를 다시 돈다.'],
   9, False, GRAY)
footer(s, 'C/V 통신 1사이클')

# ══════════════════════════════════════════════════════════════════
# 10. WCS_TASK_HOST 송수신
# ══════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
head(s, 'WCS_TASK_HOST — 수신 / 송신 함수',
     '수신은 CSrvWork(서버, EcsComA.ini LocalPort), 송신은 CCliWork(클라이언트, RemoteIP:RemotePort)가 맡는다. 서로 다른 소켓이다.')

tb(s, 0.45, 1.22, 6.1, 0.3, '수신 — CSrvWork', 13, True, NAVY)
recv = [
    ('ReadRequest()', 'CheckHeader() / CheckBody() 로 헤더·본문 길이 확인'),
    ('Parsing()', '전문 종류 판별 후 분기. 수신 로그 InsertHostIfLog()'),
    ('ParseOorR()', 'O = 작업지시(66) / R = 재지정(31). InsertJobMst() 로 작업 생성'),
    ('ParseCancel()', 'D = 작업 취소 — 상태 99 / 10 / 20 만 대상'),
    ('ParseModeChange() → SetCvDirection()', 'M = 모드 전환(7). 출고 DIR 즉시 / 입고 DIRW(대기)'),
    ('MakeResponse() → SendSock()', '응답 전문 회신 (정상 또는 반려 사유코드)'),
]
y = 1.58
for fn, desc in recv:
    box(s, 0.45, y, 2.95, 0.58, fn, PROC['HOST'][1], BLUE, NAVY, 9.5)
    tb(s, 3.52, y + 0.08, 3.0, 0.5, desc, 8.5, False, GRAY)
    if fn != recv[-1][0]:
        arrow(s, 1.92, y + 0.58, 1.92, y + 0.70, LINE, 1.25)
    y += 0.70

tb(s, 6.85, 1.22, 6.1, 0.3, '송신 — CCliWork.GetSendData() 매 주기', 13, True, NAVY)
send = [
    ('GetStatusReport()', 'S = 상태 변경 시 + 30초 주기 상태보고', 0),
    ('GetErrorReport()', 'E = 설비 에러보고. 새 에러(HOST_ERR_SEND_YN=N)만', 0),
    ('IsEquip_ERROR_Modified()', 'ErrorKind : 73·74 → 1, 75 → 3, 그 밖 0', 1),
    ('GetJobCompleteReport()', 'F = 완료보고. 대상 상태 19 / 29', 0),
    ('IsJobExist(nJobStatus)', '출고 29 는 1차 보고만(WC_STEP=1)', 1),
    ('UpdateJobStatusTo(\'09\') → RequestSrv()', '응답 시 DeleteJobMst(), 실패 시 롤백', 1),
    ('ReReportComplete09()', '09 작업 : 즉시 → RE_REPORT_09_SEC 간격 재보고', 0),
    ('GetLoadArrivalReport()', '도착보고 - 대상 상태 12', 0),
]
y = 1.58
for i, (fn, desc, ind) in enumerate(send):
    dx = 0.38 * ind
    box(s, 6.85 + dx, y, 3.15 - dx, 0.52, fn, PROC['HOST'][1], BLUE, NAVY, 9.5)
    tb(s, 10.12, y + 0.06, 2.8, 0.45, desc, 8.5, False, GRAY)
    if ind:
        tb(s, 6.87, y + 0.13, 0.36, 0.26, '└', 11, True, LINE, PP_ALIGN.CENTER)
    y += 0.62
tb(s, 6.85, y - 0.02, 6.05, 0.3, '└ 표시는 위 함수가 내부에서 부르는 함수다.', 8.5, False, LINE)

tb(s, 0.45, 6.25, 6.1, 0.7,
   ['※ 무게보고(U)·빈파렛트 요청(P)은 이 현장이 재고를 관리하지 않아 발신하지 않는다',
    '   (송신 함수 자체를 정리했다).',
    '※ 반자동·수동(작업구분 10~15, 작업번호 9000번대)은 완료 보고 대상이 아니다.'],
   9, False, GRAY)
footer(s, '상위 통신 함수')

# ══════════════════════════════════════════════════════════════════
# 11. 알람 · 에러 함수 흐름
# ══════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
head(s, '알람 · 에러 — 함수 흐름 (PLC 알람 리스트 2026-09-17)',
     '설비 통신이 알람을 코드로 기록하고, 상위 통신이 E 전문으로 보고한다. 스케줄러는 로그만 남기고 에러 설비를 지시에서 뺀다.')
legend(s)
chain(s, [
    ('VEH', '01', 'VehThread.PollObservations()  (S/C)',
     '알람코드 워드 → ERR_CODE_RD. [CNF] SC_ERR_CODE_BLOCK 0=워드 상시(현장)'),
    ('VEH', '02', 'ReadRgvAlarmBitCode()  (RGV)',
     '코드 워드가 0 이면 M5601~M560F 중 가장 낮은 켜진 비트의 코드'),
    ('VEH', '03', 'InsertEqpErrHis()',
     '새 에러 → EQP_ERR_HIS 적재, HOST_ERR_SEND_YN = N'),
    ('CV', '04', 'CvThread.CvMainAlarmBits()',
     "M5501~M550D → CV_DATA.ERROR_CODE '0119'~'0131'. [CNF] MAIN_ALM_BITS"),
    ('CV', '05', 'UpdateEQMT_ERR_LOG()',
     '메인 C/V 과부하 발생 → EQP_ERR_HIS 적재'),
    ('SCH', '06', 'MonitorAlarm() → CheckAlarm()',
     'CV / SC / RGV 의 새 알람 코드를 로그로만 남긴다'),
    ('SCH', '07', 'DriveCV() / DriveSC() / DriveRGV()',
     '에러 코드가 있는 설비·라인(IsCvError)에는 지시하지 않는다'),
    ('HOST', '08', 'GetErrorReport(EQP_TYP)',
     'CV / SC / RTV 순서로 새 에러를 찾는다'),
    ('HOST', '09', 'IsEquip_ERROR_Modified()',
     'EQP_ECD_MST(SC_LGLS) · modDefApp.IsCodeIn() 으로 ErrorKind 판정'),
    ('HOST', '10', 'RequestSrv()  — E 전문',
     'DeviceClass · DeviceNo · ErrorKind · ErrorCode · 작업번호 · 셀'),
    ('SCH', '11', '(호출 안 함) MarkErrorJobStatus()',
     '이중입고/공출고 → 08/07 상태 표시. "이 현장 없음"(09-16)'),
    ('SCH', '12', '(호출 안 함) ResumeRedirectedJobs()',
     '상위 R 재지정(06/05) → 새 셀로 재개. 정의만 보존'),
], cols=3, w=3.94, h=0.74, gap=0.95, colgap=4.24)
box(s, 0.45, 5.4, 12.45, 1.45,
    [('ErrorKind 판정 코드 (EcsComA.ini [Host])', 11, True, NAVY),
     ('ScDualCodes = 73,74 (좌/우측 렉 이중입고 → 1)     ScEmptyCodes = 75 (공출고 → 3)     ScInFailCodes / ScOutFailCodes = 없음     그 밖 → 0 (기계적)', 9.5, False, GRAY),
     ('크레인 코드표 구분 ScErrCodeType = SC_LGLS (종전 SC_SFA 의 54·55 / 58·59 는 이 현장 코드가 아니다)', 9.5, False, GRAY),
     ('설비 에러 해제는 설비(조작반)에서 한다. 운전 화면에는 에러 해제 기능이 없다.', 9.5, False, GRAY)],
    LIGHT, NAVY, GRAY, 9.5, MSO_SHAPE.ROUNDED_RECTANGLE, False)
footer(s, '알람 · 에러 흐름')

# ══════════════════════════════════════════════════════════════════
# 12. 예외 · 정리 함수
# ══════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
head(s, '예외 · 정리 함수 — 체류 / 강제완료 / 잔류 정리',
     '완료는 설비 신호로만 판정한다. 시간 기반 자동 완료(체류 복구)는 2026-09-17 폐기했다.')
exc = [
    ('CheckStalledJobs()', '1분 주기. 19/29 가 아닌 작업이 JOB_STALL_WARN_SEC(300초) 넘게 멈추면 경고만 남긴다', RED),
    ('CleanupAutoTimeSets()', 'JOB_MST 에 없는 작업번호를 시간 기반 허용·사용 이력에서 지운다', GRAY),
    ('AutoTimeProcEnabled()', "COMMON_CODE SCH_OPT/AUTO_TIME = 'Y' 일 때만 CompleteSC/CompleteRGVReal 의 경과시간 분기를 연다(기본 N)", GRAY),
    ('ConsumeForceComplete()', '운전 화면 [강제완료](FCMP) 소비 → ForceCompleteSc() 25→29 / ForceCompleteRtv() 35→39', GREEN),
    ('  ForceCompleteSc() / ForceCompleteRtv()', '설비가 화물을 들고 있으면 처리하지 않는다 (사람이 이미 내려놓았다는 전제)', GREEN),
    ('SweepOrphanVehicleData()', '지시 소비 · 차상 빔 · 유휴 · JOB_MST 에 없음 → ClearRvData() / ClearScOd() / ResetScComplete()', AMBER),
    ('CompleteRGVManual() / CompleteSCManual()', '수동지시 9998 / 9999 완료 신호 → 지시 흔적 정리, 이력 로그', AMBER),
    ('DeleteJobNow()', '반자동 입고(11) 크레인 완료 · 반자동 출고(12) RTV 도착 기록 확인 시 즉시 삭제 (상위 보고 없음)', BLUE),
]
y = 1.3
for i, (fn, desc, col) in enumerate(exc):
    box(s, 0.45, y, 4.3, 0.56, fn, WHITE, col, NAVY, 10)
    tb(s, 4.92, y + 0.08, 8.0, 0.45, desc, 9.5, False, GRAY)
    if i < len(exc) - 1:
        arrow(s, 2.60, y + 0.56, 2.60, y + 0.66, LINE, 1.25)
    y += 0.66
tb(s, 0.45, 6.62, 12.45, 0.35,
   '※ 체류 경고가 뜨면 설비 상태와 화물 위치를 사람이 확인한다. 작업정보를 자동으로 지우면 실물 위치를 되찾을 길이 없다(유령 재고 방지).',
   9.5, False, RED)
footer(s, '예외 · 정리 함수')

out = os.path.join(OUT, '09_함수단_흐름도.pptx')
prs.save(out)
print('saved', out, '|', page[0], 'slides')
