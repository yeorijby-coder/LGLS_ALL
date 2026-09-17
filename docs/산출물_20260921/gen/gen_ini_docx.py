# -*- coding: utf-8 -*-
"""12_INI_설정_정의서.docx 생성 (2026-09-21 최종판)

현재값은 실행 폴더의 INI 를 직접 읽어 채운다(정리 끝난 최종본, CP949).
기본값(코드) · 반영 시점 · 읽는 코드는 2026-09-17 소스 확인 결과를 적었다.
"""
import os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from gen_common import (new_doc, table, para, bullets, OUT, ROOT, GRAY, NAVY,
                        read_cp949, DATE)
from docx.shared import RGBColor

RED = RGBColor(0xC0, 0x00, 0x00)

INI_CV = os.path.join(ROOT, r'TASK\WCS_TASK_CV_BIN\bin\Debug\WCS_DB.INI')
INI_IO = os.path.join(ROOT, r'TASK\IO_TASK\bin\Debug\ENV_IOSCH.INI')
INI_HOST = os.path.join(ROOT, r'TASK\WCS_TASK_HOST\bin\Debug\EcsComA.ini')
INI_CLI = os.path.join(ROOT, r'EXE_NEWUI\WCS_CLIENT\Ecs.ini')


def load_ini(path):
    """{(섹션, 키): 값}. ';' '#' '--' 로 시작하는 줄은 주석. 값 뒤 공백은 지운다."""
    d = {}
    sec = ''
    for raw in read_cp949(path).splitlines():
        s = raw.strip()
        if not s or s[0] in ';#' or s.startswith('--'):
            continue
        m = re.match(r'\[([^\]]+)\]', s)
        if m:
            sec = m.group(1).strip(); continue
        if '=' in s:
            k, v = s.split('=', 1)
            d[(sec, k.strip())] = v.strip()
    return d


CUR = {'CV': load_ini(INI_CV), 'IO': load_ini(INI_IO), 'HOST': load_ini(INI_HOST), 'CLI': load_ini(INI_CLI)}


def cur(prog, sec, key):
    if ' ' in key:                      # 여러 키를 묶은 행
        return '아래 표'
    v = CUR[prog].get((sec, key))
    if v is None:
        return '(키 없음)'
    return v if v != '' else '(빈 값)'


HEAD = ['키', '현재값', '기본값(코드)', '값 범위 · 뜻', '반영 시점', '읽는 코드', '현장값 · 비고']
W = [3.4, 2.3, 2.2, 8.6, 2.2, 4.2, 2.8]


def keytable(d, prog, sec, rows):
    """rows : [키, 기본값, 뜻, 반영, 코드, 비고]"""
    out = []
    for k, dflt, mean, when, code, note in rows:
        out.append([k, cur(prog, sec, k), dflt, mean, when, code, note])
    table(d, HEAD, out, widths=W, font=8)


# 반영 시점 표기
NOW2 = '즉시(2초)'        # cDefApi.GsCnfInt : 2초 캐시
CALL = '즉시(호출마다)'   # 쓸 때마다 GetPrivateProfile*
RST = '재기동'
DBL = '재기동(DB 재접속 시 재읽기)'
HOT = '즉시(핫 리로드)'   # Client : 1초 주기 저장 감지 + 1초 확인
CYC = '즉시(주기마다)'

d = new_doc('INI 설정 정의서', '설비 통신 · 스케줄러 · 상위 통신 · 운전 화면', landscape=True, history=[])

# ════════════════════════════════════════════════════════════════════
d.add_heading('1. 개요', 1)
para(d, '이 문서는 LG 화학 1동 자동창고 ECS 리뉴얼의 실행 설정 파일(INI) 4종을 키 단위로 정의한다. '
        '현재값은 %s 기준 실행 폴더의 파일(미사용 키 정리를 마친 최종본)에서 읽었고, 기본값과 반영 시점은 '
        '소스 코드에서 확인했다. 기본값(코드)은 키가 없을 때 프로그램이 쓰는 값이다.' % '2026-09-17')
table(d, ['프로그램', '파일', '실행 폴더 (개발 PC)', '인코딩', '읽는 방식'], [
    ['설비 통신 (WCS_TASK_CV)', 'WCS_DB.INI', r'TASK\WCS_TASK_CV_BIN\bin\Debug', 'CP949',
     r'작업 폴더 기준 ./WCS_DB.INI (cDefApp.GM_ENV_INI). 주소 정의는 7_DeviceMap\PlcAddressMap.xml'],
    ['스케줄러 (IO_TASK)', 'ENV_IOSCH.INI', r'TASK\IO_TASK\bin\Debug', 'CP949',
     '작업 폴더 기준 ./ENV_IOSCH.INI (cDefApp.GM_ENV_INI)'],
    ['상위 통신 (WCS_TASK_HOST)', 'EcsComA.ini', r'TASK\WCS_TASK_HOST\bin\Debug', 'CP949',
     r'작업 폴더 기준 .\ECSCOMA.ini (modDefApp.MAIN_INI). RE_REPORT_09_SEC 만 실행 파일 폴더 기준'],
    ['운전 화면 (Ecs.exe)', 'Ecs.ini', r'EXE_NEWUI\WCS_CLIENT (운영 실행본)', 'CP949',
     r'실행 파일 폴더\Ecs.ini (ECS_INI_FILE = g_strEcsPath + "\Ecs.ini")'],
], widths=[4.2, 3.0, 6.0, 1.8, 10.7])
para(d, r'※ 통합 실행본 TASK\ALL_TASK\bin\Debug 는 위 세 TASK INI(WCS_DB.INI · ENV_IOSCH.INI · EcsComA.ini)를 '
        r'같은 내용으로 둔다(2026-09-17 비교 결과 동일). ALL_TASK 자신의 설정은 ALL_TASK.INI([LOG] DIR/KEEP_DAYS, '
        r'[AUTOSTART] EQP/IO/HOST)뿐이다.', size=9, color=GRAY)

d.add_heading('1.1 반영 시점 표기', 2)
table(d, ['표기', '뜻', '해당 코드'], [
    [NOW2, '저장 후 2초 안에 반영. 재기동 불필요', 'WCS_TASK_CV cDefApi.GsCnfInt (키별 값을 2초 캐시)'],
    [CALL, '값을 쓰는 순간마다 파일을 읽음. 재기동 불필요', 'GetPrivateProfileInt/String 직접 호출 (IO_TASK GsReadInitProfileCnf 등)'],
    [CYC, '주기 루프에서 매번 읽음. 다음 주기부터 반영', 'Client 수집 스레드, HOST 재보고(5초 캐시)'],
    [HOT, 'Client 가 1초마다 Ecs.ini 수정 시각을 보고, 바뀐 뒤 1초 더 확인하고 다시 읽음 (UI 로그 [INI] reload)',
     'CEcsView::CheckIniHotReload → ReloadIniHot'],
    [RST, '기동 시 한 번 읽음. 바꾸면 프로그램을 다시 띄워야 함', '폼 Load / 스레드 시작부'],
    [DBL, 'HOST 는 DB 로그인(DBLogIn)마다 ReadInitProfile 을 다시 부르지만, 수신 포트 바인딩 등은 기동 때 한 번', 'modCmLib.ReadInitProfile'],
], widths=[3.5, 12.0, 10.2])

# ════════════════════════════════════════════════════════════════════
d.add_heading('2. 요약 - 주요 운영 키', 1)
para(d, '현장 운전 중 조정할 가능성이 높은 키만 모았다. 상세 정의는 3~6장.')
table(d, ['프로그램 / 섹션.키', '현재값', '기본값', '요점', '반영'], [
    ['설비 통신 [CNF] SC_ERR_CODE_BLOCK', cur('CV', 'CNF', 'SC_ERR_CODE_BLOCK'), '0',
     '크레인 에러코드 판정 방식 (0/1/2/3/4, 아래 2.1). 크레인 3·5호기 유령 에러 대책으로 현장 권장 3', CALL],
    ['설비 통신 [PLC] DIR_CODE', cur('CV', 'PLC', 'DIR_CODE'), 'IN0',
     'C/V 방향 워드 부호. IN0 = 입고 0/출고 1, IN1 = 입고 1/출고 0. 시뮬레이터는 IN0 만 이해. PLC 담당 확인값은 IN1', RST + ' (화면 라디오는 즉시)'],
    ['설비 통신 [CNF] BIT_WRITE', cur('CV', 'CNF', 'BIT_WRITE'), '1',
     'M 비트 쓰기 : 1 = 단일 비트(%MX, 구 ECS 와 같음) / 0 = 워드 읽고-바꿔-되쓰기(비교용)', CALL],
    ['설비 통신 [CNF] CV_BULK_R · CV_EVT_BULK · CV_HB_MS · CV_CYCLE_SLEEP_MS · CV_CYCLE_LOG_MS · UI_LOG_ASYNC',
     '%s · %s · %s · %s · %s · %s' % tuple(cur('CV', 'CNF', k) for k in
        ('CV_BULK_R', 'CV_EVT_BULK', 'CV_HB_MS', 'CV_CYCLE_SLEEP_MS', 'CV_CYCLE_LOG_MS', 'UI_LOG_ASYNC')),
     '1 · 1 · 1000 · 200 · 1000 · 1', 'C/V 순회 속도 옵션 (일괄 읽기, 통신상태 갱신 간격, 한 바퀴 휴지, 느린 사이클 로그, 화면 로그 비동기)', NOW2],
    ['설비 통신 [CNF] MAIN_ALM_BITS', cur('CV', 'CNF', 'MAIN_ALM_BITS'), '1',
     '메인 PLC 알람 비트(M5501~/M5601~) 읽기 → C/V 모터 과부하 · RGV 알람', NOW2],
    ['설비 통신 [CNF] SC_ERR_TYP', cur('CV', 'CNF', 'SC_ERR_TYP'), 'SC_LGLS',
     '에러이력에 남길 크레인 코드표 구분 (SC_LGLS = PLC 알람 리스트 260917 / SC_SFA = 종전)', RST],
    ['스케줄러 [CNF] CV2_DIR_MODE · CV2_FAIR · CV2_STALL_SEC',
     '%s · %s · %s' % (cur('IO', 'CNF', 'CV2_DIR_MODE'), cur('IO', 'CNF', 'CV2_FAIR'), cur('IO', 'CNF', 'CV2_STALL_SEC')),
     '1 · 0 · 60', 'S/C#1 통로(C/V#2 103) 방향 전환 방식 (아래 2.2)', CALL],
    ['스케줄러 [CNF] IN_HS_STATUS', cur('IO', 'CNF', 'IN_HS_STATUS'), '15',
     '입고가 통로(H/S)에 내려진 뒤 작업상태 15(CV 구동중) / 16(통로CV 구동중). 스케줄러는 둘 다 읽음', CALL],
    ['스케줄러 [CNF] OUT_HOLD_SEC', cur('IO', 'CNF', 'OUT_HOLD_SEC'), '3',
     '1호기 출고 완료 직후 통로 출고 유지(초). CV2_DIR_MODE=1 에서는 쓰지 않음', CALL],
    ['스케줄러 [CNF] MANUAL_DIR_HOLD_SEC', cur('IO', 'CNF', 'MANUAL_DIR_HOLD_SEC'), '120',
     'Client 더블클릭 수동 방향 전환 뒤 자동 방향 정합을 멈추는 시간(초). 0 = 억제 없음', CALL],
    ['스케줄러 [CNF] SC_DUAL_CODES · SC_EMPTY_CODES',
     '%s · %s' % (cur('IO', 'CNF', 'SC_DUAL_CODES'), cur('IO', 'CNF', 'SC_EMPTY_CODES')), '73,74 · 75',
     '이중입고/공출고 크레인 에러코드. 현재는 재지정 지시 허용 판정에만 쓰임', CALL],
    ['상위 통신 [Host] ScErrCodeType · ScDualCodes · ScEmptyCodes',
     '%s · %s · %s' % (cur('HOST', 'Host', 'ScErrCodeType'), cur('HOST', 'Host', 'ScDualCodes'), cur('HOST', 'Host', 'ScEmptyCodes')),
     'SC_LGLS · 73,74 · 75', '크레인 에러의 상위 E 전문 ErrorKind 판정 코드', DBL],
    ['상위 통신 [Host] RE_REPORT_09_SEC', cur('HOST', 'Host', 'RE_REPORT_09_SEC'), '60',
     '완료(F) 보고 뒤 응답 없는 09 잔류 작업 재보고 간격(초, 5~3600)', CYC + ' (5초)'],
    ['상위 통신 [Host] StationMapMode', cur('HOST', 'Host', 'StationMapMode'), 'WMS',
     'IMS 스테이션 코드 해석 기준 (WMS = 운영 DB 기준 / ECS = 구 ECS 구축시드 기준)', DBL],
    ['운전 화면 [SC_ERR] ERR_TYP · DUAL_CODES · EMPTY_CODES',
     '%s · %s · %s' % (cur('CLI', 'SC_ERR', 'ERR_TYP'), cur('CLI', 'SC_ERR', 'DUAL_CODES'), cur('CLI', 'SC_ERR', 'EMPTY_CODES')),
     'SC_LGLS · 73,74 · 75', '상태창 알람 문구 코드표, 에러해제/삭제 확인창 판정', CALL],
    ['운전 화면 [MENU] VEH_CLEAR_MODE', cur('CLI', 'MENU', 'VEH_CLEAR_MODE'), '3',
     '크레인 · RTV 색과 작업번호를 지우는 시점 (0 종전 / 1 적재 비트 / 2 H/S 기록 / 3 표시는 0, 지움은 새 트랙 기록)', HOT],
    ['운전 화면 [ETC] COLLECT_INTERVAL_MS', cur('CLI', 'ETC', 'COLLECT_INTERVAL_MS'), '300',
     '설비 데이터 DB 수집 주기(ms, 100~5000)', CYC],
], widths=[6.2, 3.0, 2.8, 11.2, 2.5], font=8)

d.add_heading('2.1 SC_ERR_CODE_BLOCK 모드 (설비 통신 WCS_DB.INI [CNF])', 2)
table(d, ['모드', '에러 세움 (ERR_CODE_RD)', '에러 내림', 'DOWN 이면 에러색', '비고'], [
    ['0', '알람코드 워드(ALARM_SET_CODE, D0161+10k)가 0 아니면 매 주기', '워드가 0 이 되면', '아니오',
     '코드 기본값. 구 ECS 배차 규약. 보고 비트 없이 워드 값만 실리면 유령 에러(현장 3·5호기 0001, 09-15/16)'],
    ['2', 'ALARM_SET_REPORT 보고 비트가 설 때 코드 래치', 'ALARM_RESET_REPORT 보고', '아니오', '구 ECS 팝업 알람 LED 규약'],
    ['3', '2 와 같음', 'RESET 보고 또는 워드 0', '아니오',
     '현장 권장. RESET 펄스를 약 16초 순회 사이에 놓쳐도 에러가 내려감'],
    ['4', '2 와 같음', 'RESET 보고', '예 (ERR_STA_FK1_RD=3, 상위 E 전문 없음)', '구 ECS 메인 화면 규약'],
    ['1', '2 와 같음', 'RESET 보고 또는 워드 0', '예', '구 ECS 와 가장 가까움. 종전 "ErrCode 블록" 용도는 폐기'],
], widths=[1.3, 6.5, 4.5, 4.5, 8.9], font=8.5)
para(d, '※ RTV 에는 적용되지 않는다(RTV 는 워드 상시 + 래치 병행). 현재 개발 PC 값은 %s 이다. '
        '현장 반영 시 3 으로 두고 크레인 3·5호기 에러 표시를 확인한다.' % cur('CV', 'CNF', 'SC_ERR_CODE_BLOCK'),
     size=9, color=GRAY)

d.add_heading('2.2 CV2_DIR_MODE 모드 (스케줄러 ENV_IOSCH.INI [CNF])', 2)
table(d, ['모드', '동작', '관련 키'], [
    ['0', '종전 : 입고 진행 우선 → 출고 진행 → 대기 순으로 수시 전환', 'OUT_HOLD_SEC'],
    ['1 (코드 기본)', '교착 때만 전환 : 통로를 현재 방향으로 쓰는 작업(출고 25/29/16, 입고 35/39/25/16/15 실물)이 있으면 유지, '
     '없으면 지시 직전에 전환. 있어도 CV2_STALL_SEC 동안 무진행(크레인 25 제외)이면 한 번 전환(DIRF 즉시 쓰기). '
     'RGV 1호기행 입고 지시는 출고 29 · 16 도 점유로 보고 보류', 'CV2_STALL_SEC (기본 60, 최소 10)'],
    ['2 (현재값)', '통로 사용권 : 크레인 1호기 출고 지시(25) 시 출고 사용권, RTV 1호기행 입고 지시(35) 시 입고 사용권. '
     '상대 사용권이 있으면 지시하지 않음. 방향은 그 지시를 낼 때만 바꿈. 출고 사용권은 RTV 가 싣고 통로를 벗어나면, '
     '입고 사용권은 크레인 1호기 입고 지시 때 풀림. 사용권은 작업상태(DB)로 판정하므로 재기동해도 유지',
     'CV2_FAIR : 0 = 출고 우선(기본) / 1 = 번갈아 쓰기(더 오래 기다린 쪽)'],
], widths=[2.8, 16.5, 6.4], font=8.5)

d.add_heading('2.3 주의 - [CNF] 섹션 위치', 2)
para(d, '2026-09-17 전까지 스케줄러 ENV_IOSCH.INI 에서 CV2_DIR_MODE · CV2_STALL_SEC · CV2_FAIR 3키가 [VIEW] 섹션 아래에 '
        '잘못 들어가 있었다. 스케줄러는 이 키들을 [CNF] 섹션에서만 읽으므로 그동안 값이 적용되지 않고 코드 기본값'
        '(CV2_DIR_MODE=1, CV2_STALL_SEC=60, CV2_FAIR=0)으로 동작했다. 2026-09-17 에 [CNF] 로 옮겼다.', bold=True, color=RED)
bullets(d, [
    '현장 PC 의 ENV_IOSCH.INI 도 반드시 섹션 위치를 확인한다. [VIEW] 아래(LOG_FILTER_BTN 뒤)에 붙어 있으면 [CNF] 블록 안으로 옮긴다.',
    'INI 는 "마지막 [섹션] 머리 뒤" 에 있는 줄을 그 섹션 키로 본다. 파일 끝에 키를 덧붙일 때는 마지막 섹션이 무엇인지 먼저 본다 '
    '(ENV_IOSCH.INI · WCS_DB.INI 모두 마지막 섹션이 [VIEW] 다).',
    '스케줄러는 키가 바뀌면 호출마다 읽으므로 옮긴 즉시 반영된다. 로그에서 통로 사용권/교착 전환 메시지로 확인한다.',
])

# ════════════════════════════════════════════════════════════════════
d.add_heading('3. 설비 통신 - WCS_DB.INI', 1)
CVA = r'4_Static\cDefApi.cs'
d.add_heading('3.1 [DB_2] DB 접속', 2)
keytable(d, 'CV', 'DB_2', [
    ['IP', '없음', 'SQL Server 주소 "서버,포트" (접속 문자열 Server=)', RST, CVA + ':GsGetInitPorFileDB_2 ← SYS_MAIN', '현장 : 서버IP,1433'],
    ['DATABASE', '없음', 'DB 명', RST, CVA + ':GsGetInitPorFileDB_2', 'LGLS_MCS_IO'],
    ['PORT', '없음', '읽기만 하고 SQL 빌드 접속 문자열에는 쓰지 않음 (PostgreSQL 잔재). 포트는 IP 뒤 ",포트" 로 지정', '-', CVA + ':GsGetInitPorFileDB_2', '미사용'],
    ['USER', '없음', 'SQL 로그인', RST, CVA + ':GsGetInitPorFileDB_2', ''],
    ['USER_PW', '없음', 'SQL 암호', RST, CVA + ':GsGetInitPorFileDB_2', ''],
])
para(d, '※ 섹션 머리의 "--PostgreSql" 표기는 옛 주석이다. 빌드 정의가 SQL 이라 MS SQL Server 로 접속한다.', size=9, color=GRAY)

d.add_heading('3.2 [PLC] PLC 규약', 2)
keytable(d, 'CV', 'PLC', [
    ['MAKER', '1', '1 = XGT(LS 산전 FEnet) / 0 = Melsec(Q3E)', RST, CVA + ':GsReadInitProfilePlcMaker ← SYS_MAIN, CvThread', ''],
    ['R_ADDR_MODE', 'HEX', 'R(트래킹) 문서표기 해석. DEC = 10진(현행 확정) / HEX = 구 ECS 16진. '
     '주소맵 XML rAddrMode 가 단일 기준이고, XML 이 있으면 기동 시 XML 값으로 INI 를 덮어씀', RST, CVA + ':GsReadInitProfileRAddrHex (XML 없을 때만)', 'DEC'],
    ['D_ADDR_MODE', 'DOC', 'D 블록 해석. DOC = 문서표기 10진 / LEGACY = 구 ezMCS 환산. XML dAddrMode 가 기준, INI 는 동기화 사본', RST, CVA + ':GsReadInitProfileDAddrDoc', 'DOC'],
    ['DIR_CODE', 'IN0', 'C/V 방향 워드 부호. IN0 = 입고 0/출고 1 / IN1 = 입고 1/출고 0. 쓰기 · 판독 양쪽 적용. '
     '메인 화면 셋째 줄 라디오로 바꾸면 즉시 반영되고 INI 에 저장됨', RST + ' (라디오는 즉시)',
     CVA + ':GsReadInitProfileDirIn1 ← SYS_MAIN 기동', 'PLC 담당 확인값 IN1. 시뮬은 IN0'],
])

d.add_heading('3.3 [PROCESS]', 2)
keytable(d, 'CV', 'PROCESS', [
    ['CNT', '1', '읽을 [COMMn] 섹션 수. 마스터 PLC 1소켓 구조라 1', RST, CVA + ':GsReadInitProfileProcessCnt', ''],
])

d.add_heading('3.4 [CNF] 운전 옵션', 2)
CVT = r'CvThread.cs'
keytable(d, 'CV', 'CNF', [
    ['CV_BULK_R', '1', '1 = R 트래킹(작업번호)을 128워드 이하 조각으로 한꺼번에 읽음 / 0 = 설비별 개별 READ(종전)', NOW2, CVT + ':GsCnfInt', ''],
    ['CV_EVT_BULK', '1', '1 = 이벤트 확인에 사이클 시작의 M 일괄 값 사용 / 0 = 설비별 개별 READ', NOW2, CVT + ':GsCnfInt', ''],
    ['CV_HB_MS', '1000', '통신상태(EQP_MST) 갱신 간격 ms. 0 = 설비마다(종전). Client 끊김 판정은 5초', NOW2, CVT + ':GsCnfInt', ''],
    ['CV_CYCLE_SLEEP_MS', '200', '한 바퀴 뒤 쉬는 시간 ms (코드에서 0~2000 으로 자름)', NOW2, CVT + ':GsCnfInt', ''],
    ['CV_CYCLE_LOG_MS', '1000 (CYCLE_WARN_MS)', '한 바퀴가 이 값(ms)을 넘으면 30초에 한 번 [CYCLE] 로그', NOW2, CVT + ':GsCnfInt', ''],
    ['UI_LOG_ASYNC', '1', '1 = 화면 로그 비동기 / 0 = 동기(통신 스레드가 화면을 기다려 느려짐)', NOW2, r'SYS_MAIN.cs:GsCnfInt', ''],
    ['SC_ERR_TYP', 'SC_LGLS', '설비 에러이력(EQP_ERR_HIS)의 크레인 EQP_TYP. SC_LGLS = PLC 알람 리스트(260917) 지상반 코드 11~93, 101~118 / SC_SFA = 종전', RST,
     CVA + ':GsReadInitProfileErrCodeTyp ← SYS_MAIN', ''],
    ['MAIN_ALM_BITS', '1', '1 = 메인 PLC 알람 비트 읽음 (주소맵 <Global name="MainAlarm">) : M5501~M550D → CV119~131 모터 과부하 '
     '(CV_DATA.ERROR_CODE 0119~0131), M5601~M560F → RGV 알람(알람코드 워드가 0 일 때) / 0 = 읽지 않음', NOW2,
     'CvThread.cs · VehThread.cs:GsCnfInt', ''],
    ['RTV_ERR_TYP', 'RTV', '설비 에러이력의 RTV EQP_TYP', RST, CVA + ':GsReadInitProfileErrCodeTyp', ''],
    ['WH_TYP', '없음', '창고 구분 (DB 조회 조건)', RST, CVA + ':GsGetInitPorFileCNF', '10'],
    ['USERID', '없음', '읽어 두기만 하고 쓰는 곳 없음 (cDefApp.GM_USERID)', '-', CVA + ':GsGetInitPorFileCNF', '미사용'],
    ['CV_GLOBAL_ALARM', '0', 'CV 스레드의 "전 설비 공통 알람"(M786/787, Ack M1539/1540) 처리. 실은 S/C#1 알람 비트라 VehThread 와 중복이고 '
     'S/C#1 Ack 비트를 지울 수 있어 0(끔). PLC 가 별도 공통 알람으로 쓰는 것이 확인될 때만 1', RST, 'CvThread.cs (첫 사이클 1회)', ''],
    ['BIT_WRITE', '1', '1 = 단일 비트 쓰기(%MX, 구 ECS mdDevSet/mdDevRst 와 같음) / 0 = 워드 읽고-바꿔-되쓰기(비교용). '
     'V0.9 주소모드와 Melsec 은 항상 종전 방식', CALL, r'7_Socket\FenetProtocol.cs:GsReadInitProfileBitWrite', ''],
    ['SC_ERR_CODE_BLOCK', '0', '크레인 에러코드 판정 방식 0/1/2/3/4 (2.1 절). RTV 제외', CALL,
     r'VehThread.cs:GsReadInitProfileScErrCodeBlock', '현장 권장 3'],
])
para(d, '※ 코드에는 있으나 INI 에 두지 않은 키 : [PLC] CV_DTRACK_WRITE (기본 OFF, 구 트랙테이블 D 지시 쓰기), '
        '[DELAY] 계열(GsReadInitProfileDelay). 키가 없으면 기본값으로 동작한다.', size=9, color=GRAY)

d.add_heading('3.5 [COMM0] 마스터 PLC 접속', 2)
keytable(d, 'CV', 'COMM0', [
    ['EQMT', '없음', '설비 종류 표기 (CV)', RST, CVA + ':GsReadInitProfileCom', ''],
    ['IP', '없음', '마스터 PLC IP. SC/RTV 스레드도 같은 IP:2004 를 씀 (비면 127.0.0.1)', RST, CVA + ':GsReadInitProfileCom ← SYS_MAIN', '현장 192.100.1.110'],
    ['CUR_PORT', '호출 인자', '현재 접속 포트. 포트 전환 시 프로그램이 INI 에 되씀', RST, CVA + ':GsReadInitProfileCom / CvThread 기록', '2004'],
    ['FROM_PORT', '호출 인자', '포트 순환 범위 시작', RST, CVA + ':GsReadInitProfileCom', '2004'],
    ['TO_PORT', '호출 인자', '포트 순환 범위 끝', RST, CVA + ':GsReadInitProfileCom', '2004'],
    ['PORT_CNT', 'CNT 값', '포트 수', RST, CVA + ':GsReadInitProfileCom', '1'],
    ['CNT', '호출 인자', '이 COMM 의 설비 수 표기', RST, CVA + ':GsReadInitProfileCom', '1'],
    ['PLC_NO', '없음 (비면 순회 중단)', 'PLC 번호. 비어 있으면 그 COMM 부터 읽지 않음', RST, CVA + ':GsReadInitProfileCom', '11'],
    ['FR_TRACK', '1', '시작 트랙 번호', RST, CVA + ':GsReadInitProfileCom', '121'],
    ['TO_TRACK', '1', '끝 트랙 번호', RST, CVA + ':GsReadInitProfileCom', '122'],
    ['LOG_PATH', '없음', '로그 폴더', RST, CVA + ':GsReadInitProfileCom', r'C:\WCS_TASK\CV'],
    ['FILENAME', '없음', '로그 파일 이름 머리', RST, CVA + ':GsReadInitProfileCom', ''],
])
para(d, '※ 파일에는 현장 주소가 주석(--IP=192.100.1.110 ; 현장 PLC)으로 남아 있다. 현장 반영 시 IP 줄을 이 값으로 바꾼다. '
        '값 뒤 공백(IP=127.0.0.1␣)은 GetPrivateProfileString 이 잘라낸다.', size=9, color=GRAY)

d.add_heading('3.6 [VIEW] 화면', 2)
keytable(d, 'CV', 'VIEW', [
    ['LOG_FILTER_BTN', 'Y', 'Y = 메인 화면 [로그 필터] 버튼 표시 / N = 숨김 (빈 값 · 키 없음 = Y)', RST, r'SYS_MAIN.cs:GsReadInitProfileLogFilterBtn', ''],
    ['SCENARIO_BTN', 'Y', 'Y = [시나리오 테스트] 버튼 표시 / N = 숨김', RST, r'SYS_MAIN.cs:GsReadInitProfileScenarioBtn', '현장 N 검토'],
])

d.add_heading('3.7 PlcAddressMap.xml 의 역할', 2)
para(d, r'PLC 주소 정의는 INI 가 아니라 TASK\WCS_TASK_CV_BIN\7_DeviceMap\PlcAddressMap.xml 하나가 단일 기준이다. '
        r'수정하면 bin\Debug\7_DeviceMap\ 에도 같은 파일을 복사해야 한다(2026-09-17 두 파일 동일).')
bullets(d, [
    '설비 통신 자동운전 : C/V 15대 · S/C 5대 · RGV 1대의 이벤트/Ack 비트, 방향 D 워드, R 트래킹, 알람 비트를 '
    '"origin + (설비번호 - numberFrom) x stride + offset" 블록식으로 정의(<EquipGroup>/<Block>/<Signal>).',
    '전역 비트 : <Global name="Alarm">(S/C#1 대표 알람), <Global name="MainAlarm">(메인 C/V 과부하 · RGV 알람, [CNF] MAIN_ALM_BITS 로 on/off).',
    '크레인 · C/V 대응표 <CraneMap> (S/C#1 은 방향전환형 C/V#2 겸용), 설비 목록 <Equipments>.',
    '주소 해석 모드 rAddrMode(DEC) · dAddrMode(DOC) 속성이 INI 의 R_ADDR_MODE · D_ADDR_MODE 보다 우선한다. '
    '설비 통신 화면 라디오로 바꾸면 XML 과 INI 를 함께 고친다.',
    '같은 파일을 EQP_SIM(시뮬레이터), PLC 메모리 맵 화면, 시나리오 테스트(<Scenarios>)가 공유한다.',
    'Client 는 이 파일을 읽지 않는다. 방향 전환 확인창의 주소 표기는 Ecs.ini [MENU] DIR_ADDR_CV11/CV2 로 맞춘다.',
])

# ════════════════════════════════════════════════════════════════════
d.add_heading('4. 스케줄러 - ENV_IOSCH.INI', 1)
IOA = r'CLS\cDefApi.cs'
SCH = r'CLS\Thread\cThread_SCH.cs'
d.add_heading('4.1 [DB] (Oracle 빌드 전용)', 2)
keytable(d, 'IO', 'DB', [
    ['PROVIDER', '없음', 'Oracle OLE DB 공급자', '-', IOA + ':GsGetInitPorFileDB (#if ORACLE)', '미사용'],
    ['ALIAS', '없음', 'Oracle 별칭', '-', IOA + ':GsGetInitPorFileDB', '미사용'],
    ['USERID', '없음', 'Oracle 계정', '-', IOA + ':GsGetInitPorFileDB', '미사용'],
    ['PASSWORD', '없음', 'Oracle 암호', '-', IOA + ':GsGetInitPorFileDB', '미사용'],
    ['LOGPATH', '없음', '로그 폴더 (Oracle 빌드)', '-', IOA + ':GsGetInitPorFileDB', '미사용'],
    ['FILENAME', '없음', '로그 파일 머리 (Oracle 빌드)', '-', IOA + ':GsGetInitPorFileDB', '미사용'],
])
para(d, '※ 현행 빌드 정의는 SQL 이라 [DB] 섹션은 읽지 않는다. 구 Oracle 판 호환으로 남겨 둔 섹션이다.', size=9, color=GRAY)

d.add_heading('4.2 [P_DB] DB 접속 (SQL Server)', 2)
keytable(d, 'IO', 'P_DB', [
    ['IP', '없음', 'SQL Server 주소 "서버,포트"', RST, IOA + ':GsGetInitPorFilePDB ← SYS_MAIN', '현장 : 서버IP,1433'],
    ['DATABASE', '없음', 'DB 명', RST, IOA + ':GsGetInitPorFilePDB', 'LGLS_MCS_IO'],
    ['PORT', '없음', '읽기만 하고 SQL 접속 문자열에는 쓰지 않음 (PostgreSQL 잔재)', '-', IOA + ':GsGetInitPorFilePDB', '미사용'],
    ['USER', '없음', 'SQL 로그인', RST, IOA + ':GsGetInitPorFilePDB', ''],
    ['USER_PW', '없음', 'SQL 암호', RST, IOA + ':GsGetInitPorFilePDB', ''],
    ['LOGPATH', '없음', '로그 폴더 (작업 폴더 기준 상대경로)', RST, IOA + ':GsGetInitPorFilePDB → cLogThread', ''],
    ['FILENAME', '없음', '로그 파일 머리', RST, IOA + ':GsGetInitPorFilePDB', 'IOTASK_'],
])

d.add_heading('4.3 [CNF] 스케줄 옵션', 2)
keytable(d, 'IO', 'CNF', [
    ['LAND_BY_CARGO_SEC', '1', '출고 착지를 트래킹 없이 물리 상태(화물 감지)로 인정하기까지 대기(초). 0 = 즉시', CALL, SCH + ':GsReadInitProfileCnf', ''],
    ['OUT_VIA_RGV_STATE', '0', '1 = 출고 RGV 구간도 상태로 표현 15→30→35→15→19 / 0 = 구 경로(메모리 큐, RecoverOutOrphans)', RST + ' (DB 연결 직후 1회)', SCH + ':GsReadInitProfileCnf', ''],
    ['OUT_HOLD_SEC', '3', 'S/C#1 통로 출고 유지 유예(초). 음수는 0. 0 = 유예 없음(입고→출고 재전환 2회 가능). CV2_DIR_MODE=1 에서는 미사용', CALL, SCH + ':OUT_HOLD_SEC', ''],
    ['IN_HS_STATUS', '15', '입고가 RGV 로 통로(H/S)에 내려진 뒤 작업상태. 15 = CV 구동중 / 16 = 통로CV 구동중(출고와 같은 표시). 16 이 아니면 15 로 봄', CALL, SCH + ':GsReadInitProfileCnf', ''],
    ['MANUAL_DIR_HOLD_SEC', '120', 'Client 더블클릭 수동 방향 전환 뒤 자동 방향 정합 억제 시간(초). 0 = 억제 없음. 수동 전환은 설비 통신이 화물 보류 없이 즉시 씀', CALL, SCH + ':GsReadInitProfileCnf', ''],
    ['CV2_DIR_MODE', '1', 'S/C#1 통로(C/V#2 103) 방향 전환 방식 0/1/2 (2.2 절)', CALL, SCH + ':CV2_DIR_MODE', ''],
    ['CV2_STALL_SEC', '60', '모드 1 의 교착 판정 무진행 시간(초). 10 미만은 10', CALL, SCH + ':CV2_STALL_SEC', ''],
    ['CV2_FAIR', '0', '모드 2 : 0 = 출고 우선 / 1 = 번갈아 쓰기(더 오래 기다린 쪽)', CALL, SCH + ':CV2_FAIR', ''],
    ['SC_DUAL_CODES', '73,74', '크레인 이중입고 코드(쉼표 구분). 73 좌측 렉 / 74 우측 렉. MarkErrorJobStatus 가 주석 처리되어 현재는 '
     'UpdateScData 의 재지정 지시 허용 판정(bAllowRedirectErr)에만 쓰임', CALL, SCH + ':GsReadInitProfileCnfStr', ''],
    ['SC_EMPTY_CODES', '75', '크레인 공출고 코드. 용도는 위와 같음', CALL, SCH + ':GsReadInitProfileCnfStr', ''],
])

d.add_heading('4.4 [VIEW]', 2)
keytable(d, 'IO', 'VIEW', [
    ['LOG_FILTER_BTN', 'Y (키 없음 · 빈 값)', 'Y/1/T = 메인 화면 [로그 필터] 버튼 표시 / N/0/F = 숨김', RST, r'FRM\SYS\SYS_MAIN.cs:GsReadInitProfileViewYN', ''],
])
para(d, '※ [VIEW] 는 파일의 마지막 섹션이다. 운전 키를 파일 끝에 덧붙이면 [VIEW] 로 들어가 무시된다(2.3 절).', size=9, color=RED)

# ════════════════════════════════════════════════════════════════════
d.add_heading('5. 상위 통신 - EcsComA.ini', 1)
HL = 'modCmLib.cs:ReadInitProfile'
d.add_heading('5.1 [WMS] · [Network] · [Property]', 2)
table(d, HEAD, [
    ['[WMS] INFO', cur('HOST', 'WMS', 'INFO'), '빈 값', 'WmsInfo.dll DB 정보 키. [DB] DATABASE 가 비었을 때만 이 키로 DB 정보를 가져옴', DBL, HL, '현행은 [DB] 사용'],
    ['[Network] LocalPort', cur('HOST', 'Network', 'LocalPort'), '없음 (빈 값이면 기동 오류)', '상위(WMS/IMS) 명령 수신 대기 포트', RST, HL + ' → modWorkThread TcpListener', '현장 8001'],
    ['[Network] RemoteIP', cur('HOST', 'Network', 'RemoteIP'), '없음', '보고 송신 대상 IP', DBL, HL + ' → CCliWork.ConnectSock', '현장 192.100.1.191'],
    ['[Network] RemotePort', cur('HOST', 'Network', 'RemotePort'), '없음 (빈 값이면 기동 오류)', '보고 송신 대상 포트', DBL, HL, '현장 8002'],
    ['[Property] RE_DRCT_CNT', cur('HOST', 'Property', 'RE_DRCT_CNT'), '0', '이중입고 재지정 허용 횟수(구 ECS). 읽어서 g_strRE_DRCT_CNT 에 두지만 사용처 없음', '-', HL, '미사용'],
], widths=W, font=8)
para(d, '※ 시뮬레이터 시험값(127.0.0.1 / 9911 / 9910)이 들어 있다. 현장 반영 시 RemoteIP=192.100.1.191, LocalPort=8001, '
        'RemotePort=8002 로 복원한다(주석 #RemoteIP=192.100.1.191 참고). "#" 줄은 주석이다.', size=9, color=RED)

d.add_heading('5.2 [DB] DB 접속', 2)
keytable(d, 'HOST', 'DB', [
    ['IP', '빈 값', 'SQL Server 주소 "서버,포트"', DBL, HL, '현장 : 서버IP,1433'],
    ['DATABASE', '빈 값 (비면 WmsInfo.dll)', 'DB 명', DBL, HL, 'LGLS_MCS_IO'],
    ['USER', 'STD_USER (빈 값일 때)', 'SQL 로그인', DBL, HL, ''],
    ['USER_PW', 'STD_USER (빈 값일 때)', 'SQL 암호', DBL, HL, ''],
])

d.add_heading('5.3 [Host] 상위 규약', 2)
keytable(d, 'HOST', 'Host', [
    ['StationMapMode', 'WMS', 'IMS 스테이션 코드 해석. WMS = 운영 DB(MC_STA_MST AA01) 기준 : 101=TR#22, 102=출발 TR#30/도착 TR#29, '
     '103=TR#24(제품입고대), 104=TR#26(원부자재불출대), 105/110=TR#3(C/V#2) / ECS = 구 ECS 구축시드(TB_CODEMASTER 022) : '
     '101=TR#22 102=TR#24 103=TR#26 104=TR#29 105=TR#30. 하이랙 001~005/201~205 → 901~905 공통', DBL, 'modStationMap.cs:LoadFromIni', 'WMS'],
    ['RE_REPORT_09_SEC', '60', '완료(F) 보고 후 응답 없는 09 잔류 작업 재보고 간격(초). 처음 09 는 즉시 보고. 5~3600 으로 자름', CYC + ' (5초)',
     'CCliWork.cs:RE_REPORT_09_SEC', ''],
    ['ScErrCodeType', 'SC_LGLS', '크레인 에러코드 마스터 구분(EQP_ECD_MST.EQP_TYP). SC_LGLS = 알람 리스트 260917 / SC_SFA = 종전', DBL, HL, ''],
    ['ScDualCodes', '73,74', 'E 전문 ErrorKind=1(이중입고) 판정 코드. SC_SFA 로 되돌릴 때 54,55', DBL, HL + ' → CCliWork', ''],
    ['ScEmptyCodes', '75', 'ErrorKind=3(공출고) 판정 코드. SC_SFA 는 58,59', DBL, HL, ''],
    ['ScInFailCodes', '빈 값', 'ErrorKind=2(입고장애) 판정 코드. SC_SFA 는 56', DBL, HL, '새 코드표에는 해당 없음'],
    ['ScOutFailCodes', '빈 값', 'ErrorKind=4(출고장애) 판정 코드. SC_SFA 는 57', DBL, HL, '새 코드표에는 해당 없음'],
])
para(d, '※ 코드에는 [Host] WarehouseDefine(A/B, 없으면 기본 유지)도 있으나 INI 에 두지 않는다.', size=9, color=GRAY)

d.add_heading('5.4 [StationMap_ECS] · [StationMap_WMS] 개별 재정의', 2)
para(d, 'StationMapMode 에 맞는 섹션 하나만 읽는다. 형식은 "IMS코드 = 출발내부, 도착내부" 이고, "REPORT_MAIN_nnn = 내부코드" 로 '
        '상태보고 대표 설비를 바꿀 수 있다. 현재 두 섹션의 항목은 모두 ";" 주석(재정의 없음, 코드 기본 매핑 사용)이다. '
        '읽는 코드 : modStationMap.cs:LoadOverride / LoadReportMainOverride. 반영 : ' + DBL + '.')

# ════════════════════════════════════════════════════════════════════
d.add_heading('6. 운전 화면 - Ecs.ini', 1)
CF = 'Config.cpp'
d.add_heading('6.1 [WH] · [DB_2] · [COMMON]', 2)
table(d, HEAD, [
    ['[WH] WH_TYP', cur('CLI', 'WH', 'WH_TYP'), '10', '창고 구분', RST, CF + ':LoadConfigWH, MainFrm.cpp', ''],
    ['[DB_2] DRIVER', cur('CLI', 'DB_2', 'DRIVER'), 'PostgreSQL Unicode', 'ODBC 드라이버 이름', RST, CF + ':LoadConfigDATABASE → CAdoDB', 'SQL Server'],
    ['[DB_2] SERVER', cur('CLI', 'DB_2', 'SERVER'), 'LOCALHOST', 'SQL Server 주소 "서버,포트"', RST, CF + ':LoadConfigDATABASE', '현장 localhost,1433 (서버 PC) 또는 서버IP,1433'],
    ['[DB_2] TRUSTED', cur('CLI', 'DB_2', 'TRUSTED'), '0', '1 = Windows 인증(USERID/USERPASSWORD 무시) / 0 = SQL 로그인', RST, CF + ':LoadConfigDATABASE', ''],
    ['[DB_2] USERID', cur('CLI', 'DB_2', 'USERID'), 'LFC', 'SQL 로그인', RST, CF + ':LoadConfigDATABASE', ''],
    ['[DB_2] USERPASSWORD', cur('CLI', 'DB_2', 'USERPASSWORD'), 'LFC', 'SQL 암호', RST, CF + ':LoadConfigDATABASE', ''],
    ['[DB_2] DATABASE', cur('CLI', 'DB_2', 'DATABASE'), 'LFC', 'DB 명', RST, CF + ':LoadConfigDATABASE', 'LGLS_MCS_IO'],
    ['[COMMON] TabCount', cur('CLI', 'COMMON', 'TabCount'), '1', '메인 화면 탭 수', RST, 'EcsDoc.cpp', ''],
], widths=W, font=8)
para(d, '※ INI 주석대로 현장에 가져갈 때는 SERVER=127.0.0.1,1435(로컬 시험값)를 현장 서버 주소(,1433)로 되돌린다.', size=9, color=RED)

d.add_heading('6.2 [USER] 로그인 · 언어 · 색', 2)
keytable(d, 'CLI', 'USER', [
    ['DEFAULT_ID', 'LGLS', '기동 시 자동 로그인 계정. 비우면 로그인 창을 그대로 둠', RST, 'EcsDoc.cpp:GetDefaultUserId', ''],
    ['DEFAULT_PW', 'LGLS', '자동 로그인 암호', RST, 'EcsDoc.cpp:GetDefaultUserPw', ''],
    ['VIEW_ONLY_ID', 'LFC', '읽기전용 계정(로그인한 것으로 보지 않음). 안 쓰면 비움', CALL, 'EcsDoc.cpp:IsViewOnlyId', ''],
    ['LAST_LANG', '1', '마지막 언어 0 = KOR / 1 = ENG / 2 = HUN / 3 = CHIN (EN_LANG). 종료 시 프로그램이 저장', RST, CF + ':LoadConfigUSER / SaveConfigUSER', '0'],
    ['LAST_TAB_INDEX', '1', '마지막 탭. 종료 시 저장', RST, CF + ':LoadConfigUSER', ''],
    ['JOB_STALL_WARN_SEC', '300', '완료가 아닌 작업이 이 시간(초) 넘게 갱신되지 않으면 체류 경고창(CWarningDlg). 0 이하는 10초로 보정', HOT, 'WarningDlg.cpp:ReloadIni', ''],
    ['USER_COLOR_STO / RET / MOVE / RTR', 'GREEN / BLUE / CYAN / MAGENTA',
     '입고 · 출고 · 이동 · 랙간 작업색 (16진 6자리, BGR 순서 - 예 0000ff = 빨강)', RST + ' (핫 리로드 제외)', CF + ':LoadConfigUSER', ''],
    ['USER_COLOR_WTW / ATA', 'SITE / MAGENTA', 'WTW 는 사용 안 함. ATA 는 읽는 줄이 WTW 변수에 대입되는 결함이 있어 ini 값이 반영되지 않음(부록 A)', RST, CF + ':LoadConfigUSER', ''],
    ['USER_COLOR_STN_STO / STN_RET', 'GREEN / BLUE', '작업대 입고/출고 색', RST, CF + ':LoadConfigUSER', ''],
    ['USER_COLOR_HS_STO / HS_RET', '오렌지 / 진초록', '핸드셰이크 색 (싣는 곳 / 내려놓는 곳)', RST, CF + ':LoadConfigUSER', ''],
    ['USER_COLOR_SUSPEND · STO_SUSPEND · RET_SUSPEND · STORET_SUSPEND', 'RED · 주황빨강 · 자홍빨강 · 진빨강', '정지 · 입고 금지 · 출고 금지 · 입출고 정지 레일색', RST, CF + ':LoadConfigUSER', ''],
    ['USER_COLOR_RAIL_ERROR', '보라', '레일 설비 에러색', RST, CF + ':LoadConfigUSER', ''],
    ['USER_COLOR_MANUAL · DISCONNECT · ERROR', 'YELLOW · WHITE · RED', '수동 · 통신 끊김 · 에러색', RST, CF + ':LoadConfigUSER', ''],
    ['USER_COLOR_WC_PASS · TRACKING · CV_SEARCH · SC_INVK', 'SITE · 황토 · PINK · 파랑',
     '통과 · 작업번호만 남은 트랙(범례 20) · C/V 찾기 · 크레인 작업중 레일색', RST, CF + ':LoadConfigUSER', ''],
    ['USER_COLOR_SEMI_STO / RET / MOVE / RTR / ATA / PR', '자동색의 진한 색', '반자동 작업색', RST, CF + ':LoadConfigUSER', ''],
])
para(d, '※ USER_COLOR_* 의 현재값은 INI 원문을 따른다(27개). 사용자 지시로 핫 리로드 대상에서 뺐다.', size=9, color=GRAY)
table(d, ['키', '현재값'], [[k[1], v] for k, v in CUR['CLI'].items() if k[0] == 'USER' and k[1].startswith('USER_COLOR_')],
      widths=[8.0, 4.0], font=8)

d.add_heading('6.3 [SIREN] · [ETC]', 2)
table(d, HEAD, [
    ['[SIREN] SOUND', cur('CLI', 'SIREN', 'SOUND'), '0', '사이렌 소리 파일. LoadConfigSound 호출이 주석 처리되어 현재 읽지 않음', '-', CF + ':LoadConfigSound (미호출)', '미사용'],
    ['[ETC] COLLECT_INTERVAL_MS', cur('CLI', 'ETC', 'COLLECT_INTERVAL_MS'), '300', '설비 데이터 DB 수집 주기(ms). 100~5000 으로 자름. 키우면 화면이 느려지고 줄이면 조회가 잦아짐', CYC, 'CollectDB.cpp (수집 스레드 루프)', ''],
    ['[ETC] ViewRetCnt', cur('CLI', 'ETC', 'ViewRetCnt'), '0', '크레인별 출고 잔여 수 표시 1/0', HOT, CF + ':LoadConfigUSER, EcsView.cpp:ReloadIniHot', ''],
], widths=W, font=8)

d.add_heading('6.4 [MENU] 메뉴 · 화면 배치 · 표시 규칙', 2)
MF = 'MainFrm.cpp'
keytable(d, 'CLI', 'MENU', [
    ['STATUS_POS', 'RIBBON (배치) / TOP (다른 곳)', '통신상태(EQUIP/HOST/SCH) 위치. RIBBON = 리본 위쪽 / BOTTOM = 하단 상태바', RST, MF + ' (1215, 1477 행)', ''],
    ['COMM_MODE', '1', 'STATUS_POS=RIBBON 일 때 모양. 1 = 탭 줄 오른쪽 한 줄 / 2 = 탭 줄 "통신" 이름표 + 리본 오른쪽 "통신" 그룹', RST, MF + ':GetCommMode', ''],
    ['SCWAIT_VIEW', '1', '크레인별 출고 대기 건수 표시 1/0. 0 이면 조회도 하지 않음', CYC + ' (1초)', 'ScPair.cpp:RefreshScWaitCount', ''],
    ['USER_MENU', '1', 'ECS 리본 [사용자] 메뉴 표시 1/0', RST, MF, ''],
    ['SEMITEST_MENU', '1', 'MANUAL 탭 [반자동 TEST] 그룹 표시 1/0', RST, MF, ''],
    ['INI_MENU', '1', '리본 환경설정 [INI 열기] 버튼 표시 1/0', RST, MF, '설계상 0 (숨김)'],
    ['UIMODE_MENU', '1', '리본 [판넬 보기] 그룹 표시 1/0 (PANEVIEW_MENU 와 함께 판정)', RST, MF, ''],
    ['PANEVIEW_MENU', 'MAIN_UI=2 이면 0, 아니면 1', '[판넬 보기] 그룹 강제 지정 1/0. 파일에는 주석(;)으로만 있음', RST, MF, ''],
    ['PROCESS_MENU', '1', '리본 [처리] 그룹 표시 1/0 (그룹 안 버튼은 각자의 키)', RST, MF, ''],
    ['PRODINFO_MENU', '1', '창고 모니터링 [제품정보] 버튼 표시 1/0', RST, MF, ''],
    ['AUTOTIME_MENU', '1', '[처리] 그룹 [시간 기반 자동 처리] 버튼 표시 1/0 (선택 여부는 DB COMMON_CODE SCH_OPT/AUTO_TIME)', RST, MF, ''],
    ['ZOOM_BTN', '1', 'CV/SC/RTV 상태창 [확대] 버튼 표시 1/0 (숨기면 펼친 패널도 접음)', HOT, 'EcsView.cpp:ReloadIniHot, CvSkinDlg.cpp', '현장 점검 [쓰기] 절차 때 1'],
    ['MAIN_UI', '1', '메인 화면 배치. 1 = 종전(범례 큰 표) / 2 = 왼쪽에 통신 상태 · 축소 범례 · 작업정보 고정 (범례 칸 뺀 EcsLayout1.xml 필요)', RST, 'EcsView.cpp:CreateMainUi2, ' + MF, ''],
    ['MAIN_UI_LEFT_W', '0 (기본 크기)', '왼쪽 칸 폭(px). 손잡이를 끌면 프로그램이 저장', RST, 'EcsView.cpp:CreateMainUi2', ''],
    ['MAIN_UI_COMM_H', '0', '통신 칸 높이(px), 자동 저장', RST, 'EcsView.cpp', ''],
    ['MAIN_UI_LEG_H', '0', '범례 칸 높이(px), 자동 저장', RST, 'EcsView.cpp', ''],
    ['MAIN_UI_VEH_H', '0', '설비반송 칸 높이(px), 자동 저장', RST, 'EcsView.cpp', ''],
    ['UI_TRACE', '1', '크레인/RGV 대화상자 클릭 경로 단계를 LOG\\ECS_UI_yyyymmdd.log 에 기록 1/0', HOT, 'Lib.cpp:IniUiTrace', ''],
    ['LOADBIT_GATE', '1', '크레인/RGV 색 · 번호를 차상 화물 비트로 한 번 더 거름 1 = 거름(구 ECS 판정) / 0 = 끔(현장 PLC 가 비트를 안 채우면)', HOT, 'Lib.cpp:IniLoadBitGate', ''],
    ['VEH_CLEAR_MODE', '3 (0~3 밖이면 3)', '하역 뒤 색 · 작업번호 지우는 시점. 0 = 설비 완료 보고까지(약 10초 잔상) / 1 = 차상 적재 비트 꺼짐 / '
     '2 = 화물번호가 C/V 트랙에 기록 / 3 = 표시는 0 과 같게, 지움은 작업 시작 때 없던 새 트랙에 기록될 때', HOT, 'Lib.cpp:IniVehClearMode', ''],
    ['DIR_ADDR_CV11', '310', '메인 화면 "입고/출고 모드" 칸 더블클릭 확인창에 보일 C/V#11 방향 D 워드 주소(10진). 주소맵과 같은 번호', CALL, 'EcsView.cpp (더블클릭 처리)', ''],
    ['DIR_ADDR_CV2', '301', '같은 확인창의 C/V#2 방향 D 워드 주소', CALL, 'EcsView.cpp', ''],
])
para(d, '※ INI_MENU 는 현재 %s 이다. 설계상 운영 PC 에서는 0(INI 열기 버튼 숨김)으로 둔다. '
        'ZOOM_BTN 은 현재 %s 이며, INI 주석("현장 점검 절차에 필요해 켠다")과 값이 다르다 - 점검 때만 1 로 바꾼다(저장 즉시 반영).'
     % (cur('CLI', 'MENU', 'INI_MENU'), cur('CLI', 'MENU', 'ZOOM_BTN')), size=9, color=RED)

d.add_heading('6.5 [Title] · [RibbonMenu] · [SC_ERR]', 2)
table(d, HEAD, [
    ['[Title] BuildDate', cur('CLI', 'Title', 'BuildDate'), '1', '제목줄에 빌드(링크) 시각 표시 1/0', HOT, MF + ':ReloadTitleAndTipIni', ''],
    ['[Title] DbInfo', cur('CLI', 'Title', 'DbInfo'), '1', '제목줄에 접속 DB(DATABASE@SERVER) 표시 1/0', HOT, MF + ':ReloadTitleAndTipIni', ''],
    ['[Title] Path', cur('CLI', 'Title', 'Path'), '1', '제목줄에 실행 파일 전체 경로 표시 1/0. 창 폭보다 길면 흘러감(마퀴)', HOT, MF + ':ReloadTitleAndTipIni', ''],
    ['[RibbonMenu] ToolTip', cur('CLI', 'RibbonMenu', 'ToolTip'), '1', '1 = 리본 툴팁에 실제 리소스 경로(그룹 = 문구 ini, 버튼 = 아이콘 png, 없으면 "(파일 없음)") / 0 = 원래 툴팁', HOT, MF + ':ApplyRibbonToolTipIni', '운영 0 검토'],
    ['[SC_ERR] ERR_TYP', cur('CLI', 'SC_ERR', 'ERR_TYP'), 'SC_LGLS', '상태창 알람 문구용 EQP_ECD_MST 코드표 구분', CALL, 'Lib.cpp:ScErrTyp', ''],
    ['[SC_ERR] DUAL_CODES', cur('CLI', 'SC_ERR', 'DUAL_CODES'), '73,74', '이중입고 코드(에러해제 확인창). 73 좌측 렉 / 74 우측 렉', CALL, 'Lib.cpp:IsScDualErr', ''],
    ['[SC_ERR] EMPTY_CODES', cur('CLI', 'SC_ERR', 'EMPTY_CODES'), '75', '공출고 코드(에러해제 · 삭제 확인창)', CALL, 'Lib.cpp:IsScEmptyErr', ''],
], widths=W, font=8)

# ════════════════════════════════════════════════════════════════════
d.add_heading('7. 현장 반영 점검표', 1)
table(d, ['No', '파일', '확인 항목', '현장값'], [
    ['1', 'WCS_DB.INI', '[COMM0] IP', '192.100.1.110'],
    ['2', 'WCS_DB.INI · ENV_IOSCH.INI · EcsComA.ini', '[DB_2] / [P_DB] / [DB] IP', '서버IP,1433'],
    ['3', 'Ecs.ini', '[DB_2] SERVER', 'localhost,1433 또는 서버IP,1433'],
    ['4', 'EcsComA.ini', '[Network] RemoteIP / LocalPort / RemotePort', '192.100.1.191 / 8001 / 8002'],
    ['5', 'WCS_DB.INI', '[CNF] SC_ERR_CODE_BLOCK', '3 (권장)'],
    ['6', 'WCS_DB.INI', '[PLC] DIR_CODE', 'PLC 래더 확인값 (담당자 확인 IN1). 시뮬 시험 뒤 되돌림 주의'],
    ['7', 'ENV_IOSCH.INI', 'CV2_DIR_MODE · CV2_STALL_SEC · CV2_FAIR 가 [CNF] 안에 있는지', '[VIEW] 아래면 옮김'],
    ['8', 'Ecs.ini', '[MENU] INI_MENU / ZOOM_BTN / SEMITEST_MENU / USER_MENU', '0 / 점검 때만 1 / 0 / 0'],
    ['9', 'ALL_TASK\\bin\\Debug', '세 TASK INI 가 개별 실행본과 같은지', '동일'],
], widths=[1.0, 7.0, 10.5, 7.2], font=9)

# ════════════════════════════════════════════════════════════════════
d.add_heading('부록 A. 문서화 중 확인한 INI · 코드 불일치', 1)
para(d, '아래 항목은 이번 문서 작성 중 발견했으며 코드와 INI 는 고치지 않았다(기록만).')
table(d, ['No', '위치', '내용', '영향'], [
    ['1', 'ENV_IOSCH.INI [CNF] IN_HS_STATUS', '코드 기본값 15, 현재값 16', '의도된 운용값. 키가 빠지면 15 로 돌아감'],
    ['2', 'Ecs.ini [MENU] COMM_MODE', 'INI 주석은 "2 (기본)" 이나 코드 기본값은 1 (MainFrm.cpp GetCommMode)', '키가 빠지면 모양이 1 로 바뀜'],
    ['3', 'MainFrm.cpp STATUS_POS', '같은 키의 기본값이 두 곳에서 다름 (1215 행 RIBBON, 1477 행 TOP). 주석(1144 행)도 "TOP 기본"', '키가 있으면 영향 없음'],
    ['4', 'Ecs.ini [MENU] ZOOM_BTN', '주석은 "현장 점검 절차에 필요해 켠다" 인데 값은 0. 코드 기본값은 1', '주석 정리 필요'],
    ['5', 'Ecs.ini [ETC] ViewRetCnt', '코드 기본값 0, 현재값 1', '키가 빠지면 숨김'],
    ['6', 'Config.cpp LoadConfigUSER', 'USER_COLOR_ATA 를 읽어 m_clrUSER_COLOR_WTW 에 대입 - ATA 는 ini 값이 반영되지 않고 WTW 를 덮음', '현재 두 값이 같아(ff00ff) 화면 차이 없음'],
    ['7', 'Ecs.ini [SIREN] SOUND', 'LoadConfigSound 호출이 주석 처리되어 키를 읽지 않음', '정리 대상 후보'],
    ['8', 'Config.cpp SaveConfigWH', '저장 키가 "WHTYP" (읽는 키는 WH_TYP)', 'SaveConfigAll 경로로 저장하면 엉뚱한 키가 생김'],
    ['9', 'WCS_DB.INI [PLC] DIR_CODE', 'INI 주석은 "재기동 없이 다음 지시부터" 이나 INI 직접 수정은 기동 시에만 읽음. 즉시 반영은 화면 라디오로 바꿀 때만', 'INI 를 고친 뒤에는 재기동 필요'],
    ['10', 'WCS_DB.INI [DB_2] PORT, [CNF] USERID / ENV_IOSCH.INI [P_DB] PORT, [DB] 섹션', 'SQL 빌드에서 읽기만 하거나 전혀 읽지 않음', '미사용 키'],
    ['11', 'EcsComA.ini [Property] RE_DRCT_CNT', '읽어서 전역에 두지만 사용처 없음. 또 숫자 검사를 읽은 값이 아니라 기존 전역값으로 해 첫 호출에는 0 이 됨', '미사용 키'],
    ['12', 'WCS_TASK_HOST 경로', 'MAIN_INI 는 작업 폴더 기준, RE_REPORT_09_SEC 는 실행 파일 폴더 기준으로 EcsComA.ini 를 찾음', '바로가기 작업 폴더가 다르면 서로 다른 파일을 읽을 수 있음'],
    ['13', 'VehThread.cs 822 행 주석', '"SC_ERR_CODE_BLOCK=1 이면 종전대로 블록을 읽는다" - 09-16 모드 재정의 전 설명이 남아 있음', '주석만'],
    ['14', 'WCS_DB.INI [DB_2] 머리 주석', '"--PostgreSql" 표기. 실제는 SQL Server', '주석만'],
    ['15', 'WCS_DB.INI [CNF] SC_ERR_CODE_BLOCK', '코드 기본값 0, 현재값 0 - 현장 권장값 3 과 다름', '현장 반영 때 3 으로'],
], widths=[1.0, 6.2, 12.0, 6.5], font=8)

out = os.path.join(OUT, '12_INI_설정_정의서.docx')
d.save(out)
print('saved', out)
