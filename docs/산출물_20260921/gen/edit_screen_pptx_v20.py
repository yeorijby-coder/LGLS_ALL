# -*- coding: utf-8 -*-
"""06_ECS_화면설계서.pptx (사용자 수정본) 제자리 수정 - V2.0 (2026-09-21 최종판).
   사용자가 손본 09-17 본을 그대로 두고 글자 조각·그림만 바꾼다. 새로 만들지 않는다.
   캡처는 2026-09-21 Ecs 를 전체화면으로 띄워 새로 찍은 것(shots\\*_v20.png, shots\\ribbon_*.png).
   사용: python edit_screen_pptx_v20.py [빌드문자열]"""
import os, sys, copy
from pptx import Presentation

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PPT = os.path.join(BASE, '06_ECS_화면설계서.pptx')
SHOTS = os.path.join(BASE, 'shots')
BUILD = sys.argv[1] if len(sys.argv) > 1 else ''

prs = Presentation(PPT)


def replace_in_runs(shape, old, new):
    n = 0
    if not shape.has_text_frame:
        return 0
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new); n += 1
    return n


def add_para(shape, text):
    """마지막 문단의 서식을 복사해 문단을 하나 덧붙인다."""
    tf = shape.text_frame
    last = tf.paragraphs[-1]._p
    newp = copy.deepcopy(last)
    last.addnext(newp)
    p = tf.paragraphs[-1]
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    p.runs[0].text = text


def swap_picture(slide, old_shape, png):
    """그림 하나를 같은 자리·같은 폭으로 바꾼다(높이는 비율 유지, 원래 높이를 넘으면 높이에 맞춤)."""
    if not os.path.exists(png):
        print('  그림 없음', os.path.basename(png)); return False
    from PIL import Image
    w, h = Image.open(png).size
    L, T, W, H = old_shape.left, old_shape.top, old_shape.width, old_shape.height
    nw, nh = W, int(W * h / w)
    if nh > H:
        nh, nw = H, int(H * w / h)
    slide.shapes._spTree.remove(old_shape._element)
    slide.shapes.add_picture(png, L + (W - nw) // 2, T, nw, nh)
    print('  그림 교체', os.path.basename(png))
    return True


def pics(slide):
    return [s for s in slide.shapes if s.shape_type == 13]


def replace_biggest(slide, png):
    ps = pics(slide)
    if not ps:
        return False
    return swap_picture(slide, max(ps, key=lambda s: s.width * s.height), png)


def body(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and '메뉴 Path' in sh.text_frame.text:
            return sh


# 1) 모든 장 : 버전 표기
cnt = 0
for s in prs.slides:
    for sh in s.shapes:
        cnt += replace_in_runs(sh, 'Ver 1.2', 'Ver 2.0')
print('Ver 표기', cnt)

# 2) 표지
for sh in prs.slides[0].shapes:
    replace_in_runs(sh, '2026-09-17', '2026-09-21')
    if BUILD and sh.has_text_frame and 'Build 2026.09.17 08:18:34' in sh.text_frame.text:
        replace_in_runs(sh, 'Build 2026.09.17 08:18:34', 'Build ' + BUILD)

# 3) 이력 표에 V2.0 행
for sh in prs.slides[1].shapes:
    if sh.has_table:
        tbl = sh.table
        if tbl.rows[len(tbl.rows) - 1].cells[0].text != 'V2.0':
            tr = copy.deepcopy(tbl._tbl.tr_lst[-1])
            tbl._tbl.append(tr)
            vals = ['V2.0', '2026-09-21', 'LGLS ECS Renewal', '',
                    '최종판 : 크레인·RGV 상태창 정리(동작상태 파생 항목 제거, 값 칸 통합·확대), '
                    '설비 DOWN·수동을 에러색으로, 에러 칸 "[코드]문구" 표시, 상태창 긴 글자 흐름 표시, '
                    'MANUAL 수동지시 숨김 반영. 캡처 전면 갱신(전체화면)']
            row = tbl.rows[len(tbl.rows) - 1]
            for c, v in zip(row.cells, vals):
                tf = c.text_frame
                for p in tf.paragraphs[1:]:
                    p._p.getparent().remove(p._p)
                rs = tf.paragraphs[0].runs
                if rs:
                    rs[0].text = v
                    for r in rs[1:]:
                        r._r.getparent().remove(r._r)
                else:
                    tf.paragraphs[0].text = v
            tbl.rows[len(tbl.rows) - 1].height = tbl.rows[len(tbl.rows) - 2].height

# 4) 캡처 교체 ─ 메인 / CV·SC·RTV 상태
print('슬라이드 3 (메인 전체화면)');   replace_biggest(prs.slides[3], os.path.join(SHOTS, 'main_full_v20.png'))
print('슬라이드 24 (CV 상태)');       replace_biggest(prs.slides[24], os.path.join(SHOTS, 'cv_0_v20.png'))
print('슬라이드 25 (SC 상태)');       replace_biggest(prs.slides[25], os.path.join(SHOTS, 'sc_0_v20.png'))
print('슬라이드 26 (RTV 상태)');      replace_biggest(prs.slides[26], os.path.join(SHOTS, 'rtv_0_v20.png'))

# 5) 리본 4장 (슬라이드 5) : 왼쪽 3장은 top 순(ECS/MANUAL/LOG), 오른쪽 1장은 통신
print('슬라이드 5 (리본)')
sl5 = prs.slides[5]
ps = pics(sl5)
left = sorted([p for p in ps if p.left < 8000000], key=lambda p: p.top)
right = [p for p in ps if p.left >= 8000000]
for sh, name in zip(left, ('ribbon_ecs.png', 'ribbon_manual.png', 'ribbon_log.png')):
    swap_picture(sl5, sh, os.path.join(SHOTS, name))
for sh in right:
    swap_picture(sl5, sh, os.path.join(SHOTS, 'ribbon_comm.png'))

# 6) 특기사항 ─ 설비에러이력 / CV / SC / RTV
b = body(prs.slides[20])
if b is not None and '코드표' not in b.text_frame.text:
    add_para(b, '')
    add_para(b, '■ 특기사항')
    add_para(b, '· 에러 문구는 에러코드 마스터(EQP_ECD_MST)에서 설비·코드표 구분으로 찾는다. 크레인은 PLC 알람 리스트 코드표(SC_LGLS).')
    add_para(b, '· 설비구분 SC 를 고르면 SC_LGLS 등 크레인 코드표 구분으로 쌓인 이력도 함께 조회된다.')

b = body(prs.slides[24])
if b is not None and '흐름' not in b.text_frame.text:
    add_para(b, '· 칸에 다 들어가지 않는 긴 글자는 그 칸에서 한 글자씩 흘러간다(입력 중인 칸은 제외). 2026-09-21 추가.')

b = body(prs.slides[25])
if b is not None and '알람 문구' not in b.text_frame.text:
    add_para(b, '· 에러 칸은 "[코드]알람 문구" 로 표시한다 (예 [0073]좌측 렉 이중입고). 코드표 = Ecs.ini [SC_ERR] ERR_TYP(SC_LGLS).')
    add_para(b, '· 에러 해제 시 이중입고(73·74)·공출고(75) 확인창은 Ecs.ini [SC_ERR] DUAL_CODES / EMPTY_CODES 기준.')
    add_para(b, '· 2026-09-21 : 지상반·기상반·SRC상태 칸을 없앴다(셋 다 동작상태 워드에서 계산한 값이라 동작상태와 늘 같았다).')
    add_para(b, '· 화물유무·포크위치·수평주행·완료상태를 동작상태와 한 그룹으로 모으고 값 칸을 넓혔다.')

b = body(prs.slides[26])
if b is not None and '알람 문구' not in b.text_frame.text:
    add_para(b, '· 진단 칸은 "[코드]알람 문구" 로 표시한다 (예 [0033]RGV 전진 한계점, PLC 알람 비트 M5601~M560F).')
    add_para(b, '· 2026-09-21 : 운영모드·ACTIVE 칸을 없애고(동작상태 파생값), [RTV 금지] 영역에 맞춰 배치를 정리했다.')

# 7) 범례/색 설명 장에 DOWN·수동 에러색 추가 (슬라이드 6)
b = None
for sh in prs.slides[6].shapes:
    if sh.has_text_frame and '색' in sh.text_frame.text and len(sh.text_frame.paragraphs) > 2:
        b = sh; break
if b is not None and 'DOWN' not in b.text_frame.text:
    add_para(b, '· 2026-09-21 : 설비 동작상태가 DOWN 이거나 기상반이 수동이면 크레인·RGV 를 에러색으로 그린다.')

prs.save(PPT)
print('saved', PPT, len(prs.slides), '장')
