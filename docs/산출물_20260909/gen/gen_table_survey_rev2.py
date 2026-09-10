# -*- coding: utf-8 -*-
# [LGLS 2026-09-10] 테이블 사용 현황 조사표 Rev2 (밸리데이션 변경관리 첨부용)
#   Rev1 = 조사 결과(삭제 권고 7 / 판단 필요 5 / 유지 25)
#   Rev2 = 판단 필요 5개도 폐지하기로 결정하여 12개를 실제로 삭제한 뒤의 기록
#   판정 기준 : "빌드되는 프로그램의 소스에 그 테이블을 대상으로 한 SQL 이 있는가"
#   행 수는 근거로 쓰지 않는다(eqp_err_his 처럼 아직 안 쌓였을 뿐인 표가 있다).
import os, io, sys, datetime
sys.path.insert(0, r'D:\project\LGLS\Renewal\docs\산출물_20260909\gen')
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

OUT = r'D:\project\LGLS\Renewal\docs\산출물_20260909\테이블_사용현황_조사표_Rev2.docx'

ROWS = {}
try:
    for ln in io.open(r'C:\Users\USER\AppData\Local\Temp\rowcounts.txt', encoding='utf-8-sig'):
        ln = ln.strip()
        if '|' in ln:
            a, b = ln.split('|')
            ROWS[a.strip().lower()] = b.strip()
except IOError:
    pass
# 지운 표의 "삭제 직전" 행 수 (10_점검_표삭제전.sql 실행 결과, 2026-09-10)
ROWS.update({'auto_sc_work': '0', 'bcr_mst': '0', 'client_pgr_nm': '0', 'dest_def': '154',
             'event_log': '0', 'host_empty_plt': '0', 'mes_if_log': '0', 'rtv_data': '1',
             'sc_data': '5', 'wait_track': '0', 'wc_data': '0', 'wc_his': '0'})

# (테이블, 판정, 쓰는 곳/근거, 비고)
KEEP = ' 유지'
DROP = '삭제 완료'
HOLD = '삭제 완료'   # Rev1 에서는 '판단 필요' 였다

DATA = [
 ('job_mst',        KEEP, 'IO_TASK, WCS_TASK_HOST, Client', '진행 중 작업. 시스템의 중심 표'),
 ('job_mst_his',    KEEP, 'DB 트리거 trg_JOB_MST_StatusHis, WCS_TASK_HOST', '작업 상태 이력'),
 ('cv_data',        KEEP, 'WCS_TASK_CV, IO_TASK, Client', '컨베이어 트랙 상태/지시'),
 ('SC_DATA_LGLS',   KEEP, 'IO_TASK, Client', '크레인 상태/지시 (현행)'),
 ('RTV_DATA_LGLS',  KEEP, 'IO_TASK, Client', 'RGV 상태/지시 (현행)'),
 ('eqp_mst',        KEEP, '전 TASK 하트비트, Client 상태바', '설비/프로그램 마스터'),
 ('eqp_ecd_mst',    KEEP, 'Client 설비 대화상자', '설비 에러코드 명칭'),
 ('eqp_err_his',    KEEP, 'WCS_TASK_CV 기록, Client 조회', '0행은 아직 에러가 없었을 뿐이다'),
 ('common_code',    KEEP, '전 프로그램', '공통 코드'),
 ('cell_mst',       KEEP, 'Client, WCS_TASK_HOST', '랙 셀 마스터'),
 ('cell_dtl',       KEEP, 'Client CollectDataList (CELL_MST 와 LEFT JOIN)', '0행이나 조회 대상'),
 ('cv_def_inf',     KEEP, 'Client CvSkinDlg / PanelInfoDlg', '컨베이어 연결 정의'),
 ('cv_flow_def',    KEEP, 'WCS_TASK_CV', '흐름 정의'),
 ('sc_def_inf',     KEEP, 'Client ManualJob (CELL_MST 대체 조회)', '크레인 담당 구역'),
 ('sc_hs_def',      KEEP, 'Client Lib / ScManualRet / ScSkinDlg', '크레인 H/S 정의'),
 ('dest_pos_def',   KEEP, 'IO_TASK, Client', '도착지 정의'),
 ('dead_lock_zone_def', KEEP, 'IO_TASK cDefApi', '교착 구간 정의'),
 ('del_his_setting', KEEP, 'Client 로그 삭제 설정', '이력 보관 주기'),
 ('user_mst',       KEEP, 'Client 로그인', '사용자'),
 ('user_grp',       KEEP, 'Client 권한', '사용자 그룹'),
 ('user_grp_win',   KEEP, 'Client 권한(EcsDoc)', '그룹별 화면 권한'),
 ('chg_lang',       KEEP, 'Client Lang / Lib', '다국어 문구'),
 ('host_if_log',    KEEP, 'WCS_TASK_HOST 기록, Client 조회', '상위 전문 이력'),
 ('wcs_log_pgr',    KEEP, '전 TASK 기록, Client 조회', '프로그램 운전 이력'),
 ('wcs_client_log', KEEP, 'Client 기록/조회', '화면 조작 이력'),
 ('mes_if_log',     HOLD, 'Client [MES_LOG] 조회 화면만', 'MES 연동 없음. [MES_LOG] 조회 화면과 함께 폐지 (Rev1 판단 필요 -> 폐지 결정)'),
 ('wc_his',         HOLD, 'Client [WC 로그] 조회 화면만', 'WC(검량) 설비 없음. [WC 로그] 화면과 함께 폐지'),
 ('wc_data',        HOLD, 'Client WcSkinDlg/Wc, WCS_TASK_HOST 조회', 'WC 설비 없음. Client WC 화면·설비클래스, WCS_TASK_HOST 무게보고까지 함께 폐지'),
 ('bcr_mst',        HOLD, 'Client BCRSkinDlg/Bcr', 'BCR 설비 없음. Client BCR 화면·설비클래스와 함께 폐지'),
 ('host_empty_plt', HOLD, 'Client 공PLT 화면, WCS_TASK_HOST 조회', '공PLT 기능 폐지. 메뉴 항목과 Ecs.ini [MENU] 키까지 걷어냈다. WCS_TASK_HOST 빈파렛트 요청도 제거'),
 ('sc_data',        DROP, '없음 (C++ 클래스 파일명 SC_DATA.h 와 이름만 겹친다)', '구 ECS 잔재. SC_DATA_LGLS 로 대체됨. 이름이 비슷해 혼동을 부른다'),
 ('rtv_data',       DROP, '없음 (RTV_DATA.h 와 이름만 겹친다)', '구 ECS 잔재. RTV_DATA_LGLS 로 대체됨'),
 ('auto_sc_work',   DROP, '없음', '구 ECS 잔재'),
 ('event_log',      DROP, '없음', '구 ECS 잔재'),
 ('wait_track',     DROP, '없음', '구 ECS 잔재'),
 ('client_pgr_nm',  DROP, '없음 (코드의 CLIENT_PGR_NM 은 COMMON_CODE 의 코드구분 값이다)', '표 자체를 조회하는 곳이 없다'),
 ('dest_def',       DROP, '없음 (구축 스크립트만)', '154행이 들어 있으나 읽는 프로그램이 없다. dest_pos_def 가 현행이다'),
]

d = Document()
st = d.styles['Normal']
st.font.name = '맑은 고딕'
st.font.size = Pt(9)

h = d.add_heading('테이블 사용 현황 조사표 (Rev2)', level=0)
p = d.add_paragraph('LGLS_MCS_IO · 조사일 2026-09-10 · 개정 2026-09-10 (Rev2, 삭제 실행 반영) · 대상 %d개 표' % len(DATA))
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

d.add_heading('1. 조사 방법', level=1)
for t in [
    '판정 기준은 "빌드되는 프로그램의 소스에 그 표를 대상으로 한 SQL 이 있는가" 하나다.',
    '행 수는 판정 근거로 쓰지 않았다. eqp_err_his 는 0행이지만 실제로 쓰는 표다(아직 에러가 없었을 뿐이다).',
    '반대로 dest_def 는 154행이 들어 있지만 읽는 프로그램이 없다.',
    'IO_TASK 의 cThread_CV / cThread_SC / cThread_R 은 프로젝트에서 빌드 제외되어 있어(현장 구성에서 제외) '
    '그 안의 참조는 판정에서 뺐다. 실제로 컴파일되는 설비 스레드는 cThread_SCH 뿐이다.',
    '표 이름이 C++ 클래스 파일명과 겹치는 경우(SC_DATA.h, RTV_DATA.h)가 있어, 단순 이름 검색이 아니라 '
    'FROM / JOIN / INSERT INTO / UPDATE / DELETE FROM 뒤에 오는지로 확인했다.',
]:
    d.add_paragraph(t, style='List Bullet')

d.add_heading('2. 조사 결과 및 처리', level=1)
cols = ['표 이름', '행 수(삭제 직전)', '판정', '쓰는 곳 / 근거', '비고']
tbl = d.add_table(rows=1, cols=len(cols))
tbl.style = 'Table Grid'
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, c in enumerate(cols):
    cell = tbl.rows[0].cells[i]
    cell.text = c
    for r in cell.paragraphs[0].runs:
        r.bold = True

order = {DROP: 0, HOLD: 1, KEEP: 2}
for name, verdict, where, note in sorted(DATA, key=lambda x: (order[x[1]], x[0].lower())):
    row = tbl.add_row().cells
    row[0].text = name
    row[1].text = ROWS.get(name.lower(), '?')
    row[2].text = verdict
    row[3].text = where
    row[4].text = note
    if verdict == DROP:
        for c in row:
            for para in c.paragraphs:
                for r in para.runs:
                    r.font.color.rgb = RGBColor(0xC0, 0, 0)

for i, w in enumerate([Cm(3.2), Cm(1.6), Cm(2.0), Cm(6.4), Cm(6.4)]):
    for row in tbl.rows:
        row.cells[i].width = w

d.add_heading('3. 요약', level=1)
nd = sum(1 for x in DATA if x[1] != KEEP)   # DROP 과 HOLD 는 Rev2 에서 같은 판정이다
nh = 0
nk = sum(1 for x in DATA if x[1] == KEEP)
d.add_paragraph('삭제 %d개 / 유지 %d개   (Rev1 : 삭제 권고 7 + 판단 필요 5 -> 12개 모두 폐지 결정)'
                % (nd, nk))
d.add_paragraph('삭제한 표 : ' + ', '.join(sorted([x[0] for x in DATA if x[1] in (DROP, HOLD)])))
d.add_paragraph('남은 표 : %d개. 테이블 정의서 Rev4 에 전부 실려 있다(문서와 DB 가 1:1 로 맞는다).' % nk)

d.add_heading('4. 삭제 절차 및 실행 기록', level=1)
for i, t in enumerate([
    '이 조사표를 근거로 변경관리(Change Control)를 기안한다 - 대상, 사유, 영향평가, 승인.',
    '삭제 전 점검 : DB_BACKUP/10_점검_표삭제전.sql - 행 수와 외래키 참조를 확인했다. '
    '이 표들을 가리키는 외래키는 없었다.',
    '프로그램 정리 : Client 에서 WC / BCR / 공PLT / WC 로그 화면과 설비 클래스를 걷어내고'
    '(지운 소스는 Ecs\\_removed_20260910\\ 에 보관), WCS_TASK_HOST 에서 무게보고·빈파렛트 요청을 제거했다. '
    '신규 구축 DDL(mssql_schema.sql)과 시뮬레이터 시드 스크립트에서도 뺐다.',
    '삭제 실행 : DB_BACKUP/11_표삭제.sql - 12개 표 삭제. 되돌리기는 13_표복구.sql '
    '(같은 자료형의 빈 표를 다시 만든다. 값은 되살릴 수 없다).',
    '삭제 후 점검 : DB_BACKUP/12_점검_표삭제후.sql - 대상 0행, 남은 표 25개.',
    '회귀 시험 (2026-09-10 11:23~11:28) : 5개 프로그램을 재빌드하여 전체 기동하고 입출고 풀사이클을 '
    '5분 연속 운전했다. 작업 이력 124행 증가, 진행 중 작업 3건 정상 진행, CV/HOST/SCH 하트비트 정상, '
    '새 크래시 리포트와 조회 오류 없음.',
    '문서 개정 : 테이블 정의서 Rev4(남은 25개 표), ECS 화면설계서(폐지 화면 반영), 이 조사표 Rev2.',
], 1):
    d.add_paragraph('%d) %s' % (i, t))

d.add_heading('5. 판단 필요 5건에 대한 결정 (Rev2)', level=1)
for t in [
    'mes_if_log / wc_his / wc_data / bcr_mst / host_empty_plt 5건은 모두 폐지하기로 결정했다. '
    '이 현장에는 MES 연동도, WC(검량)·BCR 설비도, 공파렛트 운용도 없다.',
    '표만 지우면 화면이 오류를 내므로 화면·메뉴·설비 클래스까지 같은 변경으로 함께 걷어냈다. '
    '공PLT 는 메뉴 항목과 Ecs.ini [MENU] 키까지 제거했다(숨김이 아니라 폐지다).',
    '지운 소스는 버리지 않고 Ecs\\_removed_20260910\\ 에 보관했다. 형상관리 이력으로도 추적할 수 있다.',
    '밸리데이션 관점에서 "일부만 정리" 가 가장 불리하므로 지울 것을 한 번에 정했고, '
    '남기는 25개는 남기는 이유를 2장 표의 "쓰는 곳 / 근거" 칸에 적어 두었다.',
]:
    d.add_paragraph(t, style='List Bullet')

d.save(OUT)
print('저장', OUT)
print('삭제 %d / 유지 %d' % (nd, nk))
