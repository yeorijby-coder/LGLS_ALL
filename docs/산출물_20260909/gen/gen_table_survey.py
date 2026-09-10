# -*- coding: utf-8 -*-
# [LGLS 2026-09-10] 테이블 사용 현황 조사표 (밸리데이션 변경관리 첨부용)
#   판정 기준 : "빌드되는 프로그램의 소스에 그 테이블을 대상으로 한 SQL 이 있는가"
#   행 수는 근거로 쓰지 않는다(eqp_err_his 처럼 아직 안 쌓였을 뿐인 표가 있다).
import os, io, sys, datetime
sys.path.insert(0, r'D:\project\LGLS\Renewal\docs\산출물_20260909\gen')
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

OUT = r'D:\project\LGLS\Renewal\docs\산출물_20260909\테이블_사용현황_조사표.docx'

ROWS = {}
for ln in io.open(r'C:\Users\USER\AppData\Local\Temp\rowcounts.txt', encoding='utf-8-sig'):
    ln = ln.strip()
    if '|' in ln:
        a, b = ln.split('|')
        ROWS[a.strip().lower()] = b.strip()

# (테이블, 판정, 쓰는 곳/근거, 비고)
KEEP = ' 유지'
DROP = '삭제 권고'
HOLD = '판단 필요'

DATA = [
 ('job_mst',        KEEP, 'IO_TASK, WCS_TASK_HOST, Client', '진행 중 작업. 시스템의 중심 표'),
 ('job_mst_his',    KEEP, 'DB 트리거 trg_JOB_MST_StatusHis, WCS_TASK_HOST', '작업 상태 이력'),
 ('cv_data',        KEEP, 'WCS_TASK_CV, IO_TASK, Client', '컨베이어 트랙 상태/지시'),
 ('SC_DATA_LGLS',   KEEP, 'IO_TASK, Client', '크레인 상태/지시 (현행)'),
 ('RTV_DATA_LGLS',  KEEP, 'IO_TASK, Client', 'RGV 상태/지시 (현행)'),
 ('eqp_mst',        KEEP, '전 TASK 하트비트, Client 상태바', '설비/프로그램 마스터'),
 ('eqp_ecd_mst',    KEEP, 'Client 설비 대화상자', '설비 에러코드 명칭'),
 ('eqp_err_his',    KEEP, 'WCS_TASK_CV 기록, Client 조회', '0행은 아직 에러가 없었을 뿐이다'),
 ('common_code',    KEEP, '전 프로그램', '공통 코드'),
 ('cell_mst',       KEEP, 'Client, WCS_TASK_HOST', '랙 셀 마스터'),
 ('cell_dtl',       KEEP, 'Client CollectDataList (CELL_MST 와 LEFT JOIN)', '0행이나 조회 대상'),
 ('cv_def_inf',     KEEP, 'Client CvSkinDlg / PanelInfoDlg', '컨베이어 연결 정의'),
 ('cv_flow_def',    KEEP, 'WCS_TASK_CV', '흐름 정의'),
 ('sc_def_inf',     KEEP, 'Client ManualJob (CELL_MST 대체 조회)', '크레인 담당 구역'),
 ('sc_hs_def',      KEEP, 'Client Lib / ScManualRet / ScSkinDlg', '크레인 H/S 정의'),
 ('dest_pos_def',   KEEP, 'IO_TASK, Client', '도착지 정의'),
 ('dead_lock_zone_def', KEEP, 'IO_TASK cDefApi', '교착 구간 정의'),
 ('del_his_setting', KEEP, 'Client 로그 삭제 설정', '이력 보관 주기'),
 ('user_mst',       KEEP, 'Client 로그인', '사용자'),
 ('user_grp',       KEEP, 'Client 권한', '사용자 그룹'),
 ('user_grp_win',   KEEP, 'Client 권한(EcsDoc)', '그룹별 화면 권한'),
 ('chg_lang',       KEEP, 'Client Lang / Lib', '다국어 문구'),
 ('host_if_log',    KEEP, 'WCS_TASK_HOST 기록, Client 조회', '상위 전문 이력'),
 ('wcs_log_pgr',    KEEP, '전 TASK 기록, Client 조회', '프로그램 운전 이력'),
 ('wcs_client_log', KEEP, 'Client 기록/조회', '화면 조작 이력'),
 ('mes_if_log',     HOLD, 'Client [MES_LOG] 조회 화면만', '이 현장은 MES 연동을 하지 않는다. 화면과 함께 정리할지 결정 필요'),
 ('wc_his',         HOLD, 'Client [WC 로그] 조회 화면만', 'WC(검량) 설비가 이 현장에 없다. 화면과 함께 정리할지 결정 필요'),
 ('wc_data',        HOLD, 'Client WcSkinDlg/Wc, WCS_TASK_HOST 조회', 'WC 설비 없음. 화면·조회까지 함께 정리할지 결정 필요'),
 ('bcr_mst',        HOLD, 'Client BCRSkinDlg/Bcr', 'BCR 설비 없음. 화면과 함께 정리할지 결정 필요'),
 ('host_empty_plt', HOLD, 'Client 공PLT 화면, WCS_TASK_HOST 조회', '공PLT 메뉴는 이미 숨김(EMPTYPLT_MENU=0). 기능 폐지 여부 결정 필요'),
 ('sc_data',        DROP, '없음 (C++ 클래스 파일명 SC_DATA.h 와 이름만 겹친다)', '구 ECS 잔재. SC_DATA_LGLS 로 대체됨. 이름이 비슷해 혼동을 부른다'),
 ('rtv_data',       DROP, '없음 (RTV_DATA.h 와 이름만 겹친다)', '구 ECS 잔재. RTV_DATA_LGLS 로 대체됨'),
 ('auto_sc_work',   DROP, '없음', '구 ECS 잔재'),
 ('event_log',      DROP, '없음', '구 ECS 잔재'),
 ('wait_track',     DROP, '없음', '구 ECS 잔재'),
 ('client_pgr_nm',  DROP, '없음 (코드의 CLIENT_PGR_NM 은 COMMON_CODE 의 코드구분 값이다)', '표 자체를 조회하는 곳이 없다'),
 ('dest_def',       DROP, '없음 (구축 스크립트만)', '154행이 들어 있으나 읽는 프로그램이 없다. dest_pos_def 가 현행이다'),
]

d = Document()
st = d.styles['Normal']
st.font.name = '맑은 고딕'
st.font.size = Pt(9)

h = d.add_heading('테이블 사용 현황 조사표', level=0)
p = d.add_paragraph('LGLS_MCS_IO  ·  조사일 2026-09-10  ·  대상 %d개 표' % len(DATA))
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

d.add_heading('1. 조사 방법', level=1)
for t in [
    '판정 기준은 "빌드되는 프로그램의 소스에 그 표를 대상으로 한 SQL 이 있는가" 하나다.',
    '행 수는 판정 근거로 쓰지 않았다. eqp_err_his 는 0행이지만 실제로 쓰는 표다(아직 에러가 없었을 뿐이다).',
    '반대로 dest_def 는 154행이 들어 있지만 읽는 프로그램이 없다.',
    'IO_TASK 의 cThread_CV / cThread_SC / cThread_R 은 프로젝트에서 빌드 제외되어 있어(현장 구성에서 제외) '
    '그 안의 참조는 판정에서 뺐다. 실제로 컴파일되는 설비 스레드는 cThread_SCH 뿐이다.',
    '표 이름이 C++ 클래스 파일명과 겹치는 경우(SC_DATA.h, RTV_DATA.h)가 있어, 단순 이름 검색이 아니라 '
    'FROM / JOIN / INSERT INTO / UPDATE / DELETE FROM 뒤에 오는지로 확인했다.',
]:
    d.add_paragraph(t, style='List Bullet')

d.add_heading('2. 조사 결과', level=1)
cols = ['표 이름', '행 수', '판정', '쓰는 곳 / 근거', '비고']
tbl = d.add_table(rows=1, cols=len(cols))
tbl.style = 'Table Grid'
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, c in enumerate(cols):
    cell = tbl.rows[0].cells[i]
    cell.text = c
    for r in cell.paragraphs[0].runs:
        r.bold = True

order = {DROP: 0, HOLD: 1, KEEP: 2}
for name, verdict, where, note in sorted(DATA, key=lambda x: (order[x[1]], x[0].lower())):
    row = tbl.add_row().cells
    row[0].text = name
    row[1].text = ROWS.get(name.lower(), '?')
    row[2].text = verdict
    row[3].text = where
    row[4].text = note
    if verdict == DROP:
        for c in row:
            for para in c.paragraphs:
                for r in para.runs:
                    r.font.color.rgb = RGBColor(0xC0, 0, 0)

for i, w in enumerate([Cm(3.2), Cm(1.6), Cm(2.0), Cm(6.4), Cm(6.4)]):
    for row in tbl.rows:
        row.cells[i].width = w

d.add_heading('3. 요약', level=1)
nd = sum(1 for x in DATA if x[1] == DROP)
nh = sum(1 for x in DATA if x[1] == HOLD)
nk = sum(1 for x in DATA if x[1] == KEEP)
d.add_paragraph('삭제 권고 %d개 / 판단 필요 %d개 / 유지 %d개' % (nd, nh, nk))
d.add_paragraph('삭제 권고 : ' + ', '.join(sorted([x[0] for x in DATA if x[1] == DROP])))
d.add_paragraph('판단 필요 : ' + ', '.join(sorted([x[0] for x in DATA if x[1] == HOLD])))

d.add_heading('4. 삭제 절차 (제안)', level=1)
for i, t in enumerate([
    '이 조사표를 근거로 변경관리(Change Control)를 기안한다 - 대상, 사유, 영향평가, 승인.',
    '삭제 전 점검 스크립트로 행 수와 참조를 다시 확인한다.',
    '삭제 스크립트를 실행한다. 현장은 백업/복원을 할 수 없으므로 같은 자료형으로 되돌리는 복구 스크립트를 함께 둔다.',
    '삭제 후 회귀 시험 : 입고·출고 풀사이클, 수동작업, 각 로그 조회 화면. 시험 기록을 남긴다.',
    '테이블 정의서를 개정하고 개정이력에 삭제 내역을 적는다.',
], 1):
    d.add_paragraph('%d) %s' % (i, t))

d.add_heading('5. 판단이 필요한 것에 대한 의견', level=1)
for t in [
    'mes_if_log / wc_his / wc_data / bcr_mst 는 "표는 살아 있고 화면도 있으나 이 현장에는 그 설비·연동이 없는" 것들이다. '
    '표만 지우면 화면이 오류를 내므로, 지우려면 화면(메뉴)까지 함께 정리해야 한다.',
    'host_empty_plt(공PLT)는 이미 메뉴에서 숨겨 두었다(Ecs.ini [MENU] EMPTYPLT_MENU=0). '
    '기능을 완전히 폐지할지, 숨김 상태로 남길지 결정이 필요하다.',
    '밸리데이션 관점에서는 "일부만 정리" 가 가장 불리하다. 지울 것과 남길 것을 한 번에 정하고, '
    '남기는 것은 남기는 이유를 문서에 적어 두는 편이 다음 심사에서 반복 질문을 막는다.',
]:
    d.add_paragraph(t, style='List Bullet')

d.save(OUT)
print('저장', OUT)
print('삭제 권고 %d / 판단 필요 %d / 유지 %d' % (nd, nh, nk))
