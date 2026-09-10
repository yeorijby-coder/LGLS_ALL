# -*- coding: utf-8 -*-
# [LGLS 2026-09-10] 화면설계서의 캡처 그림을 새 캡처로 갈아 끼운다.
#   자리(위치/칸 크기)는 그대로 두고 새 그림의 가로세로비에 맞춰 칸 안에 넣는다.
#   사용 : python pptx_replace_shot.py <슬라이드번호> <새그림경로> [슬라이드번호 새그림경로 ...]
import os, shutil, sys
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

PPTX = r"D:\project\LGLS\Renewal\docs\산출물_20260909\06_ECS_화면설계서.pptx"

args = sys.argv[1:]
assert args and len(args) % 2 == 0, '슬라이드번호 그림경로 짝으로 넘겨라'
jobs = [(int(args[i]), args[i + 1]) for i in range(0, len(args), 2)]

if not os.path.exists(PPTX + '.bak_shot'):
    shutil.copy(PPTX, PPTX + '.bak_shot')

prs = Presentation(PPTX)
for idx, img in jobs:
    assert os.path.exists(img), img
    slide = prs.slides[idx]
    pics = [sh for sh in slide.shapes if sh.__class__.__name__ == 'Picture']
    assert pics, '슬라이드 %d 에 그림이 없다' % idx
    pic = max(pics, key=lambda s: s.width * s.height)   # 가장 큰 그림이 화면 캡처
    L, T, W, H = pic.left, pic.top, pic.width, pic.height

    iw, ih = Image.open(img).size
    scale = min(W / float(iw), H / float(ih))
    nw, nh = int(iw * scale), int(ih * scale)
    nl = L + (W - nw) // 2
    nt = T + (H - nh) // 2

    pic._element.getparent().remove(pic._element)
    slide.shapes.add_picture(img, Emu(nl), Emu(nt), Emu(nw), Emu(nh))
    print('슬라이드 %2d  %-24s  %dx%d -> 칸 %dx%d 안에 %dx%d'
          % (idx, os.path.basename(img), iw, ih, W, H, nw, nh))

prs.save(PPTX)
print('저장', PPTX)
