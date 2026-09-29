# -*- coding: utf-8 -*-
u"""2026-09-29 적용 안내서 (Word).

   사용자 지시 - 문서는 md 가 아니라 ★WORD★ 로 만든다(md 를 볼 장치가 없다).
   문서 스타일은 docs/산출물_20260927/gen/gen_common.py 를 그대로 쓴다.

   실행 : python gen_apply_docx.py
"""
import os, sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, 'docs', '산출물_20260927', 'gen'))

from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import gen_common as G

SHOT = os.path.join(HERE, 'shots')
NAVY, GRAY = G.NAVY, G.GRAY
RED = RGBColor(0xB0, 0x2A, 0x20)


def code(d, lines, size=9):
    u"""로그·SQL·코드 조각 - 고정폭에 옅은 바탕."""
    t = d.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    c = t.rows[0].cells[0]
    G.shade(c, 'F4F5F7')
    c.text = ''
    for i, ln in enumerate(lines):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(ln)
        r.font.name = 'D2Coding'
        r.font.size = Pt(size)
    d.add_paragraph()


def note(d, text):
    p = d.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    r = p.add_run('※ ' + text)
    r.font.size = Pt(9.5)
    r.font.color.rgb = RED
    return p


G.DATE = '2026-09-29'
G.VER = '1.0'
G.BASE_HISTORY = []
G.FINAL_NOTE = '최초 작성 - 2026-09-29 적용분 (출고 완료 지연, 신호등 PLC, EQP_MST 옵션, 범례 정리, 4K 대응)'

d = G.new_doc('2026-09-29 적용 안내서',
              '출고 완료 지연 · 신호등 PLC · EQP_MST 옵션 · 범례 정리 · 4K 대응')

# ════════════════════════════════════════════════════════════════
d.add_heading('1. 이 폴더는 무엇인가', 1)
G.para(d, '원격지(현장) 서버의 프로그램이 9월 23일 이전 판이라, 2026-09-29 에 작업한 것만 따로 모은 폴더다. '
          '현장 파일을 덮어쓰기 전에 반드시 백업한다.')
G.table(d, ['폴더', '내용'], [
    ['SQL', 'EQP_MST 정리·복구 SQL 3종'],
    ['Client', '운전 화면 Ecs.exe (Release 판) + 추가할 설정'],
    ['TASK/IO_TASK', '스케줄러 IO_TASK_SEMI_FINISH.exe + 추가할 설정'],
    ['TASK/HOST_TASK', '상위 통신 TASK_LFC10_G1_ECSCOM.exe + 추가할 설정'],
], widths=[4.5, 11.5])

d.add_heading('1.1 ini 는 통째로 넣지 않는다', 2)
G.para(d, '현장 ini 에는 DB 주소와 상위 IP 가 들어 있어 덮어쓰면 접속이 끊긴다. '
          '각 폴더의 「추가할_설정_*.txt」 에 있는 줄만 해당 섹션에 붙여넣는다.')
note(d, '붙여넣지 않아도 코드 기본값으로 동작한다. 즉 exe 만 바꿔도 돌아간다.')

d.add_heading('1.2 Client 는 Release 판이다', 2)
G.para(d, '현장 배포 폴더(EXE_NEWUI\\WCS_CLIENT)의 DLL 은 Release 로 빌드된 것이라, '
          'Debug 로 빌드한 Ecs.exe 를 넣으면 다음 오류가 난다.')
code(d, ['프로시저 시작 지점 ?Dump@?$CList@PAVCDciControl@@PAV1@@@UBEXAAVCDumpContext@@@Z 을(를)',
         'DLL ...\\Ecs.exe 에서 찾을 수 없습니다.'])
G.para(d, '이 심볼(CList<CDciControl*>::Dump)은 MFC 디버그 빌드에서만 만들어진다. '
          'Debug exe 가 그것을 요구하는데 Release DciLib.dll 에는 없어서 나는 오류다. '
          '이 폴더의 Ecs.exe 는 Release 로 빌드했으므로 DLL 과 짝이 맞는다.')
note(d, 'Ecs.exe 를 바꿀 때 DciLib.dll / EcsLib.dll / XmlLib.dll 도 같은 판으로 함께 바꾼다.')

d.add_page_break()

# ════════════════════════════════════════════════════════════════
d.add_heading('2. 출고 완료가 1~2분 걸리던 것', 1)
G.para(d, '대상 : TASK/IO_TASK', bold=True)

d.add_heading('2.1 증상', 2)
G.para(d, '자동 출고 작업이 끝났는데 작업이 상태 15(C/V 구동중)에 1~2분 머물다 완료됐다. '
          '출고대 준비 신호는 운전 화면에서 확실히 ON 이었다.')

d.add_heading('2.2 출고 완료는 원래 네 단계로 판정한다', 2)
G.para(d, 'IO_TASK 의 CompleteCV() 는 "출고 화물이 정말 출고대에 도착해서 나갔는가" 를 '
          '다음 순서로 확인한다. 위에서 성립하면 아래는 보지 않는다.')
G.table(d, ['순서', '무엇을 보나', '판정'], [
    ['①', '출고대 신호 (CV_DATA.RET_READY_RD = 1)',
     '설비가 "출고대에 화물이 놓였다" 를 직접 알린 것 → 곧바로 완료'],
    ['②', '실도착 관측 (센서 ON + 트래킹이 내 작업번호)',
     '신호를 놓쳤어도 내 화물이 출고대에 있는 것을 직접 봤다 → 기록해 두고 대기'],
    ['③', '배출 확인 (②를 본 뒤 트랙이 비었다)',
     '내 화물이 나간 것을 확인 → 완료'],
    ['④', '최후 안전망 (트래킹이 전 트랙에 없는 상태가 60초)',
     '아무것도 못 봤지만 이미 나간 것으로 인정 → 완료(로그에 경고)'],
], widths=[1.5, 6.5, 8.0])
G.para(d, '①이 가장 빠르고 정확하다. ④는 "어쩔 수 없을 때" 의 마지막 수단이라 60초를 기다린다.')

d.add_heading('2.3 왜 ①을 못 타고 ④까지 갔나 — 문턱이 둘 있었다', 2)
G.para(d, '신호가 ON 이어도 ①에 닿기까지 조건이 두 개 더 있었다.')

G.para(d, '문턱 ㉠ — 반출 시퀀스 가드', bold=True, size=10.5)
code(d, ['// 종전 : 신호를 보기도 전에 여기서 걸렀다',
         'if (OutSeqPending(luggNo)) continue;   // 반출(RTV) 시퀀스가 진행 중이면 판정 자체를 건너뜀'])
G.para(d, 'RTV 반출 시퀀스의 잔여 상태가 남아 있으면 신호를 볼 기회조차 없었다.')

G.para(d, '문턱 ㉡ — 트래킹이 내 번호여야 함', bold=True, size=10.5)
code(d, ['bool bMine = (cvLugg == "" || cvLugg == "0" || cvLugg == "0000" || cvLugg == luggNo);',
         'if (cvRetRdy == "1" && bMine) { ...출고 완료... }'])
G.para(d, '출고대 트랙의 트래킹 값(LUGG_NO_RD)이 내 작업번호이거나 비어 있어야 신호를 인정했다. '
          '그런데 앞 화물이 나간 뒤에도 설비 트래킹에 그 번호가 남는 일이 있다. '
          '그러면 cvLugg 는 "남의 번호" 가 되어 bMine 이 거짓이 된다.')

G.para(d, '두 문턱 중 하나만 걸려도 이렇게 흘러간다.', bold=True, size=10.5)
G.bullets(d, [
    '① 불성립 — 신호는 ON 인데 인정받지 못한다',
    '② 불성립 — 트래킹이 남의 번호라 "내 화물 실도착" 도 성립하지 않는다',
    '③ 불성립 — ②를 못 봤으니 배출 확인 경로로 못 간다',
    '④ 로 떨어진다 — 트래킹이 사라진 뒤 60초를 더 기다린 끝에 완료',
])
G.para(d, '설비가 트래킹을 지우는 시점이 제각각이라 60초 계측 시작도 늦어진다. '
          '이것이 현장에서 본 "1~2분" 의 정체다.')

d.add_heading('2.4 고친 내용', 2)
G.para(d, '출고대 신호를 최우선 근거로 삼는다. 다만 앞 화물이 아직 안 나간 경우를 조기 완료하면 안 되므로, '
          '"남의 번호" 가 살아 있는 작업인지 아닌지를 갈라서 본다.')
G.table(d, ['출고대 트래킹 상태', '종전', '지금'], [
    ['비었거나 내 번호', '신호로 완료', '신호로 완료 (같음)'],
    ['남의 번호인데 그 작업이 JOB_MST 에 없다',
     '완료 못 함 → 60초 대기', '유령 트래킹으로 보고 신호로 완료'],
    ['남의 번호이고 그 작업이 살아 있다',
     '완료 못 함 → 60초 대기', '앞 화물이 아직 안 나간 것 → 종전대로 기다림'],
], widths=[6.0, 5.0, 5.0])
G.para(d, '또 신호로 완료가 서면 반출 시퀀스 가드(문턱 ㉠)를 건너뛴다. '
          '신호 자체가 "출고대에 놓였다" 는 설비의 직접 통보라 시퀀스 잔여 상태보다 믿을 만하다.')
note(d, '신호를 못 본 경우의 ②③④ 경로는 한 줄도 바꾸지 않았다. 여전히 못 보면 1분 뒤 완료된다.')

d.add_heading('2.5 로그로 어느 경로인지 알 수 있다', 2)
code(d, ['[SCH][CV] 작업 0123 출고대 126 출고대 신호 ON - 출고 완료',
         '[SCH][CV] 작업 0123 출고대 126 트래킹에 끝난 작업 0119 가 남아 있습니다(유령) - 신호로 완료',
         '[SCH][CV] 작업 0123 출고대 126 실도착 관측 (신호/배출 대기)',
         '[SCH][CV] 작업 0123 출고대 126 배출 확인 - 출고 완료',
         '[SCH][CV] 작업 0123 출고대 126 - 출고대 신호도 실도착도 관측하지 못했습니다. ...'])
G.para(d, '첫 줄이나 둘째 줄이 나오면 신호로 곧바로 완료된 것이다. '
          '마지막 줄이 자주 나오면 설비 신호 배선을 점검할 일이다.')

d.add_heading('2.6 설정', 2)
G.table(d, ['키', '기본', '설명'], [
    ['OUT_DONE_BY_SIGNAL', '1',
     '1 = 위 내용대로. 0 = 전부 종전 동작(트래킹 일치 + 반출 시퀀스 완료를 요구)'],
    ['OUT_SIGNAL_MISS_SEC', '60',
     '④ 최후 안전망의 유예(초). 요청대로 1분을 유지했다. 5 미만은 5 로 올린다'],
], widths=[5.0, 1.8, 9.2])
G.para(d, '둘 다 ENV_IOSCH.INI 의 [CNF] 섹션에 넣는다. 호출할 때마다 읽으므로 고치면 곧 반영된다.')

d.add_page_break()

# ════════════════════════════════════════════════════════════════
d.add_heading('3. 신호등 아래 글자 EQP → PLC', 1)
G.para(d, '대상 : Client', bold=True)
G.para(d, '리본 [통신] 그룹의 셋째 칸 이름을 PLC 로 바꿨다.')
G.image(d, os.path.join(SHOT, 'ribbon_plc.png'), 16.0,
        '리본 [통신] 그룹 - WMS1 / WMS2 / PLC / SCH')
G.para(d, 'EQP_MST 의 EQP_TYP 값은 아직 \'EQP\' 다. 나중에 \'PLC\' 로 바꿀 것을 감안해 '
          '운전 화면의 조회는 \'EQP\' · \'PLC\' · \'CV\' 를 모두 같은 칸으로 본다. '
          '값을 바꿔도 화면은 그대로 돌고, 바꾸지 않아도 그대로 돈다.')

d.add_page_break()

# ════════════════════════════════════════════════════════════════
d.add_heading('4. EQP_MST 행을 줄여도 / 안 줄여도 동작', 1)
G.para(d, '대상 : TASK/HOST_TASK, SQL', bold=True)

d.add_heading('4.1 먼저 - CLR20r3 의 원인에 대해', 2)
G.para(d, '행 삭제가 원인이라고 보셨는데, EQP_MST 를 쓰는 곳을 모두 확인한 결과 그렇게 보기 어렵다.')
G.table(d, ['프로그램', '하는 일', '행이 0개일 때'], [
    ['HOST_TASK CCliWork', 'SELECT 후 건수 분기', 'iCnt == 0 분기가 있어 정상 동작'],
    ['HOST_TASK frmMain', '통신상태 UPDATE', '고친 행 0 - 예외 없음'],
    ['WCS_TASK_CV', '하트비트 UPDATE', '고친 행 0 - 예외 없음'],
    ['IO_TASK', 'UPDATE 후 없으면 INSERT', '스스로 행을 만든다'],
], widths=[4.5, 4.5, 7.0])
G.para(d, '행이 없어서 나는 증상은 운전 화면 신호등이 미접속으로 보이는 것이지 프로세스 종료가 아니다. '
          'CLR20r3 는 작업 스레드의 미처리 예외로 보이며, 그에 대한 안전망은 이미 넣었다(커밋 ff7d119). '
          '스레드 하나가 죽어도 프로세스는 살아남는다. ALL_TASK 를 쓰지 않기로 하셨으니 위험 자체도 크게 줄었다.')

d.add_heading('4.2 그래도 요청대로 옵션을 넣었다', 2)
G.table(d, ['키', '기본', '설명'], [
    ['EQP_MST_AUTO_ROW', '1',
     '1 = 5초 주기 통신상태 기록에서 고친 행이 없으면 그 행을 스스로 만든다. '
     '0 = UPDATE 만 한다(행이 없으면 신호등은 미접속으로 보인다)'],
], widths=[5.0, 1.8, 9.2])
G.para(d, 'EcsComA.ini 의 [Host] 섹션에 넣는다. '
          '행을 줄여도 신호등이 정상으로 보이고, 줄이지 않아도 종전대로 돈다.')

d.add_heading('4.3 SQL 3종', 2)
G.table(d, ['파일', '하는 일'], [
    ['260929_EQP_MST_01_원상복구.sql',
     '설비통신 행을 지우고 CV(15)·SC(5)·RTV(1) 21건을 되돌린다'],
    ['260929_EQP_MST_02_설비통신행_INSERT.sql',
     '설비통신 행 1건을 넣는다(이미 있으면 주소만 맞춘다)'],
    ['260929_EQP_MST_03_설비행_정리.sql',
     'CV/SC/RTV 를 지워 4행만 남긴다'],
], widths=[7.5, 8.5])
G.bullets(d, [
    '02 안의 @EQP_TYPE 한 줄만 \'PLC\' 로 고치면 종류 값까지 바꿀 수 있다',
    '@WH_TYP 은 현장 창고 구분에 맞춘다 (기본 \'10\')',
    '세 파일 모두 여러 번 돌려도 안전하다',
])

d.add_page_break()

# ════════════════════════════════════════════════════════════════
d.add_heading('5. 범례', 1)
G.para(d, '대상 : Client', bold=True)

d.add_heading('5.1 뺀 항목', 2)
G.para(d, '범례 판넬과 범례 설정 창 양쪽에서 없앴다. 설정 창은 리소스에서 실제로 지웠다.')
G.bullets(d, [
    '랙투랙 · 호기간 이동',
    '반자동 랙투랙 · 반자동 호기간 이동',
    '입고 HS · 출고 HS',
    '통신두절 · 작업번호 있음',
    '레일 관련 전부 (입고 금지 · 출고 금지 · 입출고 정지 · 레일 에러 · 작업중)',
    '창고간 이동 (이미 숨겨져 있어 같이 정리)',
])
note(d, '색 자체는 그대로 쓰인다. 범례에서만 뺀 것이고, 랙투랙 작업이 들어오면 여전히 그 색으로 칠한다.')

d.add_heading('5.2 화물만 있으면 살색이 되던 이유', 2)
G.para(d, 'TrackInfo.cpp 의 GetCvColor() 맨 끝에 이런 줄이 있었다.')
code(d, ['//화물D',
         'if (V_JOB_TYP_RD == "0" && V_SENSOR0_DATA_RD == "1")',
         '    return WHEAT;          // RGB(255,225,192) = 살색'])
G.para(d, '작업구분이 0(작업 없음)이고 화물감지 센서만 켜져 있으면 살색으로 칠했다. '
          '구 ECS 부터 있던 "작업은 없는데 화물은 놓여 있다" 표시인데, '
          '범례에 없는 색이라 화면에서만 튀었다.')
G.para(d, '요청대로 이 줄을 없앴다. 이제 화물만 있는 트랙은 빈 트랙과 같은 회색이고, '
          '화물이 있다는 것은 트랙 네 모서리의 검은 점 4개로만 보인다. '
          '이 표시는 CDciTrackCtrl 이 센서값(m_bExist)을 그대로 받아 그리므로 따로 손댈 것이 없었다.')

d.add_heading('5.3 범례에 「화물 감지」 를 넣었다', 2)
G.para(d, '색으로 칠하지 않게 됐으니 범례에도 그 표시를 넣어야 뜻이 통한다. '
          '색칩이 아니라 트랙과 같은 모양(회색 바탕 + 네 모서리 검은 점)으로 그린다.')
G.image(d, os.path.join(SHOT, 'legend_panel.png'), 15.0,
        '메인 화면 왼쪽 범례 판넬 - C/V 상태 그룹에 「화물 감지」')

d.add_heading('5.4 숨김 체크박스', 2)
G.para(d, '범례 설정 창의 색상 칸 옆에 [숨김] 체크박스를 뒀다. '
          '체크하고 [저장] 하면 메인 화면 왼쪽 범례 판넬에서 그 줄이 빠지고 나머지가 당겨진다. '
          'Ecs.ini 의 [USER] LEGEND_HIDE_* 에 저장되므로 다시 켜도 유지된다.')
G.image(d, os.path.join(SHOT, 'legend_dlg.png'), 16.0,
        '범례 설정 창 - 4그룹 12칸, 항목마다 [숨김]')

d.add_heading('5.5 도면 안의 「범 례」 표는 없앴다', 2)
G.para(d, '메인 화면 도면 위에 레이아웃 도형으로 박힌 범례표가 또 하나 있었다. '
          '왼쪽 판넬과 내용이 어긋나 있어(랙투랙·HS·레일이 그대로 남아 있었다), '
          '현장 배포본(EXE_NEWUI\\WCS_CLIENT)의 레이아웃으로 교체했다.')
G.para(d, '교체한 파일 : EcsLayout1.xml · EcsLayout2.xml · EcsLayout3.xml '
          '(종전 파일은 같은 폴더에 .bak_0929 로 남겨 두었다)')
G.para(d, '이 레이아웃에는 작업대 명칭(제품 입고대 · 원부자재불출대 · 피킹 출고/입고 작업대 · '
          '외부 입고 전용 입출고대)도 함께 들어 있다.')

d.add_page_break()

# ════════════════════════════════════════════════════════════════
d.add_heading('6. 4K(3840×2160) 대응', 1)
G.para(d, '대상 : Client', bold=True)
G.para(d, '하나의 Client 로 1920×1080 과 3840×2160 을 모두 받는다. 화면 높이를 보고 배율을 정한다.')
G.table(d, ['화면 높이', '배율'], [
    ['2160 이상', '200%'], ['1800 이상', '175%'], ['1600 이상', '150%'],
    ['1300 이상', '125%'], ['그 미만', '100%'],
], widths=[5.0, 4.0])
G.para(d, '리본 [통신] 신호등의 치수·글꼴과 왼쪽 범례 판넬의 치수·글꼴이 이 배율을 따른다.')

d.add_heading('6.1 신호등이 2줄(위 3 / 아래 1)로 접히던 이유', 2)
G.para(d, 'MFC 리본은 패널 안을 「행」 단위로 채운다. 신호등 칸이 62×56 픽셀에 박혀 있어서 '
          '4K 에서는 상대적으로 작아지고, 그러면 한 패널에 두 행이 들어갈 여유가 생겨 '
          'MFC 가 4칸을 3+1 로 나눠 담았다.')
G.para(d, '칸을 화면에 맞게 키우면 한 행만 들어가 다시 한 줄이 된다. '
          '즉 배율을 키우는 것이 곧 이 문제의 해결이다.', bold=True)

d.add_heading('6.2 설정', 2)
G.table(d, ['키', '기본', '설명'], [
    ['UI_SCALE', '0', '0 = 자동(위 표대로). 100~300 = 그 값을 그대로 쓴다'],
], widths=[5.0, 1.8, 9.2])
G.para(d, 'Ecs.ini 에 [DISPLAY] 섹션을 만들어 넣는다. 고치면 3초 안에 반영된다(다시 그리는 시점에). '
          '4K 에서 여전히 접히면 값을 더 키운다.')
note(d, '4K 화면이 없어 이 항목만은 코드 경로로만 확인했다. 현장에서 값 조정이 필요할 수 있다.')

d.add_page_break()

# ════════════════════════════════════════════════════════════════
d.add_heading('7. 적용 순서', 1)
G.numbered(d, [
    '현장 프로그램과 ini 를 백업한다.',
    'TASK 를 정지한다.',
    'IO_TASK / HOST_TASK 의 exe 를 교체한다.',
    '각 폴더의 「추가할_설정_*.txt」 에 있는 줄을 현장 ini 의 해당 섹션에 붙여넣는다. (건너뛰어도 된다)',
    'TASK 를 기동한다. 기동 순서는 EQP → IO → HOST.',
    'Client 는 Ecs.exe 와 DciLib/EcsLib/XmlLib.dll 을 함께 교체하고, 레이아웃 3종도 함께 넣는다.',
    'SQL 은 필요한 것만 돌린다. 행을 줄일 때 02 → 03, 되돌릴 때 01.',
])

d.add_heading('7.1 확인하면 좋은 것', 2)
G.table(d, ['항목', '무엇을 보나'], [
    ['출고 완료', 'IO_TASK 로그에 「출고대 신호 ON - 출고 완료」 가 나오는가, 15 에 머무는 시간이 짧아졌는가'],
    ['신호등', '리본 [통신] 셋째 칸이 PLC 로 보이는가'],
    ['4K TV', '신호등 4칸이 한 줄로 나오는가, 글자 크기가 맞는가'],
    ['범례', '뺀 항목이 안 보이는가, 「화물 감지」 가 보이는가, [숨김] 체크가 반영되는가'],
    ['도면', '도면 위의 옛 「범 례」 표가 사라졌는가, 작업대 명칭이 보이는가'],
], widths=[3.5, 12.5])

out = os.path.join(HERE, '2026-09-29_적용안내서.docx')
d.save(out)
print('생성 :', out)
